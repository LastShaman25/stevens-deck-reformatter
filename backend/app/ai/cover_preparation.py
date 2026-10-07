"""Bounded semantic artwork preparation for the native PDF reformatter.

Only the opening is classified. The native formatter still builds the deck, and
the independent ordered QA reviews every output against the original PDF.
"""
from pathlib import Path
import json

from pptx import Presentation

from slide_engine import templates, template_policy as T
from . import providers, source_decisions
from ..qa import render_verify


def needed(sess):
    pages = (getattr(sess, 'pdf_import', None) or {}).get('page_evidence', [])
    return bool(pages and pages[0].get('independent_cover_graphics'))


def run(sess, directory, progress):
    details = {'status': 'not_needed', 'calls': [], 'source_slides': [0]}
    if not needed(sess):
        return {}, details
    if not providers.role_config('element_roles')['configured']:
        details.update(status='not_configured', message='Cover artwork classification is unavailable. '
                       'Uncertain source artwork was retained for review.')
        return {}, details
    try:
        progress(stage='source_decisions', output_slide=0)
        prs = Presentation(templates.path())
        for sid, slide in list(zip(prs.slides._sldIdLst, prs.slides)):
            if slide.slide_layout.name == T.OPENING_LAYOUT:
                continue
            prs.part.drop_rel(sid.rId)
            prs.slides._sldIdLst.remove(sid)
        if not prs.slides:
            prs.slides.add_slide(next(l for l in templates.layouts(prs) if l.name == T.OPENING_LAYOUT))
        reference = Path(directory, 'cover-reference.pptx')
        prs.save(reference)
        rendered = render_verify.check(reference, Path(directory, 'cover-reference-render'))
        if rendered.get('status') == 'error' or not rendered.get('pages'):
            raise ValueError('Opening template reference could not be rendered.')
        images = [(f'APPROVED TEMPLATE: {T.OPENING_LAYOUT}', Path(page['png']))
                  for page in rendered['pages']]

        def generate(role, system, payload, images, max_tokens):
            result = providers.generate(role, system, payload, images=images, max_tokens=max_tokens)
            details['calls'].append({'role': 'source_decision' if payload.get('stage') == 'source_decisions'
                                    else 'logo_extraction_review', 'source_slide': 0,
                                    **{key: value for key, value in result.items() if key != 'data'}})
            # Processing-only evidence mirrors the full redesign pipeline. It
            # expires with this upload and makes rejected decisions diagnosable
            # without repeating paid calls; no keys or request headers are saved.
            Path(directory, f'cover-call-{len(details["calls"]):03d}.json').write_text(
                json.dumps({'request_stage': payload.get('stage'), 'source_slide': 0,
                            'validation_error': payload.get('validation_error'), **result}), encoding='utf-8')
            return result

        decisions, _ = source_decisions.run(sess, progress, generate=generate,
                                           template_images=images, indices={0})
        details['status'] = 'completed'
        return decisions, details
    except Exception as exc:
        # Preserve the reviewable native candidate, without authorizing any
        # partial or invalid removal. The caller exposes this review finding.
        details.update(status='error', message=f'Cover artwork preparation failed: {type(exc).__name__}: {str(exc)[:1000]}')
        return {}, details
