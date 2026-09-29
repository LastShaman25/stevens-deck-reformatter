"""Versioned review, generation and download API."""
import hashlib
import os
from pathlib import Path
from typing import Literal
from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import FileResponse, Response
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel, ConfigDict
from pptx import Presentation
from . import sessions, grounded, rendering, generations
from .ai import providers, optional_review

router = APIRouter(prefix='/api')


def session(sid):
    sess = sessions.get(sid)
    if not sess or not sess.analysis:
        raise HTTPException(404, 'Session not found or expired.')
    return sess


def slide_index(sess, index):
    if index < 0 or index >= sess.analysis['slide_count']:
        raise HTTPException(404, 'Slide not found.')


def capabilities():
    return {'libreoffice': rendering.libreoffice_available(), 'renderer': rendering.renderer_info(),
            'vercel':providers.configured('vercel'),
            'openai':providers.configured('openai'),
            'gemini': providers.configured('gemini'), 'claude': providers.configured('anthropic'),
            'ai':providers.capabilities()}


def info(sess):
    return {'session_id': sess.id, 'name': sess.original_name, 'expires_at': sess.expires,
            'workflow': sess.workflow, **sess.analysis,
            'capabilities': capabilities(), 'revisions': sess.revisions,
            'revision_version': sess.revision_version, 'generation': generations.public(sess.generation),
            'preview_generation': generations.public(next((r for r in reversed(list(sess.history.values()))
                if r.get('candidate_sha256') and r.get('state') != 'checking'), None)),
            'ai_check': sorted(sess.ai_check), 'ai_results': sess.ai_results, 'benchmarked': sess.benchmarked,
            'generated': bool(sess.generation and sess.generation['state'] == 'ready')}


@router.get('/health')
def health():
    return {'ok': True, 'capabilities': capabilities()}


@router.post('/sessions')
async def create_session(file: UploadFile = File(...)):
    if not (file.filename or '').lower().endswith(('.pptx','.pdf')):
        raise HTTPException(400, 'Please upload a .pptx or .pdf file.')
    data = await file.read(sessions.MAX_UPLOAD_BYTES + 1)
    if len(data) > sessions.MAX_UPLOAD_BYTES:
        raise HTTPException(413, 'File too large (60 MB max).')
    return await run_in_threadpool(analyze_upload, file.filename, data)


def analyze_upload(filename, data):
    sess = sessions.create()
    try:
        with sessions.job(sess):
            sess.original_name = os.path.basename(filename)
            is_pdf=filename.lower().endswith('.pdf')
            if is_pdf:
                from .pdf_import import convert
                Path(sess.source_pdf).write_bytes(data)
                sess.pdf_import=convert(sess.source_pdf,sess.source_path,lambda i:sess.preview_path(i,'before'))
            else: Path(sess.source_path).write_bytes(data)
            import zipfile
            with zipfile.ZipFile(sess.source_path) as archive:
                entries = archive.infolist()
                if len(entries)>20000 or sum(item.file_size for item in entries)>300*1024*1024:
                    raise ValueError('Expanded presentation exceeds the processing limit.')
            prs = Presentation(sess.source_path)
            if not len(prs.slides):
                raise ValueError('Presentation contains no slides.')
            if len(prs.slides)>100:
                raise ValueError('This processing workflow supports at most 100 source slides.')
            sess.analysis = grounded.analyze(sess.source_path)
            # Source previews are optional. Their absence is visible, never QA evidence.
            try:
                if not is_pdf:
                    rendering.render_to_pdf(sess.source_path, sess.source_pdf)
                    rendering.rasterize_pdf(sess.source_pdf, lambda i: sess.preview_path(i, 'before'))
            except Exception:
                pass
    except Exception as exc:
        sessions.delete(sess.id)
        raise HTTPException(400, f'Could not analyze presentation: {type(exc).__name__}: {exc}')
    return info(sess)


@router.get('/sessions/{sid}')
def get_session(sid: str):
    return info(session(sid))


@router.get('/sessions/{sid}/slides/{index}/preview')
def preview(sid: str, index: int, variant: Literal['before','after']='after', generation_id: str | None=None):
    sess = session(sid)
    try:
        with sessions.read_job(sess):
            if variant == 'before':
                slide_index(sess, index)
                path = Path(sess.preview_path(index, 'before'))
            else:
                # Historical images are reference-only. Downloads and decisions
                # still require the current generation and revision identity.
                record = sess.history.get(generation_id)
                if not record or record.get('state') == 'checking':
                    raise HTTPException(409, 'No matching generated preview.')
                if not Path(record['candidate']).is_file() or generations.sha256(record['candidate']) != record.get('candidate_sha256'):
                    raise HTTPException(409, 'Generated artifact changed.')
                count = (record.get('report') or {}).get('slide_count', 0)
                if index < 0 or index >= count:
                    raise HTTPException(404, 'Output slide not found.')
                path = Path(record['directory'], 'render', f'slide-{index}.png')
                for item in record.get('output_manifest', []):
                    if Path(item['image']) == path and (not path.is_file() or generations.sha256(path) != item['sha256']):
                        raise HTTPException(409, 'Rendered artifact changed.')
            if not path.exists():
                raise HTTPException(404, 'Rendered preview unavailable.')
            return Response(path.read_bytes(), media_type='image/png', headers={'Cache-Control':'no-store'})
    except ValueError as exc:
        raise HTTPException(409, str(exc))


@router.post('/sessions/{sid}/slides/{index}/revise')
def revise(sid: str, index: int, body: generations.Revision):
    sess = session(sid)
    slide_index(sess, index)
    try:
        with sessions.job(sess):
            value = body.model_dump()
            if sess.revisions.get(str(index)) != value:
                sess.revisions[str(index)] = value
                sess.revision_version += 1
                sess.generation = None
                sess.generated = False
                sess.benchmarked = False
                sess.ai_results = {}
            return {'ok': True, 'revision_version': sess.revision_version, 'revisions': sess.revisions,
                    'rerendered': False, 'result': None,
                    'message': 'Revision saved. Generate to apply and verify it.'}
    except ValueError as exc:
        raise HTTPException(409, str(exc))


@router.post('/sessions/{sid}/generate')
def generate(sid: str, body: generations.GenerateRequest = generations.GenerateRequest()):
    sess = session(sid)
    try:
        with sessions.job(sess):
            record = generations.build(sess,mode=body.mode,repair_passes=body.repair_passes)
            return {'ok': record['state'] != 'error', 'generation': generations.public(record),
                    'built_slides': (record.get('report') or {}).get('slide_count', 0)}
    except ValueError as exc:
        raise HTTPException(409, str(exc))


@router.post('/ai/test')
def test_ai_configuration():
    """Explicit UI action: tiny synthetic requests, never a user's presentation."""
    results={}
    for role in ('planner','reviewer'):
        result=providers.generate(role,'Return only JSON: {"ok":true}.',{'purpose':'Slide Studio connection test',
            'schema':{'type':'object','properties':{'ok':{'type':'boolean','enum':[True]}},
                      'required':['ok'],'additionalProperties':False}},max_tokens=256)
        if result['status']=='completed' and result.get('data')!={'ok':True}:
            result={'status':'invalid_response','message':'Provider returned an unexpected connection-test response.'}
        results[role]={k:v for k,v in result.items() if k!='data'}
    return {'configuration':providers.capabilities(),'results':results}


@router.post('/sessions/{sid}/decisions')
def decisions(sid: str, body: generations.Decision):
    sess = session(sid)
    try:
        with sessions.job(sess):
            return {'generation': generations.public(generations.decide(sess, body))}
    except ValueError as exc:
        raise HTTPException(409, str(exc))


class AiRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    enabled: bool = True


def run_ai(sess, index):
    record = sess.generation
    outputs = record['source_to_output_slides'].get(str(index), [])
    results = []
    for oi in outputs:
        png = Path(record['directory'], 'render', f'slide-{oi}.png')
        result = optional_review.review_slide_result(str(png))
        results.append({'output_slide': oi, **result})
    if not results:
        results = [{'status':'invalid_response', 'findings':[], 'message':'No rendered output available.'}]
    record['optional_ai'][str(index)] = results
    # A new provider response cannot inherit an approval of an earlier response.
    old_ids = {f['id'] for f in record['findings'] if f.get('ai_source_slide') == index}
    record['human_decisions'] = [d for d in record['human_decisions']
                                 if not old_ids.intersection(d['finding_ids'])]
    record['findings'] = [f for f in record['findings'] if f.get('ai_source_slide') != index]
    for n, result in enumerate(results):
        if result['status'] != 'completed' or result['findings']:
            record['findings'].append({'id': f'ai:{index}:{n}', 'code':'OPTIONAL_AI_PENDING',
                'severity':'optional_pending' if result['status'] != 'completed' else 'review',
                'ai_source_slide': index, 'output_slide': result.get('output_slide'),
                'message': result.get('message') or str(result['findings']), 'check':'optional_ai'})
    sess.ai_results[index] = results
    generations.settle(record)
    generations.save(record)
    return results


@router.post('/sessions/{sid}/slides/{index}/ai-check')
def ai_check(sid: str, index: int, body: AiRequest):
    sess = session(sid)
    slide_index(sess, index)
    try:
        with sessions.job(sess):
            if body.enabled:
                sess.ai_check.add(index)
                if not sess.generation:
                    return {'ok':True, 'message':'Selected for the next generation.', 'generation':None}
                generations.verify_identity(sess, sess.generation, ready=False)
                run_ai(sess, index)
            else:
                sess.ai_check.discard(index)
                sess.ai_results.pop(index, None)
                if sess.generation:
                    r = sess.generation
                    r['optional_ai'].pop(str(index), None)
                    r['findings'] = [f for f in r['findings'] if f.get('ai_source_slide') != index]
                    r.setdefault('optional_deselections', []).append({'source_slide':index, 'candidate_sha256':r['candidate_sha256']})
                    generations.settle(r)
                    generations.save(r)
            return {'ok':True, 'generation':generations.public(sess.generation)}
    except ValueError as exc:
        raise HTTPException(409, str(exc))


def artifact_bytes(sess, ready, format='pptx'):
    record = sess.generation
    # A legacy draft=true request must never bypass mandatory QA.
    generations.verify_identity(sess, record, ready=True)
    data = Path(record['candidate']).read_bytes()
    if hashlib.sha256(data).hexdigest() != record['candidate_sha256']:
        raise ValueError('ARTIFACT_CHANGED')
    if format == 'pdf':
        export = record.get('pdf_export', {})
        path = Path(export.get('path', ''))
        if not path.is_file() or export.get('candidate_sha256') != record['candidate_sha256'] or generations.sha256(path) != export.get('sha256'):
            raise ValueError('Verified PDF render is unavailable or changed. Regenerate this candidate.')
        return path.read_bytes()
    return data


@router.get('/sessions/{sid}/download')
def download(sid: str, draft: bool=False, generation_id: str | None=None, format: Literal['pptx','pdf']='pptx'):
    sess = session(sid)
    try:
        with sessions.job(sess):
            if not sess.generation or generation_id != sess.generation['generation_id']:
                raise ValueError('GENERATION_MISMATCH')
            data = artifact_bytes(sess, ready=not draft, format=format)
            sess.close_after = min(__import__('time').time()+600, sess.expires)
            filename = ('Stevens-unverified-draft.' if draft else 'Stevens-verified.')+format
            return Response(data, media_type='application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
                headers={'Content-Disposition':f'attachment; filename="{filename}"', 'Cache-Control':'no-store'})
    except ValueError as exc:
        raise HTTPException(409, str(exc))


class BenchmarkRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    generation_id: str
    note: str = ''


@router.post('/sessions/{sid}/benchmark')
def send_to_benchmark(sid: str, body: BenchmarkRequest):
    raise HTTPException(410, 'Permanent benchmark capture is disabled. Files are used only for processing.')


@router.delete('/sessions/{sid}')
def purge(sid: str):
    try:
        return {'deleted': sessions.delete(sid)}
    except ValueError as exc:
        raise HTTPException(409, str(exc))


@router.post('/sessions/{sid}/finalize')
def finalize(sid: str):
    return {'deleted': sessions.delete(sid)}
