"""Native PPTX composition plus independent reopened-artifact verification."""
import hashlib
import json
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from .. import grounded, brand
from .graphics import render_plot, render_equation
from . import diagrams
from slide_engine import template_policy as T, templates


def text(shape, lines, size):
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.04)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Arial'; p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(brand.INK)
        p.space_after = Pt(14)


def compose(spec, path, assets, pages=(), kinds=None):
    assets = Path(assets); assets.mkdir(parents=True, exist_ok=True)
    prs = Presentation(templates.path())
    T.retain_opening_artwork(prs)
    for item in list(prs.slides._sldIdLst):
        prs.part.drop_rel(item.rId); prs.slides._sldIdLst.remove(item)
    layout = next(l for l in templates.layouts(prs) if l.name == T.CONTENT_LAYOUT)
    cover_layout = next(l for l in templates.layouts(prs) if l.name == T.OPENING_LAYOUT)
    closing_layout = next(l for l in templates.layouts(prs) if l.name == T.CLOSING_LAYOUT)
    # The closing layout contains sample text as artwork, not placeholders.
    # Replace that sample copy in this output package, retaining photo/logo art.
    for shape in list(closing_layout.shapes):
        if shape.has_text_frame and shape.text.strip():
            shape._element.getparent().remove(shape._element)
    if kinds is not None and (len(kinds) != len(spec.slides) or kinds[0] != 'opening' or kinds[-1] != 'closing'):
        raise ValueError('Opening and closing must match the approved outline.')
    manifest = []
    for index, content in enumerate(spec.slides):
        closing = kinds is not None and kinds[index] == 'closing'
        bookend = index == 0 or closing
        slide = prs.slides.add_slide(cover_layout if index==0 else closing_layout if closing else layout)
        for sh in list(slide.shapes):
            sh._element.getparent().remove(sh._element)
        title = slide.shapes.add_textbox(*(Inches(v) for v in (T.CLOSING_TITLE if closing else T.COVER_TITLE if bookend else (*T.CONTENT[:3],1.4))))
        title.name = 'authored-title'
        text(title, [content.title], 40)
        content_height=4.4 if templates.current_id()=='stevens' else T.CONTENT[1]+T.CONTENT[3]-2
        visual = any(v is not None for v in (content.chart, content.plot, content.equation, content.figure_page, content.table, content.diagram))
        body = slide.shapes.add_textbox(*(Inches(v) for v in (T.CLOSING_DETAILS if closing else T.COVER_DETAILS if bookend else (.75,2,4.0 if visual else min(11.5,T.CONTENT[0]+T.CONTENT[2]-.75),content_height))))
        body.name = 'authored-body'
        text(body, content.bullets, 20 if sum(map(len, content.bullets)) < 600 else 18)
        if bookend:
            from .text_fit import fit_bookend_body
            fit_bookend_body(body)
            for shape in (title,body):
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.color.rgb=RGBColor.from_string(brand.WHITE)
        vx,vy,vw,vh=T.COVER_SUPPORT if index==0 else (5,2,min(7.2,T.CONTENT[0]+T.CONTENT[2]-5),min(4.3,content_height))
        image_hash = None
        diagram_shapes = []
        if content.diagram:
            diagram_shapes = diagrams.compose(slide, content.diagram, (vx,vy,vw,vh))
        elif content.chart:
            chart = content.chart
            if chart.kind == 'scatter':
                data = XyChartData()
                for series in chart.series:
                    target = data.add_series(series.name)
                    for x, y in zip(chart.categories, series.values): target.add_data_point(float(x), y)
            else:
                data = CategoryChartData(); data.categories = chart.categories
                for series in chart.series: data.add_series(series.name, series.values)
            chart_kinds = {'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'column': XL_CHART_TYPE.COLUMN_CLUSTERED,
                     'line': XL_CHART_TYPE.LINE, 'pie': XL_CHART_TYPE.PIE, 'area': XL_CHART_TYPE.AREA,
                     'scatter': XL_CHART_TYPE.XY_SCATTER}
            native = slide.shapes.add_chart(chart_kinds[chart.kind], *(Inches(v) for v in (vx,vy,vw,vh)), data).chart
            native.font.name = 'Arial'; native.font.size = Pt(14)
            native.has_title = False
            native.has_legend = len(chart.series) > 1 or chart.kind == 'pie'
            if native.has_legend:
                native.legend.position = XL_LEGEND_POSITION.TOP
                native.legend.include_in_layout = False
            if chart.kind != 'pie':
                for axis, label in ((native.category_axis, chart.x_label), (native.value_axis, chart.y_label)):
                    if label:
                        axis.has_title = True; axis.axis_title.text_frame.text = label
        elif content.table:
            data = [content.table.headers]+content.table.rows
            table = slide.shapes.add_table(len(data),len(data[0]),*(Inches(v) for v in (vx,vy,vw,vh))).table
            for ri, row in enumerate(data):
                for ci, value in enumerate(row):
                    cell = table.cell(ri,ci); cell.text = value
                    # Do not combine explicit dark text with the template's red header fill.
                    cell.fill.solid()
                    cell.fill.fore_color.rgb=RGBColor.from_string(brand.RED if ri==0 else brand.WHITE)
                    for p in cell.text_frame.paragraphs:
                        p.font.name='Arial'; p.font.size=Pt(14)
                        p.font.color.rgb=RGBColor.from_string(brand.WHITE if ri==0 else brand.INK)
        elif visual:
            image = assets/f'visual-{index}.png'
            if content.plot: render_plot(content.plot, image)
            elif content.equation: render_equation(content.equation, image)
            else:
                page = next((p for p in pages if p['page'] == content.figure_page), None)
                if not page: raise ValueError('Figure source page does not exist.')
                image = Path(page['image'])
            with Image.open(image) as im: w, h = im.size
            scale = min(vw/w, vh/h)
            picture = slide.shapes.add_picture(str(image), Inches(vx+(vw-w*scale)/2), Inches(vy+(vh-h*scale)/2),
                                             width=Inches(w*scale), height=Inches(h*scale))
            picture.name = 'authored-visual'
            picture._element.nvPicPr.cNvPr.set('descr', content.equation or ('Plot: '+', '.join(content.plot.functions) if content.plot else f'Source PDF page {content.figure_page}'))
            image_hash = hashlib.sha256(image.read_bytes()).hexdigest()
        references = '\n'.join(f'PDF page {c.page}: {c.quote}' for c in content.citations)
        notes = content.notes+'\n'+references+'\nAuthoring specification:\n'+content.model_dump_json()
        slide.notes_slide.notes_text_frame.text = notes
        manifest.append({'id': content.id, 'title': content.title, 'bullets': content.bullets,
                         'kind':kinds[index] if kinds else ('opening' if index==0 else 'content'),
                         'layout':slide.slide_layout.name,
                         'notes': notes, 'chart': content.chart.model_dump() if content.chart else None,
                         'image_sha256': image_hash, 'table':content.table.model_dump() if content.table else None,
                         'diagram':content.diagram.model_dump() if content.diagram else None, 'diagram_shapes':diagram_shapes})
    prs.save(path)
    return manifest


def audit(path, manifest):
    prs = Presentation(path)
    findings = []
    def fail(i, message):
        findings.append({'code': 'AUTHORED_CONTENT_MISMATCH', 'severity': 'blocking', 'output_slide': i, 'message': message})
    if len(prs.slides) != len(manifest): fail(0, 'Slide count differs from approved authored content.')
    for i, (slide, wanted) in enumerate(zip(prs.slides, manifest)):
        if wanted.get('layout') and slide.slide_layout.name != wanted['layout']: fail(i, 'Layout differs from the approved slide role.')
        named = {s.name: s for s in slide.shapes}
        if named.get('authored-title') is None or named['authored-title'].text != wanted['title']: fail(i, 'Title differs.')
        if named.get('authored-body') is None or named['authored-body'].text != '\n'.join(wanted['bullets']): fail(i, 'Body text differs.')
        if slide.notes_slide.notes_text_frame.text != wanted['notes']: fail(i, 'Notes or source specification differs.')
        charts = [s.chart for s in slide.shapes if s.has_chart]
        if wanted['chart']:
            if len(charts) != 1: fail(i, 'Expected one native chart.')
            else:
                actual = [{'name': s.name, 'values': list(s.values)} for s in charts[0].series]
                if actual != wanted['chart']['series']: fail(i, 'Chart series values differ.')
                if wanted['chart']['kind'] == 'scatter':
                    from lxml import etree
                    ns = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
                    for series in charts[0]._chartSpace.findall('.//c:ser', ns):
                        xs = [float(node.text) for node in series.findall('.//c:xVal/c:numRef/c:numCache/c:pt/c:v',ns)]
                        if xs != [float(x) for x in wanted['chart']['categories']]: fail(i, 'Scatter x values differ.')
                if wanted['chart']['kind'] != 'scatter':
                    if [c.label for c in charts[0].plots[0].categories] != wanted['chart']['categories']: fail(i, 'Chart categories differ.')
        elif charts: fail(i, 'Unexpected chart.')
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        if wanted['image_sha256']:
            if len(pictures) != 1 or hashlib.sha256(pictures[0].image.blob).hexdigest() != wanted['image_sha256']: fail(i, 'Figure bytes differ.')
        elif pictures: fail(i, 'Unexpected picture.')
        tables = [s.table for s in slide.shapes if s.has_table]
        if wanted.get('table'):
            expected = [wanted['table']['headers']]+wanted['table']['rows']
            if len(tables)!=1 or [[c.text for c in row.cells] for row in tables[0].rows] != expected: fail(i, 'Table values differ.')
        elif tables: fail(i, 'Unexpected table.')
        if not diagrams.matches(slide, wanted.get('diagram_shapes', [])): fail(i, 'Diagram labels or relationships differ.')
    return {'status': 'failed' if findings else 'passed', 'findings': findings}
