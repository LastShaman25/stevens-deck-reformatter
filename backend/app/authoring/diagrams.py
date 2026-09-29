"""Small, editable evidence diagrams inside the template content region."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from .. import brand


def compose(slide, spec, box):
    x, y, w, h = box
    count = len(spec.nodes)
    step = h / count
    compact = count > 4
    shapes = []
    for index, node in enumerate(spec.nodes):
        top = y + index * step
        label = slide.shapes.add_textbox(Inches(x + .5), Inches(top), Inches(w - .6), Inches(.33 if compact else .38))
        label.name = f'authored-diagram-label-{index}'
        label.text = node.label
        p = label.text_frame.paragraphs[0]
        p.font.name = 'Arial'; p.font.size = Pt(18 if compact else 20); p.font.bold = True
        p.font.color.rgb = RGBColor.from_string(brand.RED)
        label.text_frame.margin_top = label.text_frame.margin_bottom = 0
        shapes.append(label)
        offset = .34 if compact else .39
        detail = slide.shapes.add_textbox(Inches(x + .5), Inches(top + offset), Inches(w - .6), Inches(step - offset - .01))
        detail.name = f'authored-diagram-detail-{index}'
        detail.text = node.detail
        detail.text_frame.word_wrap = True
        detail.text_frame.margin_top = detail.text_frame.margin_bottom = 0
        p = detail.text_frame.paragraphs[0]
        p.font.name = 'Arial'; p.font.size = Pt(17)
        p.font.color.rgb = RGBColor.from_string(brand.INK)
        shapes.append(detail)
        if spec.kind == 'process' and index < count - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x + .1), Inches(top + .28), Inches(.2), Inches(step - .12))
            arrow.name = f'authored-diagram-arrow-{index}'
            arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor.from_string(brand.RED)
            arrow.line.fill.background()
            shapes.append(arrow)
    # Reopened-artifact verification checks labels and relation shapes, not just notes.
    return [{'name':s.name, 'text':s.text if s.has_text_frame else '',
             'shape_type':int(s.shape_type), 'xml':s._element.xml} for s in shapes]


def matches(slide, expected):
    actual = [s for s in slide.shapes if s.name.startswith('authored-diagram-')]
    return [{'name':s.name, 'text':s.text if s.has_text_frame else '',
             'shape_type':int(s.shape_type), 'xml':s._element.xml} for s in actual] == expected
