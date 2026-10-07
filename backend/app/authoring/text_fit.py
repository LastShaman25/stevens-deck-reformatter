"""Measured fitting for the template's compact bookend text regions."""
from functools import lru_cache
from PIL import ImageFont
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE


@lru_cache(maxsize=16)
def font(size):
    from matplotlib.font_manager import findfont, FontProperties
    return ImageFont.truetype(findfont(FontProperties(family='Arial')), size * 4)


def wrapped_lines(text, face, width):
    # Measure at 4x resolution. Leave a small allowance for renderer differences.
    limit = width * 4 * .96
    count = 0
    for line in text.replace('\v', '\n').split('\n'):
        current = ''
        for word in line.split():
            proposed = current + (' ' if current else '') + word
            if face.getlength(proposed) <= limit:
                current = proposed
                continue
            if current:
                count += 1
            current = ''
            for char in word:
                if current and face.getlength(current + char) > limit:
                    count += 1
                    current = ''
                current += char
        count += 1
    return count


def fit_bookend_body(shape):
    tf = shape.text_frame
    width = (shape.width - tf.margin_left - tf.margin_right) / 12700
    height = (shape.height - tf.margin_top - tf.margin_bottom) / 12700
    paragraphs = list(tf.paragraphs)
    for size in range(20, 15, -1):
        lines = sum(wrapped_lines(p.text, font(size), width) for p in paragraphs)
        needed = lines * size * 1.2 + max(0, len(paragraphs)-1) * 8
        if needed > height * .96:
            continue
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for i, p in enumerate(paragraphs):
            p.font.size = Pt(size)
            p.line_spacing = Pt(size * 1.2)
            p.space_before = Pt(0)
            p.space_after = Pt(8 if i+1 < len(paragraphs) else 0)
        return
    raise ValueError('The details do not fit the selected template at a readable size. '
                     'Condense the visible introduction or takeaways and retain the full explanation in speaker notes. '
                     'Keep the approved title and meaning; do not add slides.')


def fit_content_body(shape, preferred_size=20):
    """Fit the actual body column; character count alone misses narrow columns."""
    tf = shape.text_frame
    width = (shape.width-tf.margin_left-tf.margin_right)/12700
    height = (shape.height-tf.margin_top-tf.margin_bottom)/12700
    paragraphs = list(tf.paragraphs)
    for size in range(preferred_size, 15, -1):
        lines = sum(wrapped_lines(p.text, font(size), width) for p in paragraphs)
        # Explicit line/paragraph spacing prevents inherited template metrics
        # from spilling otherwise bounded text across the protected footer.
        for gap in (10, 8, 6):
            if lines*size*1.2 + max(0,len(paragraphs)-1)*gap > height*.94:
                continue
            tf.auto_size = MSO_AUTO_SIZE.NONE
            for i,p in enumerate(paragraphs):
                p.font.size = Pt(size)
                p.line_spacing = Pt(size*1.2)
                p.space_before = Pt(0)
                p.space_after = Pt(gap if i+1<len(paragraphs) else 0)
            return
    raise ValueError('Body text cannot fit readably in its template region. Shorten the points or split the slide.')


def fit_cover_title(shape):
    """Fit complete native title copy inside its template region, including wraps."""
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    width, height = shape.width / 12700, shape.height / 12700
    for size in range(40, 19, -1):
        count = sum(wrapped_lines(p.text, font(size), width) for p in tf.paragraphs)
        if count * size * 1.22 <= height * .94:
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
            for p in tf.paragraphs:
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = Pt(size * 1.22)
                for r in p.runs: r.font.size = Pt(size)
            return
    raise ValueError('The full cover title cannot fit readably in the selected template.')
