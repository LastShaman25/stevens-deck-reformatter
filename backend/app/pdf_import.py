"""Processing-only PDF import: editable text and faithful raster graphic regions.

The original PDF remains the paired visual authority. Pages without a text layer
are preserved as page images, explicitly identified as not individually editable.
"""
from io import BytesIO
from pathlib import Path
import hashlib
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from . import pdf_regions
from .pdf_paths import remove_paths, path_bounds


def compact(text):
    return ''.join(text.split())


def needs_glyph_image(line):
    # Some embedded TeX fonts map visible glyphs to XML-illegal controls.
    # PowerPoint would show python-pptx's literal _x0000_ escapes instead.
    return any(ord(c)<32 and c not in '\t\n\r' for span in line['spans'] for c in span['text'])


def regions(boxes):
    merged=[]
    for rect in boxes:
        rect=fitz.Rect(rect)
        if rect.is_empty: continue
        again=True
        while again:
            again=False
            for other in list(merged):
                if (rect+(-1,-1,1,1)).intersects(other):
                    rect |= other; merged.remove(other); again=True
        merged.append(rect)
    return merged


def link_underlines(lines, drawings, links):
    """Recognize standalone, span-wide PDF link rules as text decoration.

    Keeping these as independent pictures lets font metrics move the text while
    the rule stays behind, turning an underline into an apparent strike-through.
    Ambiguous rules, partial spans, chart lines and ordinary artwork stay native
    source graphics.
    """
    spans = [s for line in lines if not needs_glyph_image(line) for s in line['spans'] if s['text'].strip()]
    marked, converted = set(), []
    for drawing in drawings:
        if len(drawing['items']) != 1 or drawing['items'][0][0] != 're' or drawing.get('type') != 'f':
            continue
        rect = fitz.Rect(drawing['rect'])
        if not 0 < rect.height <= 2 or rect.width < 3 or not drawing.get('fill') or drawing.get('fill_opacity',1) != 1:
            continue
        rgb = tuple(round(c*255) for c in drawing['fill'])
        for span in spans:
            box = fitz.Rect(span['bbox']); size = span['size']; baseline = span['origin'][1]
            color = tuple((span['color'] >> shift) & 255 for shift in (16,8,0))
            if (rgb == color and abs(rect.x0-box.x0)<.75 and abs(rect.x1-box.x1)<.75
                    and baseline < rect.y0 <= baseline+size*.3 and rect.y1 <= box.y1+.75
                    and any(link['kind']==fitz.LINK_URI and (fitz.Rect(link['from'])+(-1,-1,1,1)).contains(rect) for link in links)):
                marked.add(id(span)); converted.append(drawing); break
    return marked, converted


def add_links(slide, links, box, page_number, page_count):
    for link in links:
        # Beamer named GoTo destinations expose a resolved page as LINK_NAMED.
        if link['kind'] not in (fitz.LINK_URI,fitz.LINK_GOTO,fitz.LINK_NAMED):
            raise ValueError(f'PDF page {page_number} contains an unsupported link type.')
        if link['kind'] != fitz.LINK_URI:
            target = link.get('page')
            if not isinstance(target,int) or not 0 <= target < page_count:
                raise ValueError(f'PDF page {page_number} has an internal link whose target page could not be resolved.')
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,*box(link['from']))
        shape.fill.background(); shape.line.fill.background()
        if link['kind'] == fitz.LINK_URI: shape.click_action.hyperlink.address = link['uri']
        else: shape.name = f"pdf-internal-link:{link['page']}"


def convert(pdf_path, pptx_path, preview_path):
    with fitz.open(pdf_path) as doc:
        if doc.needs_pass: raise ValueError('Upload an unlocked PDF.')
        if not 1<=len(doc)<=100: raise ValueError('PDF redesign supports 1–100 pages.')
        if doc.embfile_count(): raise ValueError('PDF attachments cannot be preserved in this workflow. Remove attachments before uploading.')
        from slide_engine.template_policy import CANVAS
        prs=Presentation();prs.slide_width,prs.slide_height=CANVAS
        canvas_w,canvas_h=(v/914400 for v in CANVAS)
        evidence=[];characters=0
        repeated = pdf_regions.repeated_footer(doc)
        for i,page in enumerate(doc):
            if page.first_widget or page.first_annot:
                raise ValueError(f'PDF page {i+1} has form fields or annotations; flatten them in your PDF editor first.')
            blocks=page.get_text('dict')['blocks']
            lines=[line for block in blocks if block['type']==0 for line in block['lines']]
            source_text=''.join(span['text'] for line in lines for span in line['spans'])
            characters+=len(source_text)
            if len(lines)>500 or characters>220000:
                raise ValueError('PDF exceeds the text processing limit; upload a shorter section.')
            if any(line.get('dir',(1,0))!=(1,0) for line in lines):
                raise ValueError(f'PDF page {i+1} has rotated text; normalize its orientation before redesign.')
            size=page.rect
            scale=min(canvas_w/size.width,canvas_h/size.height)
            dx=(canvas_w-size.width*scale)/2;dy=(canvas_h-size.height*scale)/2
            def box(rect):
                r=fitz.Rect(rect)
                return [Inches(dx+r.x0*scale),Inches(dy+r.y0*scale),Inches(max(.001,r.width*scale)),Inches(max(.001,r.height*scale))]
            slide=prs.slides.add_slide(prs.slide_layouts[6])
            slide._element.cSld.set('name','sss:pdf-import')
            page.get_pixmap(matrix=fitz.Matrix(min(2,1600/max(size.width,size.height)),min(2,1600/max(size.width,size.height))),alpha=False).save(preview_path(i))
            if not source_text.strip():
                # Visible equations, diagrams and scans need not have PDF text.
                # Preserve their composition before any text/chrome heuristics.
                # Bound the raster size for unusually large PDF page dimensions.
                resolution = min(3,3000/max(size.width,size.height))
                pix = page.get_pixmap(matrix=fitz.Matrix(resolution,resolution),alpha=False)
                blob = pix.tobytes('png')
                picture = slide.shapes.add_picture(BytesIO(blob),*box(size))
                picture.name = f'PDF page {i+1} preserved page image (no text layer)'
                links = page.get_links()
                add_links(slide,links,box,i+1,len(doc))
                evidence.append({'page':i+1,'text':source_text,'editable_text':'',
                    'graphics':1,'links':len(links),'native_link_underlines':0,
                    'rasterized_text_lines':0,'source_font_regions':0,'source_font_evidence':[],
                    'page_image_regions':1,
                    'page_image_evidence':{'bounds':list(size),'pixel_size':[pix.width,pix.height],
                                           'image_sha256':hashlib.sha256(blob).hexdigest()},
                    'visible_content_bounds':list(size),
                    'excluded_chrome':{'repeated_footer_text':[],'navigation_links':0,
                        'decorative_paths':0,'navigation_path_bounds':[],
                        'removed_native_paths':0,'title_shadow_images':0},
                    'preservation':'No extractable PDF text layer. Entire page preserved as an image; labels and equations are not individually editable. No OCR was performed.'})
                continue
            drawings=page.get_drawings()
            links=page.get_links()
            footer_lines, footer_drawings, nav_links = pdf_regions.chrome(page, blocks, drawings, repeated)
            page_numbers = [line for line in lines if compact(pdf_regions.text(line)) == str(i+1)
                            and fitz.Rect(line['bbox']).y0 > size.height * .94
                            and fitz.Rect(line['bbox']).width < size.width * .04]
            footer_lines += page_numbers
            lines = [line for line in lines if line not in footer_lines]
            links = [link for link in links if link not in nav_links]
            title_lines = pdf_regions.cover_title(blocks, footer_lines) if i == 0 else []
            backdrop, shadows = pdf_regions.title_backdrop(page, title_lines, lines, drawings,
                [b for b in blocks if b['type'] == 1]) if i == 0 else ([], [])
            excluded_drawings = footer_drawings + backdrop
            content_rect = fitz.Rect(size)
            bands = [fitz.Rect(d['rect']) for d in footer_drawings if d.get('fill') is not None
                     and d.get('fill_opacity',1) == 1
                     and fitz.Rect(d['rect']).width >= size.width*.3
                     and fitz.Rect(d['rect']).y0 >= size.height*.94
                     and fitz.Rect(d['rect']).y1 >= size.height-.5]
            covered_width = 0
            for band in sorted(bands,key=lambda r:r.x0):
                if band.x0 > covered_width+.5: break
                covered_width = max(covered_width,band.x1)
            if bands and covered_width >= size.width-.5:
                # The opaque full-width old footer hid this area in the source.
                # Do not reveal hidden article text or white running copy when
                # replacing it. Keep the original visible content boundary.
                content_rect.y1 = max(r.y0 for r in bands)
            underlined_spans, underline_drawings = link_underlines(lines, drawings, links)
            full_fills=[d for d in drawings if d.get('fill') and d.get('fill_opacity',1)==1
                        and len(d['items'])==1 and d['items'][0][0]=='re'
                        and (fitz.Rect(d['rect']) & size).get_area()>=size.get_area()*.99]
            if full_fills:
                rgb=full_fills[-1]['fill'];slide.background.fill.solid()
                slide.background.fill.fore_color.rgb=RGBColor(*(round(c*255) for c in rgb))
            # Fully transparent PDF path bounds carry no visible artwork. Using
            # them as region seeds can join unrelated screenshots and underlines
            # into one giant bitmap, making independent layout repair impossible.
            visible_drawings=[d for d in drawings if
                (d.get('fill') is not None and (d.get('fill_opacity') or 0)>0) or
                (d.get('color') is not None and (d.get('stroke_opacity') or 0)>0)]
            graphic_boxes=[fitz.Rect(d['rect'])+(-.5,-.5,.5,.5) for d in visible_drawings if d not in full_fills and d not in underline_drawings and d not in excluded_drawings]
            graphic_boxes += [b['bbox'] for b in blocks if b['type']==1 and b not in shadows]
            font_seeds = [line for line in lines if i > 0 and pdf_regions.needs_source_font(line)]
            # Keep the original narrow non-Unicode fallback when no font region
            # is needed. Complex embedded fonts use whole-object closure.
            faithful = pdf_regions.fidelity_regions(lines, graphic_boxes, font_seeds) if font_seeds else []
            faithful_lines = [line for line in lines if any(r.contains(fitz.Rect(line['bbox'])) for r in faithful)]
            graphic_boxes = [b for b in graphic_boxes if not any(r.contains(fitz.Rect(b)) for r in faithful)]
            faithful = [r & content_rect for r in faithful if not (r & content_rect).is_empty]
            # Work on a private page copy: strip only native text. Images and
            # vector graphics retain source pixels and layer order within each region.
            layer=fitz.open();layer.insert_pdf(doc,from_page=i,to_page=i)
            graphic_page=layer[0]
            remove_paths(graphic_page, excluded_drawings)
            for line in pdf_regions.lines(blocks): graphic_page.add_redact_annot(line['bbox'],fill=False)
            graphic_page.apply_redactions(images=0,graphics=0,text=0)
            for drawing in underline_drawings:
                graphic_page.add_redact_annot(fitz.Rect(drawing['rect'])+(-.01,-.01,.01,.01),fill=False)
            if underline_drawings:
                graphic_page.apply_redactions(images=0,graphics=1,text=1)
            graphic_count=0
            for rect in regions(graphic_boxes):
                rect=(rect+(-.5,-.5,.5,.5)) & content_rect
                if rect.is_empty: continue
                resolution=min(2,3000/max(size.width,size.height))
                pix=graphic_page.get_pixmap(matrix=fitz.Matrix(resolution,resolution),clip=rect,alpha=True)
                # White-on-white background drawings have no visible information.
                # Do not turn them into a large opaque picture on the new cover.
                from PIL import Image
                pixels=Image.open(BytesIO(pix.tobytes('png'))).convert('RGBA')
                bg=tuple(round(c*255) for c in full_fills[-1]['fill']) if full_fills else (255,255,255)
                flattened=Image.new('RGBA',pixels.size,(*bg,255));flattened.alpha_composite(pixels)
                if all(lo==hi==c for (lo,hi),c in zip(flattened.convert('RGB').getextrema(),bg)): continue
                # PDF graphic extents often include an almost page-sized white
                # rectangle around a tiny logo. Trim only uniform outside padding;
                # retain every visible pixel and update its exact page coordinates.
                from PIL import ImageChops
                visible=ImageChops.difference(flattened.convert('RGB'),Image.new('RGB',pixels.size,bg)).getbbox()
                if not visible: continue
                crop=BytesIO();pixels.crop(visible).save(crop,format='PNG')
                rect=fitz.Rect((pix.x+visible[0])/resolution,(pix.y+visible[1])/resolution,
                               (pix.x+visible[2])/resolution,(pix.y+visible[3])/resolution)
                picture=slide.shapes.add_picture(BytesIO(crop.getvalue()),*box(rect))
                picture.name=f'PDF page {i+1} graphic {graphic_count+1}'
                graphic_count+=1
            layer.close()
            font_evidence = []
            content_layer = fitz.open(); content_layer.insert_pdf(doc,from_page=i,to_page=i)
            content_page = content_layer[0]
            # A source text region can intersect obsolete navigation paths.
            # Remove only confirmed furniture on a private copy, preserving all
            # text and images; otherwise raster fallback reintroduces that chrome.
            removed_paths = remove_paths(content_page, footer_drawings)
            for rect in faithful:
                rect = rect & size
                pix = content_page.get_pixmap(matrix=fitz.Matrix(3,3),clip=rect,alpha=False)
                actual = fitz.Rect(pix.x/3,pix.y/3,(pix.x+pix.width)/3,(pix.y+pix.height)/3)
                blob = pix.tobytes('png')
                picture = slide.shapes.add_picture(BytesIO(blob),*box(actual))
                picture.name = f'PDF page {i+1} source-font region {graphic_count+1}'
                font_evidence.append({'bounds':list(actual),'image_sha256':hashlib.sha256(blob).hexdigest()})
                graphic_count += 1
            content_layer.close()
            glyph_lines=[line for line in lines if line not in faithful_lines and needs_glyph_image(line)]
            for line in glyph_lines:
                rect=fitz.Rect(line['bbox']) & size
                pix=page.get_pixmap(matrix=fitz.Matrix(3,3),clip=rect,alpha=True)
                # Use the rounded raster bounds so the source glyphs stay aligned.
                rect=fitz.Rect(pix.x/3,pix.y/3,(pix.x+pix.width)/3,(pix.y+pix.height)/3)
                picture=slide.shapes.add_picture(BytesIO(pix.tobytes('png')),*box(rect))
                picture.name=f'PDF page {i+1} preserved text glyphs {graphic_count+1}'
                graphic_count+=1
            editable_lines=[line for line in lines if line not in faithful_lines and not needs_glyph_image(line)]
            textboxes=[]
            title_shape = None
            for line in editable_lines:
                if line in title_lines and line != title_lines[0] and title_shape is not None:
                    shape = title_shape; tf = shape.text_frame; paragraph = tf.add_paragraph()
                    bounds = fitz.Rect(title_lines[0]['bbox'])
                    for title_line in title_lines: bounds |= fitz.Rect(title_line['bbox'])
                    shape.left,shape.top,shape.width,shape.height = box(bounds)
                else:
                    shape=slide.shapes.add_textbox(*box(line['bbox']));textboxes.append(shape)
                    tf=shape.text_frame;tf.clear();tf.word_wrap=False
                    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                    paragraph=tf.paragraphs[0]
                    if line == (title_lines[0] if title_lines else None): title_shape = shape
                paragraph.space_before=paragraph.space_after=Pt(0)
                # Preserve PDF run baselines, not merely Unicode text. Sub/super-
                # scripts often have ordinary digits at a displaced PDF origin.
                main=max(line['spans'],key=lambda s:s['size'])
                baseline=main['origin'][1]
                for span in line['spans']:
                    run=paragraph.add_run();run.text=span['text']
                    run.font.name=span['font'].split('+')[-1]
                    run.font.size=Pt(max(1,span['size']*scale*72))
                    run.font.bold=bool(span['flags'] & 16);run.font.italic=bool(span['flags'] & 2)
                    run.font.color.rgb=RGBColor.from_string(f"{span['color']:06X}")
                    if id(span) in underlined_spans: run.font.underline=True
                    shift=baseline-span['origin'][1]
                    if abs(shift)>.25:
                        # PowerPoint automatically renders shifted runs at 2/3
                        # of sz. The PDF already contains the reduced glyph size;
                        # compensate or every exponent/subscript is shrunk twice.
                        nominal=span['size']*1.5
                        run.font.size=Pt(max(1,nominal*scale*72))
                        run._r.get_or_add_rPr().set('baseline',str(round(shift/max(1,nominal)*100000)))
            add_links(slide,links,box,i+1,len(doc))
            evidence.append({'page':i+1,'text':source_text,'graphics':graphic_count,'links':len(links),
                             'editable_text':''.join(s['text'] for line in editable_lines for s in line['spans']),
                             'native_link_underlines':len(underlined_spans),
                             'rasterized_text_lines':len(glyph_lines)+len(faithful_lines),
                             'source_font_regions':len(faithful),
                             'source_font_evidence':font_evidence,
                             'visible_content_bounds':list(content_rect),
                             'excluded_chrome':{'repeated_footer_text':[pdf_regions.text(l) for l in footer_lines],
                                 'navigation_links':len(nav_links),'decorative_paths':len(excluded_drawings),
                                 'navigation_path_bounds':[list(path_bounds(d)) for d in footer_drawings],
                                 'removed_native_paths':removed_paths,
                                 'title_shadow_images':len(shadows)},
                             'preservation':'Complex source fonts are preserved as image regions; cover and supported text remain editable.'})
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.name.startswith('pdf-internal-link:'):
                    target=int(shape.name.split(':')[1])
                    if not 0<=target<len(prs.slides): raise ValueError('PDF internal link target is missing.')
                    shape.click_action.target_slide=prs.slides[target]
        prs.save(pptx_path)
        reopened=Presentation(pptx_path)
        for page,slide in zip(evidence,reopened.slides):
            text=''.join(s.text for s in slide.shapes if s.has_text_frame)
            if compact(text)!=compact(page['editable_text']): raise ValueError('PDF import text verification failed.')
        return {'status':'passed','pages':len(evidence),'source_pdf_sha256':hashlib.sha256(Path(pdf_path).read_bytes()).hexdigest(),
                'source_pptx_sha256':hashlib.sha256(Path(pptx_path).read_bytes()).hexdigest(),
                'method':'editable supported text; complex fonts and graphics preserved as image regions; pages without extractable text preserved as page images; repeated navigation/footer chrome removed where separable',
                'page_evidence':[{'page':p['page'],'text_sha256':hashlib.sha256(p['text'].encode()).hexdigest(),
                                  'graphics':p['graphics'],'links':p['links'],
                                  'native_link_underlines':p['native_link_underlines'],
                                  'source_font_regions':p['source_font_regions'],
                                  'source_font_evidence':p['source_font_evidence'],
                                  'page_image_regions':p.get('page_image_regions',0),
                                  'page_image_evidence':p.get('page_image_evidence'),
                                  'visible_content_bounds':p['visible_content_bounds'],
                                  'excluded_chrome':p['excluded_chrome'],
                                  'preservation':p['preservation'],
                                  'rasterized_text_lines':p['rasterized_text_lines']} for p in evidence]}
