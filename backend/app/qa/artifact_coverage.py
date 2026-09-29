"""Reopen and inspect actual output objects; provenance only locates them."""
from collections import Counter
from pptx import Presentation
from slide_engine import inventory
from slide_engine.preserve import marker


def audit(source_path, candidate_path, report):
    source = inventory.inspect(source_path)
    prs = Presentation(candidate_path)
    mapping = report['source_to_output_slides']
    reverse = {o: int(s) for s, outputs in mapping.items() for o in outputs}
    numbers = {str(s.part.partname): reverse.get(i, -1) for i,s in enumerate(prs.slides)}
    found, findings = {}, []
    covered_outputs = [o for outputs in mapping.values() for o in outputs]
    from slide_engine import bookends
    findings.extend(bookends.check(prs,report))
    try: added=bookends.validate_added(prs,report)
    except ValueError: added=set()
    if sorted(covered_outputs+list(added)) != list(range(len(prs.slides))) or set(mapping) != {str(i) for i in range(source.slide_count)}:
        findings.append({'code': 'INVALID_SLIDE_MAPPING'})
    source_prs = Presentation(source_path)
    decisions=report.get('source_decisions',{})
    removals=set()
    extractions={}
    if decisions:
        from ..ai.source_decisions import validate
        for key,value in decisions.items():
            decision=validate(value,source_prs,int(key),source.source_sha256)
            removals.update(decision.remove_ids)
            for extraction in decision.logo_extractions:
                item=next(x for x in source.items if x.id==extraction.source_id)
                original=next(s for origin,_,s in inventory.source_objects(source_prs.slides[int(key)])
                              if origin==item.origin and s.shape_id==item.shape_id)
                from slide_engine import logo_assets
                # Recompute from reopened source pixels, never trust placement metadata.
                extractions[item.id]=logo_assets.extract(original,extraction.region)[0]
        for item in source.items:
            if item.id in removals: item.disposition='excluded_by_policy'
    source_numbers = {str(s.part.partname):i for i,s in enumerate(source_prs.slides)}
    repeats = set()
    for oi, slide in enumerate(prs.slides):
        if oi in added: continue
        from slide_engine import template_policy as T
        si=reverse.get(oi)
        kept=decisions.get(str(si),{}).get('action')=='keep_original'
        if kept:
            original=source_prs.slides[si]
            if (mapping.get(str(si))!=[oi] or not T.is_preserved(slide)
                    or not T.canvas_matches(source_prs.slide_width,source_prs.slide_height,(prs.slide_width,prs.slide_height))
                    or inventory.relationship_signature(original.part,normalize_blob=T.unchanged_blob)!=inventory.relationship_signature(slide.part,normalize_blob=T.unchanged_blob)):
                findings.append({'code':'UNCHANGED_SLIDE_ALTERED','output_slide':oi,'severity':'blocking'})
            for origin,part,shape in inventory.source_objects(slide):
                sid=f'{source.source_sha256[:16]}/{si}/{origin}/{shape.shape_id}'
                item=next((x for x in source.items if x.id==sid),None)
                if item and item.disposition!='excluded_by_policy':
                    found.setdefault(sid,[]).append((oi,inventory.signature(shape,part,numbers)))
            continue
        if T.is_preserved(slide):
            findings.append({'code':'UNAUTHORIZED_UNCHANGED_SLIDE','output_slide':oi})
        for extraction in decisions.get(str(si),{}).get('logo_extractions',[]):
            if extraction.get('placement')!='template_logo': continue
            from slide_engine import logo_assets
            sid=extraction['source_id']
            try:
                _,reuse=logo_assets.template_logo(extractions[sid],si==0,slide)
                expected=[p for p in report['placements'] if p['source_id']==sid and p['output_slide']==oi]
                if len(expected)!=1 or expected[0].get('bounds')!=reuse['bounds']:
                    raise ValueError('Template reuse placement differs from approved artwork.')
                found.setdefault(sid,[]).append((oi,{'extracted_logo':True}))
            except ValueError:
                findings.append({'code':'REUSED_TEMPLATE_LOGO_ALTERED','id':sid,'output_slide':oi,
                                 'criterion':'brand_consistency','severity':'blocking'})
        panels=[s for s in slide.shapes if s.name==T.BACKGROUND_NAME]
        color=T.solid_background(source_prs.slides[reverse[oi]]) if oi in reverse else None
        if oi in reverse:
            removed={(x.origin,x.shape_id) for x in source.items if x.slide==reverse[oi] and x.disposition=='excluded_by_policy'}
            removed.update((x.origin,x.shape_id) for x in source.items if x.slide==reverse[oi] and x.id in extractions)
            if not T.needs_source_background(source_prs.slides[reverse[oi]],slide,source_prs.slide_height,removed,decisions.get(str(si))): color=None
        if color and len(panels)!=1:
            findings.append({'code':'SOURCE_BACKGROUND_MISSING','output_slide':oi})
        connection_ids = {x.get('id'): x.get('name','').split('|',1)[0].rsplit('/',1)[-1] for x in slide._element.xpath('.//p:cNvPr') if x.get('name','').startswith('sss')}
        for sh in slide.shapes:
            name = marker(sh)
            if name==T.BACKGROUND_NAME:
                # Verify against the independently reopened source, never a report claim.
                from pptx.util import Inches
                fills=sh._element.xpath('./p:spPr/a:solidFill/a:srgbClr')
                valid=(color and len(panels)==1 and list(slide.shapes)[0]._element is sh._element
                       and not sh.text.strip() and len(fills)==1 and fills[0].get('val')==color
                       and not list(fills[0]) and sh._element.xpath('./p:spPr/a:ln/a:noFill')
                       and tuple((sh.left,sh.top,sh.width,sh.height))==tuple(Inches(v) for v in T.background_region(slide)))
                if not valid: findings.append({'code':'SOURCE_BACKGROUND_ALTERED','output_slide':oi})
                continue
            if name.startswith('sss-repeat:'):
                sid = name[len('sss-repeat:'):]
                item = next((x for x in source.items if x.id == sid), None)
                title = source_prs.slides[item.slide].shapes.title if item else None
                valid_repeat = item and title is not None and item.origin == 'slide' and item.shape_id == title.shape_id and oi in mapping[str(item.slide)][1:] and (sid,oi) not in repeats
                repeats.add((sid,oi))
                if not valid_repeat or inventory.signature(sh, slide.part, numbers, connection_ids) != item.signature:
                    findings.append({'code': 'INVALID_STRUCTURAL_REPEAT', 'output_slide': oi, 'id': sid})
                continue
            if not name.startswith('sss:'):
                findings.append({'code': 'UNEXPLAINED_OBJECT', 'output_slide': oi})
                continue
            sid = name[4:]
            if sid in extractions:
                from slide_engine import logo_assets
                expected=[p for p in report['placements'] if p['source_id']==sid and p['output_slide']==oi]
                if (not sh.name.endswith('|logo') or not logo_assets.verify_picture(sh,extractions[sid])
                        or len(expected)!=1 or list((sh.left,sh.top,sh.width,sh.height))!=expected[0].get('bounds')
                        or sh._element.xpath('.//a:hlinkClick | .//a:hlinkMouseOver')):
                    findings.append({'code':'EXTRACTED_LOGO_ALTERED','id':sid,'output_slide':oi,
                                     'criterion':'brand_consistency','severity':'blocking'})
                found.setdefault(sid,[]).append((oi,{'extracted_logo':True}))
                continue
            role=next((e for e in decisions.get(str(si),{}).get('elements',[]) if e['id']==sid),{})
            if role.get('role') in ('code','logo') or role.get('contains_logo'):
                item=next((x for x in source.items if x.id==sid),None)
                if item:
                    original=next(s for origin,_,s in inventory.source_objects(source_prs.slides[si])
                                  if origin==item.origin and s.shape_id==item.shape_id)
                    from slide_engine.text_style import protected_text
                    if protected_text(original)!=protected_text(sh):
                        findings.append({'code':'PROTECTED_TYPOGRAPHY_ALTERED','id':sid,'output_slide':oi,
                            'criterion':'content_accuracy' if role.get('role')=='code' else 'brand_consistency',
                            'severity':'blocking','message':'Protected source text, whitespace or typeface was altered.'})
            if sid in removals:
                findings.append({'code':'REJECTED_ARTWORK_RETAINED','id':sid,'output_slide':oi,'severity':'blocking'})
            expected_placements = [p for p in report['placements'] if p['source_id']==sid and p['output_slide']==oi]
            if len(expected_placements)!=1 or list((sh.left,sh.top,sh.width,sh.height)) != expected_placements[0].get('bounds'):
                findings.append({'code':'PLACEMENT_ALTERED','id':sid,'output_slide':oi})
            if sh._element.xpath('.//a:rPr/a:solidFill//a:alpha[@val="0"] | .//a:defRPr/a:solidFill//a:alpha[@val="0"]'):
                findings.append({'code':'INVISIBLE_TEXT','id':sid,'output_slide':oi})
            found.setdefault(sid, []).append((oi, inventory.signature(sh, slide.part, numbers, connection_ids)))
            if any(el.get('hidden') == '1' for el in sh._element.xpath('.//p:cNvPr')):
                findings.append({'code': 'HIDDEN_CONTENT', 'id': sid, 'output_slide': oi})
    for item in source.items:
        if item.disposition == 'excluded_by_policy':
            continue
        records = found.pop(item.id, [])
        if not records:
            findings.append({'code': 'MISSING', 'id': item.id, 'source_slide': item.slide})
            continue
        if any(oi not in mapping.get(str(item.slide), []) for oi,_ in records):
            findings.append({'code': 'WRONG_SLIDE', 'id': item.id})
        if item.id in extractions:
            if len(records)!=1:
                findings.append({'code':'EXTRACTED_LOGO_DUPLICATED','id':item.id,'severity':'blocking'})
            continue
        expected = item.signature
        if len(records) == 1:
            actual = records[0][1]
        else:
            actual = dict(records[0][1])
            actual['text'] = [t for _,r in records for t in r['text']]
            actual['emphasis'] = [t for _,r in records for t in r['emphasis']]
            actual['links'] = [t for _,r in records for t in r['links']]
            if 'table' in expected:
                for _,r in records[1:]:
                    if not r.get('table') or r['table'][0]!=expected['table'][0]:
                        findings.append({'code':'TABLE_HEADER_ALTERED','id':item.id})
                actual['table'] = [row for index,(_,r) in enumerate(records) for row in (r['table'] if index==0 else r['table'][1:])]
            if any({k:v for k,v in r.items() if k not in ('text','links','emphasis','table')} !=
                   {k:v for k,v in expected.items() if k not in ('text','links','emphasis','table')} for _,r in records):
                findings.append({'code': 'DUPLICATED_OR_ALTERED_OBJECT', 'id': item.id})
        if actual != expected:
            changed=sorted(k for k in set(expected)|set(actual) if expected.get(k)!=actual.get(k))
            findings.append({'code': 'ALTERED_OR_DUPLICATED', 'id': item.id, 'source_slide': item.slide,
                             'message': f'Source slide {item.slide+1}, shape {item.shape_id}: exported {", ".join(changed)} differs from the source.',
                             'expected': expected, 'actual': actual})
    findings.extend({'code': 'INVENTED', 'id': sid} for sid in found)
    for si, expected in source.notes.items():
        outputs = mapping.get(str(si), [])
        if not outputs or inventory.user_notes(prs.slides[outputs[0]]) != expected:
            findings.append({'code': 'NOTES_ALTERED', 'source_slide': si})
        elif inventory.notes_links(prs.slides[outputs[0]],numbers) != inventory.notes_links(source_prs.slides[si],source_numbers):
            findings.append({'code':'NOTES_LINK_ALTERED','source_slide':si})
        for oi in outputs[1:]:
            if inventory.user_notes(prs.slides[oi]):
                findings.append({'code': 'NOTES_DUPLICATED', 'source_slide': si, 'output_slide': oi})
    findings.extend({'code': 'UNSUPPORTED', **x} for x in source.unsupported)
    return {'status': 'failed' if findings else 'passed', 'findings': findings,
            'candidate_sha256': inventory.sha256(candidate_path), 'source_sha256': source.source_sha256,
            'counts': dict(Counter(x.signature['kind'] for x in source.items)), 'items': len(source.items)}
