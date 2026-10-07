"""Deterministic geometry QA: catch overlap / off-slide / text-overflow before
export. Runs without rendering, so it works even when LibreOffice is absent.

This directly serves the locked rule: no text out of the box, no overlap.
"""
from __future__ import annotations

from io import BytesIO
import math

from pptx.util import Emu
from pptx.enum.text import MSO_AUTO_SIZE

# Rough per-character width factor (as a fraction of font size) for Arial, and
# line-height factor. Used only to *estimate* overflow so we can shrink/flag;
# real fit is enforced by PowerPoint autofit at open time as a backstop.
_CHAR_W = 0.52
_LINE_H = 1.22


def _box(sh):
    try:
        return (int(sh.left or 0), int(sh.top or 0),
                int(sh.width or 0), int(sh.height or 0))
    except Exception:
        return (0, 0, 0, 0)


def _intersect_area(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    return ix * iy


def overlaps(a, b, tol_frac=0.12):
    """True if two boxes overlap by more than tol_frac of the smaller area."""
    area = _intersect_area(a, b)
    if area <= 0:
        return False
    sa = max(1, a[2] * a[3])
    sb = max(1, b[2] * b[3])
    return area / min(sa, sb) > tol_frac


def white_picture_background(picture, text_bounds, picture_bounds):
    """Prove that the *whole* text box lies on blank white image pixels.

    PDF graphics can include a page-sized frame around a diagram, with native
    title text deliberately placed in its empty white area. Its bounding box
    intersects the title, but its visible artwork does not. This narrow check
    never treats transparency, near-white pixels, cropped artwork, or unknown
    picture effects as evidence of a clear background. Callers must also check
    z-order, text color and untransformed geometry.
    """
    try:
        x, y, w, h = text_bounds
        px, py, pw, ph = picture_bounds
        if min(w, h, pw, ph) <= 0 or x < px or y < py or x+w > px+pw or y+h > py+ph:
            return False
        # Effects/color transforms can change pixels after the embedded image
        # is decoded. Keep the ordinary warning when their result is unknown.
        if picture._element.xpath('./p:blipFill/a:blip/* | ./p:spPr/a:effectLst/* | ./p:spPr/a:effectDag'):
            return False
        crop = (picture.crop_left, picture.crop_top, picture.crop_right, picture.crop_bottom)
        cl, ct, cr, cb = crop
        if min(crop) < 0 or cl+cr >= 1 or ct+cb >= 1:
            return False
        from PIL import Image
        with Image.open(BytesIO(picture.image.blob)) as image:
            # Source-rectangle cropping scales the remaining pixels to the
            # picture frame. Include a pixel beyond each text edge so borders
            # and antialiased artwork touching the text still require review.
            iw, ih = image.size
            left = math.floor((cl+(x-px)/pw*(1-cl-cr))*iw)-1
            top = math.floor((ct+(y-py)/ph*(1-ct-cb))*ih)-1
            right = math.ceil((cl+(x+w-px)/pw*(1-cl-cr))*iw)+1
            bottom = math.ceil((ct+(y+h-py)/ph*(1-ct-cb))*ih)+1
            if left < 0 or top < 0 or right > iw or bottom > ih:
                return False
            pixels = image.crop((left,top,right,bottom)).convert('RGBA')
            return pixels.getextrema() == ((255,255),)*4
    except (AttributeError, OSError, ValueError, TypeError, OverflowError):
        return False


def estimate_overflow(sh):
    """Estimate whether text overflows its box. Returns (overflow, fill_ratio)."""
    if not sh.has_text_frame:
        return (False, 0.0)
    tf = sh.text_frame
    left, top, width, height = _box(sh)
    if width <= 0 or height <= 0:
        return (False, 0.0)
    ml = int(getattr(tf, "margin_left", Emu(91440)) or 0)
    mr = int(getattr(tf, "margin_right", Emu(91440)) or 0)
    mt = int(getattr(tf, "margin_top", Emu(45720)) or 0)
    mb = int(getattr(tf, "margin_bottom", Emu(45720)) or 0)
    avail_w = max(1, width - ml - mr)
    avail_h = max(1, height - mt - mb)

    total_lines = 0.0
    max_size = 12.0
    for p in tf.paragraphs:
        size = None
        text = ""
        for r in p.runs:
            text += r.text or ""
            if r.font.size:
                size = max(size or 0, r.font.size.pt)
        size = size or 18.0
        max_size = max(max_size, size)
        if not text.strip():
            total_lines += 0.6
            continue
        char_w_emu = size * _CHAR_W * 12700  # pt -> EMU (1pt = 12700 EMU)
        chars_per_line = max(1, int(avail_w / char_w_emu))
        total_lines += 1 if tf.word_wrap is False else max(1, -(-len(text) // chars_per_line))

    needed_h = total_lines * max_size * _LINE_H * 12700
    ratio = needed_h / avail_h if avail_h else 0.0
    return (ratio > 1.02, ratio)


def check_slide(slide, slide_w, slide_h):
    """Return a list of geometry issue dicts for one slide."""
    issues = []
    shapes = [sh for sh in slide.shapes]

    # off-slide bleed
    for sh in shapes:
        x, y, w, h = _box(sh)
        if w == 0 and h == 0:
            continue
        if x < -9144 or y < -9144 or (x + w) > slide_w + 9144 or (y + h) > slide_h + 9144:
            issues.append({"type": "offslide", "shape_id": sh.shape_id})

    # overflow
    for sh in shapes:
        over, ratio = estimate_overflow(sh)
        if over:
            issues.append({"type": "overflow", "shape_id": sh.shape_id,
                           "ratio": round(ratio, 2)})

    # overlap (text-bearing shapes only, to avoid decorative bands)
    text_shapes = [sh for sh in shapes
                   if sh.has_text_frame and sh.text_frame.text.strip()]
    for i in range(len(text_shapes)):
        for j in range(i + 1, len(text_shapes)):
            if overlaps(_box(text_shapes[i]), _box(text_shapes[j])):
                issues.append({"type": "overlap",
                               "shape_id": text_shapes[i].shape_id,
                               "other_id": text_shapes[j].shape_id})
    return issues


def enforce_autofit(slide):
    """Backstop: make every text box wrap + shrink-to-fit so PowerPoint will not
    render overflowing text. Deterministic and safe."""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        tf = sh.text_frame
        try:
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except Exception:
            pass
