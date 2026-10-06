"""Required evidence from a fresh render of the candidate bytes."""
import json
import unicodedata
from collections import Counter
from pathlib import Path
import fitz
from pptx import Presentation
from slide_engine.inventory import sha256, normalize, walk_shapes
from slide_engine.ir import _group_xf
from pptx.oxml.ns import qn
from .. import rendering


def reordered_math_glyphs(expected, rendered):
    """Some PDF exporters emit fallback math glyphs after surrounding text.

    This only identifies uncertainty for visual review, never a verified match.
    Every symbol must be present in the same object region and all remaining
    wording must still occur in order. Missing or changed symbols stay failures.
    """
    symbols = {c for c in expected if ord(c)>127 and unicodedata.category(c)=='Sm'}
    if not symbols or any(rendered.count(c)<expected.count(c) for c in symbols):
        return False
    strip = lambda text: ''.join(c for c in text if c not in symbols)
    wording = strip(expected)
    return bool(wording) and wording in strip(rendered)


def check(candidate, directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=False)
    pdf = directory / 'candidate.pdf'
    digest = sha256(candidate)
    rendering.render_to_pdf(candidate, str(pdf))
    rendering.rasterize_pdf(str(pdf), lambda i: str(directory / f'slide-{i}.png'))
    prs = Presentation(candidate)
    findings, pages = [], []
    with fitz.open(pdf) as doc:
        if len(doc) != len(prs.slides):
            findings.append({'code': 'RENDER_PAGE_COUNT', 'severity': 'blocking'})
        for i, slide in enumerate(prs.slides):
            if i >= len(doc):
                continue
            page = doc[i]
            rendered = normalize(page.get_text(sort=True).replace('-\n', '-'))
            full_text=''.join(page.get_text(sort=False).split())
            image_regions=[fitz.Rect(info['bbox']) for info in page.get_image_info()]
            counts=Counter()
            for _,shape in walk_shapes(slide.shapes):
                if shape.has_text_frame:
                    counts.update(''.join(p.text.split()) for p in shape.text_frame.paragraphs if normalize(p.text))
                if shape.has_table:
                    counts.update(''.join(c.text.split()) for row in shape.table.rows for c in row.cells if not c.is_spanned and normalize(c.text))
            expected = []
            def visit(shapes, transform=lambda l,t,w,h:(l,t,w,h)):
                for sh in shapes:
                    if sh._element.tag == qn('p:grpSp'):
                        group = _group_xf(sh)
                        visit(sh.shapes, lambda l,t,w,h: transform(*group(l,t,w,h)))
                        continue
                    if any(v is None for v in (sh.left,sh.top,sh.width,sh.height)):
                        continue
                    l,t,w,h=transform(sh.left,sh.top,sh.width,sh.height)
                    rect=fitz.Rect(l/prs.slide_width*page.rect.width,t/prs.slide_height*page.rect.height,
                                   (l+w)/prs.slide_width*page.rect.width,(t+h)/prs.slide_height*page.rect.height)
                    image_rect=fitz.Rect(rect)
                    rect += (-2,-2,2,2)
                    if sh._element.tag == qn('p:pic'):
                        # PowerPoint may emit vector/metafile pictures as drawing paths;
                        # ordinary raster pictures must have a matching PDF image placement.
                        area=max(1,image_rect.get_area())
                        matched=any((image_rect & box).get_area()/area > .7 and
                                    .5 < box.get_area()/area < 1.5 for box in image_regions)
                        if not matched:
                            try:raster=sh.image.content_type in ('image/png','image/jpeg','image/gif','image/tiff','image/bmp')
                            except (AttributeError,ValueError):raster=False
                            findings.append({'code':'RENDER_IMAGE_UNVERIFIED','severity':'blocking' if raster else 'review',
                                'output_slide':i,'message':'No matching rendered image placement was found; inspect the picture.',
                                'evidence':str(directory/f'slide-{i}.png')})
                    units=[]
                    if sh.has_text_frame:
                        units.append((rect,[normalize(p.text) for p in sh.text_frame.paragraphs if normalize(p.text)]))
                    if sh.has_table:
                        # Table extraction interleaves adjacent columns. Verify each cell in its own region.
                        table=sh.table;y=sh.top
                        for ri,row in enumerate(table.rows):
                            x=sh.left
                            for ci,c in enumerate(row.cells):
                                cw=table.columns[ci].width
                                if not c.is_spanned and normalize(c.text):
                                    ww=sum(table.columns[j].width for j in range(ci,min(len(table.columns),ci+c.span_width)))
                                    hh=sum(table.rows[j].height for j in range(ri,min(len(table.rows),ri+c.span_height)))
                                    cl,ct,cw2,ch=transform(x,y,ww,hh)
                                    cr=fitz.Rect(cl/prs.slide_width*page.rect.width,ct/prs.slide_height*page.rect.height,
                                        (cl+cw2)/prs.slide_width*page.rect.width,(ct+ch)/prs.slide_height*page.rect.height)
                                    cr+=(-2,-2,2,2)
                                    units.append((cr,[normalize(c.text)]))
                                x+=cw
                            y+=row.height
                    # Match each object in its actual region. Ignore PDF-only wrapping whitespace;
                    # exact wording/punctuation/order is independently checked in the PPTX.
                    for region,texts in units:
                        remainder=''.join(page.get_textbox(region).split())
                        for text in sorted(texts,key=len,reverse=True):
                            expected.append(text)
                            compact=''.join(text.split())
                            if compact in remainder:
                                remainder=remainder.replace(compact,'',1)
                            else:
                                visible = full_text.count(compact) >= counts[compact]
                                glyph = len(compact)>1 and compact[0] in '▪■•·' and compact[1:] in remainder
                                math_order = reordered_math_glyphs(compact, remainder)
                                findings.append({'code':'RENDER_TEXT_OUTSIDE_OBJECT' if visible else 'RENDER_GLYPH_UNVERIFIED' if glyph else 'RENDER_MATH_ORDER_UNVERIFIED' if math_order else 'RENDER_TEXT_UNVERIFIED',
                                    'severity':'review' if visible or glyph or math_order else 'blocking',
                                    'output_slide':i,'message':'Text is present in the PDF but extends outside its object region; inspect fit and occlusion.' if visible else 'Text is present but its leading list glyph could not be verified; inspect the rendered marker.' if glyph else 'The wording and math symbols are present in this region, but PDF extraction reordered the symbols. Inspect their rendered positions.' if math_order else 'Expected text was not found inside its rendered object region.',
                                    'expected':text,'evidence':str(directory/f'slide-{i}.png')})
            visit(slide.shapes)
            spans = [s for b in page.get_text('dict')['blocks'] if 'lines' in b for line in b['lines'] for s in line['spans']]
            if any(not s['font'].startswith(('Arial','Wingdings','Symbol')) for s in spans if normalize(s['text'])):
                findings.append({'code':'RENDER_FONT_SUBSTITUTION','severity':'review','output_slide':i,
                                 'message':'The renderer used a font outside the supported Arial/bullet families.'})
            for span in spans:
                if normalize(span['text']) and span['size'] < 10.5:
                    findings.append({'code': 'RENDER_SMALL_TEXT', 'severity': 'review', 'output_slide': i,
                                     'message': 'Rendered text is below the supported minimum; inspect the slide.'})
                if not page.rect.contains(fitz.Rect(span['bbox'])):
                    findings.append({'code': 'RENDER_TEXT_BOUNDS', 'severity': 'blocking', 'output_slide': i})
            pages.append({'output_slide': i, 'success': True, 'png': str(directory / f'slide-{i}.png'),
                          'text_units': len(expected), 'fonts': sorted({s['font'] for s in spans})})
    if sha256(candidate) != digest:
        raise RuntimeError('Candidate changed during rendering')
    metadata = json.loads(pdf.with_suffix('.renderer.json').read_text(encoding='utf-8'))
    return {'status': 'failed' if any(f['severity'] == 'blocking' for f in findings) else
            'needs_review' if findings else 'passed', 'findings': findings, 'pages': pages,
            'candidate_sha256': digest, 'pdf_sha256': sha256(pdf), 'renderer': metadata}
