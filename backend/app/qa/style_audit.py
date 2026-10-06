"""Recursive formatting evidence, including cells and transformed groups."""
import math
from pptx.oxml.ns import qn
from slide_engine.ir import _group_xf
from .. import brand as B


def rgb(color):
    try:
        return str(color.rgb).upper()
    except (AttributeError, TypeError):
        return None


def contrast(a,b):
    def lum(color):
        channels=[int(color[i:i+2],16)/255 for i in (0,2,4)]
        channels=[x/12.92 if x<=.04045 else ((x+.055)/1.055)**2.4 for x in channels]
        return sum(x*y for x,y in zip(channels,(.2126,.7152,.0722)))
    x,y=sorted((lum(a),lum(b)))
    return (y+.05)/(x+.05)


def audit(slide, width, height):
    issues, text_boxes, visuals = [], [], []

    def finding(sev, kind, note, object_id=None):
        issues.append({'sev':sev,'type':kind,'note':note,
                       **({'object_ids':[object_id]} if object_id else {})})

    def visit(shapes, transform=lambda l,t,w,h:(l,t,w,h), depth=0, protected=False):
        for sh in shapes:
            protected_style=protected or sh.name.endswith(('|logo','|code'))
            raw=(sh.left,sh.top,sh.width,sh.height)
            if any(v is None for v in raw):
                finding('warn','unresolved_geometry',f'Geometry unavailable for shape {sh.shape_id}')
                continue
            l,t,w,h=transform(*raw)
            angle=math.radians(sh.rotation or 0)
            bw=abs(w*math.cos(angle))+abs(h*math.sin(angle))
            bh=abs(w*math.sin(angle))+abs(h*math.cos(angle))
            bounds=(l+(w-bw)/2,t+(h-bh)/2,bw,bh)
            if min(bounds[:2]) < -9144 or bounds[0]+bw > width+9144 or bounds[1]+bh > height+9144:
                finding('fail','offslide',f'Shape {sh.shape_id} extends outside the slide.')
            if sh._element.tag == qn('p:grpSp'):
                if sh.rotation:
                    finding('warn','group_rotation','Rotated group bounds require visual inspection.')
                local=_group_xf(sh)
                visit(sh.shapes,lambda a,b,c,d:transform(*local(a,b,c,d)),depth+1,protected_style)
                continue
            frames=[]
            if sh.has_table and (abs(sum(c.width for c in sh.table.columns)-sh.width)>9144 or
                                 abs(sum(r.height for r in sh.table.rows)-sh.height)>9144):
                finding('fail','table_geometry','Table grid differs from its frame; resize rows and columns with the table.',sh.name)
            fill=None
            try:
                fill=rgb(sh.fill.fore_color)
            except (AttributeError,TypeError,ValueError):pass
            if sh.has_text_frame:
                frames.append((sh.text_frame,fill,False))
                if sh.text_frame.text.strip():text_boxes.append((bounds,sh.shape_id,depth))
                if sh._element.xpath('./p:spPr/a:gradFill | ./p:spPr/a:blipFill | ./p:spPr/a:pattFill'):
                    finding('warn','uncertain_background',f'Inspect contrast on the patterned/gradient background of shape {sh.shape_id}.')
            if sh._element.tag == qn('p:pic'):
                visuals.append((bounds,sh.shape_id))
            if sh.has_chart:
                finding('warn','chart_style','Native chart data and workbook are preserved; inspect chart typography and visual legibility.')
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.is_spanned:continue
                        try:cell_fill=rgb(cell.fill.fore_color)
                        except (AttributeError,TypeError,ValueError):cell_fill=None
                        frames.append((cell.text_frame,cell_fill,True))
            for tf,background,table in frames:
                for p in tf.paragraphs:
                    if not p.text.strip():continue
                    ppr=p._p.find(qn('a:pPr'))
                    if ppr is not None and ppr.find(qn('a:buChar')) is not None:
                        bullet=ppr.find(qn('a:buChar')).get('char')
                        expected=B.BULLETS.get(min(p.level,2))
                        font=ppr.find(qn('a:buFont'))
                        if expected and (bullet != expected[0] or font is None or font.get('typeface') != expected[1]):
                            finding('warn','bullet','Source list marker differs from the approved hierarchy.')
                    for r in p.runs:
                        if not r.text.strip():continue
                        name=r.font.name or p.font.name
                        size=r.font.size or p.font.size
                        color=rgb(r.font.color) or rgb(p.font.color)
                        if not name or size is None or color is None:
                            finding('warn','unresolved_style',f'Inherited font/size/color needs review on shape {sh.shape_id}.')
                        if name and name != B.FONT and not protected_style:
                            finding('warn','font',f'Font {name} differs from Arial.')
                        from slide_engine import template_policy as T
                        fitted_bookend = (T.is_cover(slide) or T.is_closing(slide) or T.is_section(slide)) and size and 20 <= size.pt <= B.TITLE_PT
                        if sh.name.endswith('|title') and size and abs(size.pt-B.TITLE_PT)>.1 and not fitted_bookend:
                            finding('fail','title_size',f'Title must use {B.TITLE_PT} pt text.')
                        if size and size.pt < B.BODY_MIN_PT:
                            finding('warn','font_size',f'{size.pt:g} pt text is below the body minimum; verify its citation/caption role.')
                        if color and color not in B.ALLOWED_TEXT | {B.BLUE} and not protected_style:
                            finding('warn','source_color',f'Preserved source text color #{color} requires a scoped brand exception.')
                        if background==B.RED and color and color!=B.WHITE:
                            finding('fail','contrast','Text on the approved red fill must be white.')
                        if color and background and color==background:
                            finding('fail','contrast','Text and background have the same color.')
                        elif color and background and contrast(color,background)<3:
                            finding('fail','contrast','Text contrast against the explicit solid fill is below 3:1.')
                if tf._txBody.xpath('.//a:normAutofit'):
                    finding('warn','autofit','PowerPoint may shrink this text; check its rendered size.')
            if sh.has_text_frame and sh.text.strip():
                from .geometry import estimate_overflow
                over,ratio=estimate_overflow(sh)
                if over:finding('warn','overflow',f'Text fit estimate {ratio:.2f}; enlarge the text box and inspect the actual render.',sh.name)
    visit(slide.shapes)
    from .geometry import overlaps
    for i,(a,aid,ad) in enumerate(text_boxes):
        for b,bid,bd in text_boxes[i+1:]:
            if overlaps(a,b):
                finding('warn','occlusion',f'Text shapes {aid} and {bid} intersect; inspect for intentional containment or occlusion.')
        for b,bid in visuals:
            if overlaps(a,b):
                finding('warn','visual_occlusion',f'Text shape {aid} intersects image {bid}; inspect its legibility against the visual.')
    return issues
