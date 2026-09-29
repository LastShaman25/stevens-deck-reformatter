"""Explicitly authorized closing page, separate from source-slide provenance."""
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from . import template_policy as T

TITLE = 'Thank you!'
MARKER = 'template-required-closing|title'
AUTHORIZATION = 'User-required final Thank you slide using the approved Stevens statue-photo closing template; added after all original slides.'


def has_thanks(slide):
    return any(s.has_text_frame and re.fullmatch(r'thank\s+you[!?.\s]*',s.text.strip(),re.I)
               for s in slide.shapes)


def ensure(candidate, report):
    prs=Presentation(candidate)
    report['require_closing']=True
    if len(prs.slides) and T.is_closing(prs.slides[-1]) and has_thanks(prs.slides[-1]):
        return report
    layout=next(l for l in prs.slide_layouts if l.name==T.CLOSING_LAYOUT)
    for shape in list(layout.shapes):
        if shape.has_text_frame and shape.text.strip(): shape._element.getparent().remove(shape._element)
    slide=prs.slides.add_slide(layout)
    for shape in list(slide.shapes): slide.shapes._spTree.remove(shape._element)
    title=slide.shapes.add_textbox(*(Inches(v) for v in T.CLOSING_TITLE))
    title.name=MARKER;title.text=TITLE
    tf=title.text_frame;tf.word_wrap=False;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    run=tf.paragraphs[0].runs[0]
    run.font.name='Arial';run.font.size=Pt(40);run.font.color.rgb=RGBColor.from_string('FFFFFF')
    index=len(prs.slides)-1
    report['added_slides']=[{'output_slide':index,'kind':'closing','text':TITLE,'authorization':AUTHORIZATION}]
    report['slide_count']=len(prs.slides)
    report['slides'].append({'index':index,'source_index':None,'kind':'closing','layout':T.CLOSING_LAYOUT,
                             'template_contract':T.contract(slide),'hard_issues':0})
    prs.save(candidate)
    return report


def validate_added(prs,report):
    """Only one exact, bounded closing addition is exempt from source mapping."""
    added=report.get('added_slides',[])
    if not added:return set()
    expected={'output_slide':len(prs.slides)-1,'kind':'closing','text':TITLE,'authorization':AUTHORIZATION}
    if not report.get('require_closing') or added!=[expected]:raise ValueError('Invalid added-slide authorization.')
    slide=prs.slides[-1]
    if not T.is_closing(slide) or len(slide.shapes)!=1 or not has_thanks(slide):raise ValueError('Required closing altered.')
    shape=slide.shapes[0]
    if shape.name!=MARKER or shape.text!=TITLE or shape._element.xpath('.//a:hlinkClick | .//a:hlinkMouseOver'):
        raise ValueError('Required closing content altered.')
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():raise ValueError('Unexpected added notes.')
    return {len(prs.slides)-1}


def check(prs,report):
    findings=[]
    try: validate_added(prs,report)
    except ValueError as exc:findings.append({'code':'INVALID_ADDED_CLOSING','severity':'blocking','message':str(exc)})
    if report.get('require_closing') and (not len(prs.slides) or not T.is_closing(prs.slides[-1]) or not has_thanks(prs.slides[-1])):
        findings.append({'code':'REQUIRED_CLOSING_MISSING','severity':'blocking','criterion':'instruction_compliance',
                         'message':'End the deck with Thank you! on the approved Stevens statue-photo closing layout.'})
    return findings
