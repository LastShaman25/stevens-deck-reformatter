"""Source-owned inventory. No builder or planner claims are used here."""
from __future__ import annotations

import hashlib
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from types import SimpleNamespace

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normalize(value):
    return re.sub(r'\s+', ' ', value or '').strip()


def canonical(element):
    # Copied shapes inherit the destination slide's namespace declarations.
    # Only namespaces used by this subtree are part of its geometry; unrelated
    # declarations (such as r on a slide with links) must not imply alteration.
    return etree.tostring(element, method='c14n', exclusive=True).decode()


def walk_shapes(shapes, prefix=''):
    for shape in shapes:
        path = f'{prefix}/{shape.shape_id}'
        yield path, shape
        if shape._element.tag == qn('p:grpSp'):
            yield from walk_shapes(shape.shapes, path)


def source_objects(slide):
    """Include inherited non-placeholder content, not authoring prompts."""
    for owner, label in [(slide.slide_layout.slide_master, 'master'),
                         (slide.slide_layout, 'layout'), (slide, 'slide')]:
        if label != 'slide' and slide._element.get('showMasterSp') == '0':
            continue
        for shape in owner.shapes:
            if label != 'slide' and shape.is_placeholder:
                continue
            yield label, owner.part, shape


def paragraphs(shape):
    if not shape.has_text_frame:
        return []
    return [normalize(p.text) for p in shape.text_frame.paragraphs if normalize(p.text)]


def emphasis(run, paragraph, shape, key):
    from .text_style import resolve
    return bool(resolve(run,paragraph,shape,key))


def text_emphasis(shape):
    if not shape.has_text_frame:
        return []
    result = []
    for p in shape.text_frame.paragraphs:
        if not normalize(p.text):
            continue
        chars = []
        for r in p.runs:
            style = tuple(emphasis(r,p,shape,k) for k in ('bold','italic','underline'))
            chars.extend((char, *style) for char in r.text if not char.isspace())
        result.append(chars)
    return result


def user_notes(slide):
    if not slide.has_notes_slide:
        return ''
    tf = slide.notes_slide.notes_text_frame
    return normalize(tf.text) if tf is not None else ''


def notes_links(slide, numbers):
    if not slide.has_notes_slide or slide.notes_slide.notes_text_frame is None:
        return []
    element = slide.notes_slide.notes_text_frame._txBody.getparent()
    return links(SimpleNamespace(_element=element), slide.notes_slide.part, numbers)


def relationship_signature(part, seen=None):
    """Content fingerprints for chart/workbook dependencies, independent of rIds."""
    seen = set() if seen is None else seen
    if part in seen:
        return 'cycle'
    seen = seen | {part}
    return {
        'content_type': part.content_type,
        'sha256': hashlib.sha256(part.blob).hexdigest(),
        'dependencies': sorted([
            {'type': rel.reltype, 'target': rel.target_ref if rel.is_external else
             relationship_signature(rel.target_part, seen)}
            for rel in part.rels.values()
        ], key=lambda x: str(x)),
    }


def links(shape, part, slide_numbers):
    result = []
    for element in shape._element.iter():
        if element.tag not in (qn('a:hlinkClick'), qn('a:hlinkMouseOver')):
            continue
        rid = element.get(qn('r:id'))
        if not rid:
            result.append({'target': element.get('action', ''), 'text': ''})
            continue
        rel = part.rels[rid]
        target = rel.target_ref if rel.is_external else str(rel.target_part.partname)
        if not rel.is_external and target in slide_numbers:
            target = f'slide:{slide_numbers[target]}'
        parent = element.getparent()
        # Range text lives in the run enclosing rPr.
        run = parent.getparent() if parent is not None else None
        result.append({'target': target, 'text': normalize(''.join(run.itertext())) if run is not None else '',
                       'event': etree.QName(element).localname, 'action': element.get('action', '')})
    return result


def signature(shape, part, slide_numbers, connection_ids=None):
    el = shape._element
    connection_ids = connection_ids or {}
    value = {'kind': etree.QName(el).localname, 'text': paragraphs(shape), 'emphasis': text_emphasis(shape),
             'links': links(shape, part, slide_numbers)}
    if el.tag == qn('p:grpSp'):
        value['children'] = [signature(s, part, slide_numbers, connection_ids) for s in shape.shapes]
    if shape.has_table:
        value['table'] = [[{'text': normalize(c.text), 'emphasis': text_emphasis(SimpleNamespace(has_text_frame=True,text_frame=c.text_frame)), 'span': [c.span_height, c.span_width],
                            'merged': c.is_spanned} for c in row.cells] for row in shape.table.rows]
    if shape.has_chart:
        value['chart'] = relationship_signature(shape.chart.part)
    for image in el.xpath('.//a:blip'):
        rid = image.get(qn('r:embed')) or image.get(qn('r:link'))
        if rid:
            rel = part.rels[rid]
            value.setdefault('images', []).append(rel.target_ref if rel.is_external else
                hashlib.sha256(rel.target_part.blob).hexdigest())
    value['connections'] = [{**dict(x.attrib), 'id': connection_ids.get(x.get('id'), x.get('id'))} for x in el.xpath('.//a:stCxn | .//a:endCxn')]
    value['geometry'] = [canonical(x) for x in el.xpath('./p:spPr/a:prstGeom | ./p:spPr/a:custGeom')]
    value['crops'] = [dict(x.attrib) for x in el.xpath('.//a:srcRect')]
    value['rotations'] = [dict((k, v) for k, v in x.attrib.items() if k in ('rot', 'flipH', 'flipV'))
                          for x in el.xpath('.//a:xfrm') if any(k in x.attrib for k in ('rot','flipH','flipV'))]
    return value


@dataclass
class Item:
    id: str
    slide: int
    origin: str
    shape_id: int
    signature: dict
    unsupported: list[str] = field(default_factory=list)
    disposition: str = 'preserved_native'
    policy_reason: str | None = None


@dataclass
class Inventory:
    source_sha256: str
    width: int
    height: int
    items: list[Item]
    notes: dict[int, str]
    slide_count: int
    unsupported: list[dict]

    def to_dict(self):
        return asdict(self)


def inspect(path):
    prs = Presentation(path)
    digest = sha256(path)
    numbers = {str(s.part.partname): i for i, s in enumerate(prs.slides)}
    items, unsupported, notes = [], [], {}
    for i, slide in enumerate(prs.slides):
        notes[i] = user_notes(slide)
        for origin, part, shape in source_objects(slide):
            sid = f'{digest[:16]}/{i}/{origin}/{shape.shape_id}'
            reasons = []
            for el in shape._element.iter():
                local = etree.QName(el).localname
                if local in {'oleObj', 'videoFile', 'audioFile', 'relIds', 'AlternateContent', 'oMath', 'oMathPara'}:
                    reasons.append(local)
            if shape.has_chart:
                # Standard OOXML chart parts are retained natively; extended charts need a separate adapter.
                if 'chartEx' in shape.chart.part.content_type:
                    reasons.append('extended_chart')
            if shape._element.tag not in {qn('p:sp'), qn('p:pic'), qn('p:grpSp'), qn('p:cxnSp'), qn('p:graphicFrame')}:
                reasons.append('unknown_shape')
            if shape._element.tag == qn('p:graphicFrame') and not shape.has_table and not shape.has_chart:
                reasons.append('unknown_graphic_frame')
            item = Item(sid, i, origin, shape.shape_id, signature(shape, part, numbers), sorted(set(reasons)))
            fields = shape._element.xpath('.//a:fld')
            if fields and all(f.get('type') == 'slidenum' for f in fields) and not shape._element.xpath('.//a:r/a:t') and not shape.has_table:
                item.disposition = 'excluded_by_policy'
                item.policy_reason = 'template-policy-1: generated slide-number field'
            items.append(item)
            unsupported.extend({'id': sid, 'slide': i, 'reason': r} for r in item.unsupported)
        if slide._element.find(qn('p:timing')) is not None:
            unsupported.append({'id': f'{digest[:16]}/{i}/timing', 'slide': i, 'reason': 'animation'})
        bg = slide._element.find('.//' + qn('p:bg'))
        if bg is not None and bg.xpath('.//a:blip'):
            unsupported.append({'id': f'{digest[:16]}/{i}/background', 'slide': i, 'reason': 'image_background'})
    return Inventory(digest, int(prs.slide_width), int(prs.slide_height), items, notes, len(prs.slides), unsupported)


def validate_placements(inventory, placements):
    from collections import Counter
    expected = {x.id for x in inventory.items if x.disposition != 'excluded_by_policy'}
    counts = Counter(p['source_id'] for p in placements)
    errors = [{'code': 'MISSING_PLACEMENT', 'id': i} for i in sorted(expected - counts.keys())]
    errors += [{'code': 'INVENTED_PLACEMENT', 'id': i} for i in sorted(counts.keys() - expected)]
    by_id = {x.id:x for x in inventory.items}
    for p in placements:
        item = by_id.get(p['source_id'])
        if item and (p.get('source_slide')!=item.slide or not isinstance(p.get('output_slide'),int) or p['output_slide']<0):
            errors.append({'code':'INVALID_PLACEMENT_SLIDE','id':item.id})
    for sid, count in counts.items():
        parts = [p for p in placements if p['source_id'] == sid]
        item = by_id.get(sid)
        ranges = [p.get('row_range') or p.get('paragraph_range') for p in parts]
        if count>1 or any(r is not None for r in ranges):
            if not item or any(r is None for r in ranges) or sorted(ranges)!=ranges:
                errors.append({'code':'DUPLICATED_PLACEMENT','id':sid})
                continue
            units = item.signature.get('table',[]) if parts[0].get('row_range') else item.signature['text']
            if ranges[0][0]!=0 or ranges[-1][1]!=len(units) or any(r[0]>=r[1] for r in ranges) or any(a[1]!=b[0] for a,b in zip(ranges,ranges[1:])):
                errors.append({'code':'INVALID_SPLIT','id':sid})
    errors += [{'code': 'UNSUPPORTED', **u} for u in inventory.unsupported]
    return {'status': 'failed' if errors else 'passed', 'findings': errors, 'source': len(expected)}
