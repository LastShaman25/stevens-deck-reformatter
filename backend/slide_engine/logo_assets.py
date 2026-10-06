"""Bounded, source-pixel logo extraction. Never redraw or infer missing lettering.

The vision agent proposes a region; this module verifies its raster boundaries.
Semantic completeness is separately reviewed against the full image before use.
"""
from hashlib import sha256
from io import BytesIO
from functools import lru_cache
from pathlib import Path
from zipfile import ZipFile
import math

import numpy as np
from PIL import Image, ImageDraw
from pptx.oxml.ns import qn

TEMPLATE = Path(__file__).resolve().parents[1] / 'assets' / 'stevens_template.pptx'
APPROVED_REGIONS = {'ppt/media/image2.png':(805,40,975,340),
                    'ppt/media/image4.png':(800,40,975,355),
                    'ppt/media/image7.png':(800,40,975,355),
                    'ppt/media/image9.png':(800,40,975,355)}


def raster(shape):
    if shape._element.tag != qn('p:pic'):
        raise ValueError('Logo extraction supports a standalone raster picture only; preserve native logo groups.')
    if shape._element.xpath('.//a:srcRect[@l or @t or @r or @b] | .//a:xfrm[@rot or @flipH or @flipV]'):
        raise ValueError('Cropped, flipped or rotated source pictures require inspection before extraction.')
    if shape._element.xpath('.//a:blip/* | .//a:effectLst/* | .//a:effectDag'):
        raise ValueError('Source picture effects cannot be discarded during logo extraction.')
    try:
        image = Image.open(BytesIO(shape.image.blob))
        if image.format not in ('PNG', 'JPEG', 'BMP', 'TIFF') or getattr(image, 'n_frames', 1) != 1:
            raise ValueError('Unsupported logo raster format.')
        if not shape.width or not shape.height or abs((shape.width/shape.height)/(image.width/image.height)-1)>.02:
            raise ValueError('The source picture is stretched; inspect its displayed logo proportions before extraction.')
        return image.convert('RGBA')
    except (OSError, AttributeError) as exc:
        raise ValueError('Source image is not a supported raster.') from exc


def _crop(image, region):
    """Region is left/top/right/bottom in 0..1000 ORIGINAL IMAGE coordinates."""
    if len(region) != 4 or any(type(v) is not int or not 0 <= v <= 1000 for v in region):
        raise ValueError('Logo region must contain four integer coordinates from 0 to 1000.')
    l,t,r,b = region
    if l >= r or t >= b or (r-l)*(b-t) > 350000:
        raise ValueError('Identify a bounded logo region, not another miniature background.')
    box = (math.floor(l*image.width/1000), math.floor(t*image.height/1000),
           math.ceil(r*image.width/1000), math.ceil(b*image.height/1000))
    crop = image.crop(box)
    if min(crop.size) < 36:
        raise ValueError('Logo crop has too few source pixels; keep the original for review.')
    a = np.asarray(crop).astype(np.int16)
    border = np.concatenate((a[:3].reshape(-1,4), a[-3:].reshape(-1,4),
                             a[:,:3].reshape(-1,4), a[:,-3:].reshape(-1,4)))
    if np.all(border[:,3] <= 2):
        foreground = a[:,:,3] > 8
    else:
        bg = np.median(border, axis=0)
        if bg[3] < 250 or np.max(np.abs(border-bg)) > 10:
            raise ValueError('Logo touches the crop boundary or has a textured background; extraction is uncertain.')
        foreground = (np.max(np.abs(a[:,:,:3]-bg[:3]), axis=2) > 18) & (a[:,:,3] > 8)
    ys,xs = np.where(foreground)
    if len(xs) < 100 or len(xs) > foreground.size*.65:
        raise ValueError('Region does not isolate a legible logo with clear space.')
    bounds = (int(xs.min()),int(ys.min()),int(xs.max())+1,int(ys.max())+1)
    if bounds[2]-bounds[0] < 48 or bounds[3]-bounds[1] < 32:
        raise ValueError('Logo lettering has insufficient source resolution.')
    # Leave a real source-pixel border. No masking, recoloring, upscaling or inpainting.
    pad = max(3, round(max(bounds[2]-bounds[0],bounds[3]-bounds[1])*.025))
    if min(bounds[0],bounds[1],crop.width-bounds[2],crop.height-bounds[3]) < pad:
        raise ValueError('Logo crop lacks verified clear space; enlarge its proposed region.')
    tight = (bounds[0]-pad,bounds[1]-pad,bounds[2]+pad,bounds[3]+pad)
    final_box = [box[0]+tight[0],box[1]+tight[1],box[0]+tight[2],box[1]+tight[3]]
    return image.crop(tuple(final_box)), final_box


def _features(image):
    """Foreground silhouette plus original RGB; matching cannot change a variant."""
    a = np.asarray(image.convert('RGB')).astype(float)
    bg = np.median(np.concatenate((a[0],a[-1],a[:,0],a[:,-1])),axis=0)
    mask = np.max(np.abs(a-bg),axis=2) > 18
    y,x = np.where(mask)
    if not len(x): return None
    crop = image.crop((int(x.min()),int(y.min()),int(x.max())+1,int(y.max())+1)).convert('RGB')
    thumb = np.asarray(crop.resize((96,96),Image.Resampling.LANCZOS)).astype(float)
    return crop.width/crop.height, thumb, bg


@lru_cache(maxsize=2)
def _approved_catalog(template_digest):
    # Approved artwork is sourced only from the bundled template, never an AI path.
    # Regions encompass the complete tower+wordmark+small text, with clear space.
    result=[]
    with ZipFile(TEMPLATE) as z:
        for name,box in APPROVED_REGIONS.items():
            if name not in z.namelist(): continue
            try:
                crop,_ = _crop(Image.open(BytesIO(z.read(name))).convert('RGBA'),box)
                result.append((name,_features(crop)))
            except ValueError:
                continue  # A changed template needs new approved regions, never guessed crops.
    return result


def approved_match(crop):
    value = _features(crop)
    if not value: return None
    matches=[]
    for name,reference in _approved_catalog(sha256(TEMPLATE.read_bytes()).hexdigest()):
        if reference is None: continue
        aspect,pixels,bg = reference
        if abs(value[0]/aspect-1) > .025 or np.max(np.abs(value[2]-bg)) > 12: continue
        error = float(np.mean(np.abs(value[1]-pixels)))
        # Identity hint only: output always retains exact SOURCE pixels, not the reference.
        if error <= 8: matches.append((error,name))
    return min(matches)[1] if matches else None


@lru_cache(maxsize=2)
def _raw_references(template_digest):
    result=[]
    with ZipFile(TEMPLATE) as z:
        for name,region in APPROVED_REGIONS.items():
            if name not in z.namelist(): continue
            with Image.open(BytesIO(z.read(name))) as im:
                result.append((name,region,im.width/im.height,np.asarray(im.convert('RGBA').resize((128,128))).astype(float)))
    return result


def suggested_regions(image):
    """Locate known artwork before asking vision to guess a tiny wordmark's bounds."""
    pixels=np.asarray(image.convert('RGBA').resize((128,128))).astype(float)
    matches=[]
    for name,region,aspect,reference in _raw_references(sha256(TEMPLATE.read_bytes()).hexdigest()):
        error=float(np.mean(np.abs(pixels-reference)))
        if abs((image.width/image.height)/aspect-1)<.01 and error<3:
            matches.append((error,name,region))
    if not matches:return []
    error,name,region=min(matches)
    return [{'region':list(region),'reference_asset':name,'normalized_rgba_error':round(error,3),
             'basis':'Full image visually matches bundled template artwork; suggested region still requires extraction checks and independent review.'}]


def template_logo(crop_blob, cover, actual_slide=None):
    """Reuse an identical approved variant already on this layout, without doubling it.

Requires BOTH a strict image match and the independent semantic crop review. When
auditing an exported slide, also verify its inherited artwork against the bundle.
"""
    from pptx import Presentation
    from . import inventory, template_policy as T
    from . import templates
    if templates.current_id()=='cpe':
        raise ValueError('CPE logo reuse is not verified; use source_crop or retain the original logo.')
    source=_features(Image.open(BytesIO(crop_blob)))
    p=Presentation(TEMPLATE)
    wanted = actual_slide.slide_layout.name if actual_slide is not None else T.OPENING_LAYOUT if cover else 'Title Only'
    reference=p.slides.add_slide(next(l for l in p.slide_layouts if l.name==wanted))
    with ZipFile(TEMPLATE) as z:
        approved={sha256(z.read(n)).hexdigest():(n,region) for n,region in APPROVED_REGIONS.items() if n in z.namelist()}
    for origin,_,shape in inventory.source_objects(reference):
        if origin=='slide' or shape._element.tag!=qn('p:pic'): continue
        entry=approved.get(sha256(shape.image.blob).hexdigest())
        if not entry: continue
        try:
            crop,box=_crop(raster(shape),entry[1]);feature=_features(crop)
        except ValueError: continue
        if not source or not feature or abs(source[0]/feature[0]-1)>.025 or np.max(np.abs(source[2]-feature[2]))>12 or np.mean(np.abs(source[1]-feature[1]))>8:
            continue
        if actual_slide is not None:
            actual=next((s for o,_,s in inventory.source_objects(actual_slide) if o==origin and s.shape_id==shape.shape_id),None)
            if actual is None or actual._element.tag!=qn('p:pic'): continue
            try:
                raster(actual)
                if (actual.image.blob!=shape.image.blob or
                    (actual.left,actual.top,actual.width,actual.height)!=(shape.left,shape.top,shape.width,shape.height) or
                    inventory.canonical(actual._element)!=inventory.canonical(shape._element)): continue
            except ValueError: continue
        stream=BytesIO();crop.save(stream,format='PNG')
        iw,ih=shape.image.size
        bounds=[int(shape.left+box[0]/iw*shape.width),int(shape.top+box[1]/ih*shape.height),
                int((box[2]-box[0])/iw*shape.width),int((box[3]-box[1])/ih*shape.height)]
        return stream.getvalue(),{'asset':entry[0],'origin':origin,'shape_id':shape.shape_id,'bounds':bounds,
                                 'pixel_size':list(crop.size),'relative_aspect_error':round(abs(source[0]/feature[0]-1),6),
                                 'normalized_rgb_error':round(float(np.mean(np.abs(source[1]-feature[1]))),3),
                                 'crop_sha256':sha256(stream.getvalue()).hexdigest()}
    raise ValueError('No identical approved logo variant is verifiably present on this template layout; use a verified source crop or retain the original.')


def extract(shape, region):
    crop,box = _crop(raster(shape),region)
    stream=BytesIO();crop.save(stream,format='PNG');blob=stream.getvalue()
    return blob, {'method':'source_pixel_crop','source_image_sha256':sha256(shape.image.blob).hexdigest(),
                  'crop_sha256':sha256(blob).hexdigest(),'pixel_box':box,'pixel_size':list(crop.size),
                  'approved_template_match':approved_match(crop),
                  'background':'Original flat background/transparent padding retained; no generated pixels.'}


def verify_picture(shape, blob):
    """Independent gate: exact pixels, no crop/effects/rotation, original aspect ratio."""
    try:
        actual=raster(shape)
        expected=Image.open(BytesIO(blob)).convert('RGBA')
        return (actual.size == expected.size and actual.tobytes() == expected.tobytes()
                and abs((shape.width/shape.height)/(expected.width/expected.height)-1) <= .005
                and min(expected.width/(shape.width/914400),expected.height/(shape.height/914400)) >= 96
                and not shape._element.xpath('.//p:cNvPr[@hidden="1"]'))
    except (ValueError, ZeroDivisionError):
        return False


def comparison_sheet(source_blob, template_blob):
    """Review-only equal-scale views; never substitutes for exported source pixels."""
    sheet=Image.new('RGB',(800,440),'#eeeeee');draw=ImageDraw.Draw(sheet)
    for x,label,blob in [(0,'SOURCE CROP',source_blob),(400,'EXISTING TEMPLATE MARK',template_blob)]:
        im=Image.open(BytesIO(blob)).convert('RGBA')
        scale=min(340/im.width,350/im.height)
        im=im.resize((round(im.width*scale),round(im.height*scale)),Image.Resampling.LANCZOS)
        sheet.paste(im,(x+(400-im.width)//2,45+(350-im.height)//2),im)
        draw.text((x+20,15),label,fill='black')
    draw.text((20,410),'Review views only: uniform scaling preserves each mark\'s proportions.',fill='black')
    stream=BytesIO();sheet.save(stream,format='PNG');return stream.getvalue()
