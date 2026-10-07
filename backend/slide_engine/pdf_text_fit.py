"""Bounded native PDF text fitting after conversion to the template font."""
from functools import lru_cache
import math

from PIL import ImageFont
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn


@lru_cache(maxsize=128)
def _face(size, bold, italic):
    from matplotlib.font_manager import findfont, FontProperties
    path=findfont(FontProperties(family='Arial',weight='bold' if bold else 'normal',
                                 style='italic' if italic else 'normal'))
    return ImageFont.truetype(path,size)


def _line_width(runs):
    cursor=0; left=0; right=0
    for run in runs:
        face=_face(max(1,round(run.font.size.pt*4)),bool(run.font.bold),bool(run.font.italic))
        bounds=face.getbbox(run.text)
        left=min(left,cursor+bounds[0]);right=max(right,cursor+bounds[2])
        cursor+=face.getlength(run.text)
    return max(right,cursor)-left


def fit_native_lines(slide):
    """Fit isolated PDF lines whose original font has wider Arial replacements.

    Keep each line's position and all run emphasis/relative sizes. Wrapping a
    caption or equation independently can collide with its figure, so shrink
    only measurable native Arial lines, never below the readable 11pt floor.
    Uncertain/too-small cases remain for the ordinary render verifier.
    """
    for shape in slide.shapes:
        if not shape.has_text_frame or shape.name.endswith(('|logo','|code')): continue
        tf=shape.text_frame
        if tf.word_wrap is not False or len(tf.paragraphs)!=1: continue
        paragraph=tf.paragraphs[0]
        runs=[r for r in paragraph.runs if r.text]
        if not runs or any(r.font.name!='Arial' or not r.font.size or
                           any(c in r.text for c in '\t\n\v') or
                           (r._r.rPr is not None and r._r.rPr.get('baseline') not in (None,'0'))
                           for r in runs): continue
        width=(shape.width-tf.margin_left-tf.margin_right)/12700*4
        if width<=0: continue
        measured=_line_width(runs)
        if measured<=width*.98: continue
        factor=width*.96/measured
        if min(r.font.size.pt for r in runs)*factor<11: continue
        for run in runs:
            run.font.size=Pt(math.floor(run.font.size.pt*factor*100)/100)


def enlarge_small_text(slide, placements, region):
    # Mixed figures/diagrams require their original spatial relationships. Never
    # reflow those pages or disturb equations composed of overlapping objects.
    if any(s._element.tag in (qn('p:pic'), qn('p:grpSp')) or s.has_table or s.has_chart for s in slide.shapes):
        return
    text = [s for s in slide.shapes if s.has_text_frame and s.text.strip()]
    x,y,w,h = [Inches(v) for v in region]
    ordered = sorted(text,key=lambda s:(s.top,s.left))
    for shape in ordered:
        runs=[r for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
        size=max((r.font.size.pt for r in runs),default=11)
        if size >= 11: continue
        factor=11/size
        # Overlay links are handled by the uniform content transform. Avoid
        # independent line growth that could separate link text and its hitbox.
        if any(s._element.xpath('.//a:hlinkClick') and s.left < shape.left+shape.width
               and s.left+s.width > shape.left and s.top < shape.top+shape.height
               and s.top+s.height > shape.top for s in slide.shapes if s != shape): continue
        width=int(shape.width*factor)
        height=max(int(shape.height*factor),int(Pt(11*1.3)))
        if shape.left+width > x+w: continue
        following=[s for s in ordered if s.top > shape.top]
        overlapping=[s for s in following if s.left < shape.left+width and s.left+s.width > shape.left]
        shift=max(0,shape.top+height+Pt(2)-min((s.top for s in overlapping),default=y+h))
        if max([shape.top+height]+[s.top+s.height+shift for s in following]) > y+h: continue
        # Only grow a stand-alone line, never an inline fragment/superscript.
        if any(s != shape and s.top <= shape.top < s.top+s.height
               and s.left < shape.left+width and s.left+s.width > shape.left for s in text): continue
        shape.width=width; shape.height=height
        from app.pdf_text_spacing import scale_tabs
        for paragraph in shape.text_frame.paragraphs: scale_tabs(paragraph,factor)
        for r in runs: r.font.size=Pt(r.font.size.pt*factor)
        for s in following: s.top+=int(shift)
        # Link hitboxes must follow the same native text, not remain at old y.
        for s in slide.shapes:
            if s in text or not s._element.xpath('.//a:hlinkClick'): continue
            if s.top >= shape.top+height: s.top+=int(shift)
    for placement in placements:
        shape=next((s for s in slide.shapes if s.name.split('|',1)[0]=='sss:'+placement['source_id']),None)
        if shape is not None: placement['bounds']=[shape.left,shape.top,shape.width,shape.height]
