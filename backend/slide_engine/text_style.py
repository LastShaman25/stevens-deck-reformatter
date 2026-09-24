"""Resolve supported inherited run properties before moving between templates."""
from pptx.oxml.ns import qn
from pptx.util import Pt
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
from copy import deepcopy
from pptx.oxml.xmlchemy import OxmlElement


def defaults(shape, level, placeholder=False):
    if not shape.has_text_frame:
        return []
    tf=shape.text_frame
    result=[]
    if placeholder and tf.paragraphs:
        ppr=tf.paragraphs[0]._p.find(qn('a:pPr'))
        if ppr is not None:
            result.append(ppr.find(qn('a:defRPr')))
    style=tf._txBody.find(qn('a:lstStyle'))
    if style is not None:
        for tag in (f'a:lvl{level+1}pPr','a:defPPr'):
            ppr=style.find(qn(tag))
            if ppr is not None:result.append(ppr.find(qn('a:defRPr')))
    return [x for x in result if x is not None]


def inherited(shape, level):
    result=defaults(shape,level)
    try:
        if not shape.is_placeholder:return result
        slide=shape.part.slide
        ph=shape.placeholder_format
        layout=slide.slide_layout
        for candidate in layout.placeholders:
            if candidate.placeholder_format.idx==ph.idx:
                result+=defaults(candidate,level,placeholder=True)
                break
        master=layout.slide_master
        for candidate in master.placeholders:
            if candidate.placeholder_format.type==ph.type:
                result+=defaults(candidate,level,placeholder=True)
                break
        role='titleStyle' if int(ph.type) in (1,3) else 'bodyStyle'
        for node in master._element.xpath(f'./p:txStyles/p:{role}/a:lvl{level+1}pPr/a:defRPr'):
            result.append(node)
    except (AttributeError,KeyError,ValueError):
        pass
    return result


def resolve(run, paragraph, shape, key, fallback=False):
    value=getattr(run.font,key)
    if value is not None:return value
    value=getattr(paragraph.font,key)
    if value is not None:return value
    attr={'bold':'b','italic':'i','underline':'u','size':'sz'}[key]
    for node in inherited(shape,paragraph.level):
        value=node.get(attr)
        if value is not None:
            return Pt(int(value)/100) if key=='size' else value not in ('0','false','none')
    return fallback


def materialize(original, copied):
    from .inventory import walk_shapes
    for (_,source),(_,target) in zip(walk_shapes([original]),walk_shapes([copied])):
        if not source.has_text_frame or not target.has_text_frame:continue
        for sp,dp in zip(source.text_frame.paragraphs,target.text_frame.paragraphs):
            for sr,dr in zip(sp.runs,dp.runs):
                for key in ('bold','italic','underline'):
                    setattr(dr.font,key,bool(resolve(sr,sp,source,key)))
                dr.font.size=resolve(sr,sp,source,'size',Pt(20))
    # Resolve theme tokens against the source before attaching its XML to a new theme.
    # Leave the standard OOXML color transforms intact so tint/shade/alpha are retained.
    part=original.part
    seen=set()
    theme=None
    while part not in seen:
        seen.add(part)
        rels=list(part.rels.values())
        rel=next((r for r in rels if r.reltype==RT.THEME),None)
        if rel:
            theme=etree.fromstring(rel.target_part.blob);break
        rel=next((r for r in rels if r.reltype in (RT.SLIDE_LAYOUT,RT.SLIDE_MASTER)),None)
        if not rel:break
        part=rel.target_part
    if theme is None:return
    aliases={'bg1':'lt1','tx1':'dk1','bg2':'lt2','tx2':'dk2'}
    colors={}
    scheme=theme.find('.//'+qn('a:clrScheme'))
    if scheme is not None:
        for color in scheme:
            if len(color):colors[etree.QName(color).localname]=color[0].get('val') if color[0].tag==qn('a:srgbClr') else color[0].get('lastClr')
    for node in list(copied._element.iter(qn('a:schemeClr'))):
        key=node.get('val');value=colors.get(aliases.get(key,key))
        if value:
            node.tag=qn('a:srgbClr');node.set('val',value)
    for (_,source),(_,target) in zip(walk_shapes([original]),walk_shapes([copied])):
        if not source.has_text_frame or not target.has_text_frame:continue
        ref=target._element.find('.//'+qn('a:fontRef'))
        inherited_color=None
        for node in inherited(source,0):
            fill=node.find(qn('a:solidFill'))
            if fill is not None and len(fill):inherited_color=deepcopy(fill[0]);break
        if inherited_color is None and ref is not None and len(ref):inherited_color=deepcopy(ref[0])
        if inherited_color is not None and inherited_color.tag==qn('a:schemeClr'):
            key=inherited_color.get('val');value=colors.get(aliases.get(key,key))
            if value:inherited_color.tag=qn('a:srgbClr');inherited_color.set('val',value)
        if inherited_color is None:continue
        for p in target.text_frame.paragraphs:
            for r in p.runs:
                if r.font.color.type is None:
                    r.font.fill.solid()
                    fill=r._r.get_or_add_rPr().find(qn('a:solidFill'))
                    for old in list(fill):fill.remove(old)
                    fill.append(deepcopy(inherited_color))
