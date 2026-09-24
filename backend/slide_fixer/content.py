"""Content extraction: turn a source slide into a structured content model that
the rebuild engine can pour into Stevens template layouts.
"""
import re

from pptx.oxml.ns import qn

from .engine import iter_shapes

_THANKS = re.compile(r"\bthank(s| you)\b", re.I)
_NUMERIC = re.compile(r"^[\d\W]{1,4}$")          # stray page numbers / bullets
_HAS_LETTER = re.compile(r"[A-Za-z]")


def _is_junk(text):
    """True for stray page numbers or symbol-only fragments."""
    return bool(_NUMERIC.match(text.strip())) or not _HAS_LETTER.search(text)


def _is_pic(sh):
    return sh._element.tag == qn("p:pic")


def _emu(v):
    try:
        return int(v)
    except Exception:
        return 0


def extract(slide, idx):
    """Return a dict content model for one source slide."""
    model = {
        "index": idx,
        "title": None,
        "body": [],          # list of (level, text)
        "images": [],        # list of dicts: blob, ext, left, top, width, height
        "n_tables": 0,
        "notes": "",
    }

    title_sh = slide.shapes.title
    title_el = title_sh._element if title_sh is not None else None
    if title_sh is not None and (title_sh.text or "").strip():
        model["title"] = title_sh.text.strip()

    shapes = sorted(
        iter_shapes(slide.shapes),
        key=lambda s: (_emu(s.top), _emu(s.left)),
    )
    for sh in shapes:
        if title_el is not None and sh._element is title_el:
            continue
        if _is_pic(sh):
            try:
                img = sh.image
                model["images"].append({
                    "blob": img.blob, "ext": img.ext,
                    "left": _emu(sh.left), "top": _emu(sh.top),
                    "width": _emu(sh.width), "height": _emu(sh.height),
                })
            except Exception:
                pass
            continue
        if sh._element.tag == qn("p:graphicFrame"):
            try:
                if sh.has_table:
                    model["n_tables"] += 1
            except Exception:
                pass
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            for p in sh.text_frame.paragraphs:
                t = (p.text or "").strip()
                if t and not _is_junk(t):
                    model["body"].append((min(p.level, 4), t))

    # If no title placeholder, promote a short leading body line to title
    # (must be a real phrase, not a number/fragment).
    if not model["title"] and model["body"]:
        lvl, txt = model["body"][0]
        if 1 <= len(txt.split()) <= 12 and _HAS_LETTER.search(txt):
            model["title"] = txt
            model["body"] = model["body"][1:]

    try:
        if slide.has_notes_slide:
            model["notes"] = slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        pass

    return model


def classify_kind(model, idx, total):
    """One of: title, section, thankyou, picture, content."""
    text_all = " ".join(t for _, t in model["body"]) + " " + (model["title"] or "")
    if _THANKS.search(text_all) and len(model["body"]) <= 2 and not model["images"]:
        return "thankyou"
    if idx == 0:
        return "title"
    has_body = len(model["body"]) > 0
    has_img = len(model["images"]) > 0
    if not has_body and not has_img and model["title"]:
        return "section"
    if has_img and not has_body:
        return "picture"
    return "content"
