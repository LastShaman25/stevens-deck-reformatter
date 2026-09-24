"""Stevens re-skin engine (P1).

Applies the locked Stevens rulebook to an arbitrary .pptx *in place* (Option A):
theme retint, Arial everywhere, black body text, role-based type scale,
dark-blue square bullets, red 1pt borders on standalone text boxes, and
shrink-to-fit. Flags dense slides as split candidates. Rendering / vision QA
live in separate modules and are not required to run this.
"""
import os
import re
import tempfile
import zipfile

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.enum.shapes import PP_PLACEHOLDER as PPP

from . import palette
from .rules import DEFAULT_RULES

TITLE_PH = {PPP.TITLE, PPP.CENTER_TITLE}
SUBTITLE_PH = {PPP.SUBTITLE}
_BULLET_TAGS = [
    "a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct", "a:buSzPts",
    "a:buFontTx", "a:buFont", "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip",
]


# --------------------------------------------------------------------------- #
# shape traversal + classification
# --------------------------------------------------------------------------- #
def iter_shapes(shapes):
    for sh in shapes:
        if sh._element.tag == qn("p:grpSp"):
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def slide_kind(slide, idx):
    """'title' for slide 1 / title-or-section layouts, else 'content'."""
    if idx == 0:
        return "title"
    name = (slide.slide_layout.name or "").lower().strip()
    # Only genuine title/section dividers get the larger header size.
    if name in ("title slide", "title", "section header") or "section" in name:
        return "title"
    return "content"


def classify(sh):
    if sh.is_placeholder:
        try:
            t = sh.placeholder_format.type
        except Exception:
            t = None
        if t in TITLE_PH:
            return "title"
        if t in SUBTITLE_PH:
            return "subtitle"
        return "body"
    if sh.has_text_frame and sh.text_frame.text.strip():
        return "textbox"
    return "shape"


# --------------------------------------------------------------------------- #
# bullets
# --------------------------------------------------------------------------- #
def _clear_bullets(pPr):
    for tag in _BULLET_TAGS:
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)


def _insert_bullet_group(pPr, elems):
    """Insert bullet elements in schema order (before tabLst/defRPr/extLst)."""
    anchor = None
    for tag in ("a:tabLst", "a:defRPr", "a:extLst"):
        found = pPr.find(qn(tag))
        if found is not None:
            anchor = found
            break
    for el in elems:
        if anchor is not None:
            anchor.addprevious(el)
        else:
            pPr.append(el)


def _bare(tag):
    from lxml import etree
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    return etree.Element(qn(tag), nsmap={"a": a})


def set_square_bullet(p, cfg):
    pPr = p._p.get_or_add_pPr()
    _clear_bullets(pPr)
    pPr.set("marL", "285750")
    pPr.set("indent", "-285750")
    buClr = _bare("a:buClr")
    srgb = _bare("a:srgbClr")
    srgb.set("val", cfg.bullet_color)
    buClr.append(srgb)
    buFont = _bare("a:buFont")
    buFont.set("typeface", cfg.bullet_font)
    buChar = _bare("a:buChar")
    buChar.set("char", cfg.bullet_char)
    _insert_bullet_group(pPr, [buClr, buFont, buChar])


def set_no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    _clear_bullets(pPr)
    _insert_bullet_group(pPr, [_bare("a:buNone")])


# --------------------------------------------------------------------------- #
# text styling
# --------------------------------------------------------------------------- #
def _set_run(run, cfg, size_pt, color_hex):
    if cfg.normalize_fonts:
        run.font.name = cfg.body_font
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color_hex is not None:
        run.font.color.rgb = RGBColor.from_string(color_hex)


def style_text_frame(sh, role, slide_kind_, cfg, stats):
    tf = sh.text_frame

    if role == "title":
        size = cfg.title_section_pt if slide_kind_ == "title" else cfg.title_content_pt
    elif role == "subtitle":
        size = cfg.subtitle_pt
    else:
        size = cfg.body_pt

    if cfg.shrink_to_fit and role in ("body", "textbox"):
        try:
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        except Exception:
            pass

    for p in tf.paragraphs:
        runs = list(p.runs)
        has_text = any(r.text.strip() for r in runs)
        pPr = p._pPr
        existing_bullet = pPr is not None and (
            pPr.find(qn("a:buChar")) is not None or pPr.find(qn("a:buAutoNum")) is not None
        )
        # bullets
        if role == "body" and has_text:
            # body placeholders are inherently list content
            if cfg.square_bullets:
                set_square_bullet(p, cfg)
            stats["bullets"] += 1
        elif role == "textbox" and has_text and existing_bullet:
            # standalone text boxes: only convert paragraphs that already bullet
            if cfg.square_bullets:
                set_square_bullet(p, cfg)
            stats["bullets"] += 1
        elif role in ("title", "subtitle"):
            if cfg.square_bullets:
                set_no_bullet(p)

        for r in runs:
            if not r.text:
                continue
            color = cfg.text_color if cfg.force_black_text else None
            if cfg.auto_emphasis and role == "body" and r.font.bold:
                color = cfg.emphasis_color
            _set_run(r, cfg, size, color)
            if role in ("body", "textbox"):
                stats["words"] += len(r.text.split())


# --------------------------------------------------------------------------- #
# shape fills / borders
# --------------------------------------------------------------------------- #
def snap_shape_fill(sh, cfg):
    """Snap explicit shape fill/line srgb colors to the nearest Stevens color."""
    spPr = sh._element.find(qn("p:spPr"))
    if spPr is None:
        return
    for el in spPr.iterfind(".//" + qn("a:srgbClr")):
        val = el.get("val")
        if not val:
            continue
        snapped = palette.nearest_allowed(val)
        if snapped and snapped.upper() != val.upper():
            el.set("val", snapped)


def add_red_border(sh, cfg):
    try:
        ln = sh.line
        ln.color.rgb = RGBColor.from_string(cfg.box_border_color)
        ln.width = Pt(cfg.box_border_pt)
    except Exception:
        pass


# --------------------------------------------------------------------------- #
# theme retint (post-process the saved zip)
# --------------------------------------------------------------------------- #
def _retint_theme_xml(xml, cfg):
    for slot, hexv in cfg.theme_scheme.items():
        xml = re.sub(
            r"(<a:%s>).*?(</a:%s>)" % (slot, slot),
            r'\g<1><a:srgbClr val="%s"/>\g<2>' % hexv,
            xml, count=1, flags=re.S,
        )
    for block in ("majorFont", "minorFont"):
        xml = re.sub(
            r'(<a:%s>\s*<a:latin[^>]*typeface=")[^"]*(")' % block,
            r"\g<1>%s\g<2>" % cfg.body_font,
            xml, count=1, flags=re.S,
        )
    return xml


def retint_theme(src_pptx, dst_pptx, cfg):
    with zipfile.ZipFile(src_pptx, "r") as zin:
        items = zin.infolist()
        data = {i.filename: zin.read(i.filename) for i in items}
    for name in list(data):
        if re.match(r"ppt/theme/theme\d+\.xml$", name):
            data[name] = _retint_theme_xml(
                data[name].decode("utf-8", "ignore"), cfg).encode("utf-8")
    with zipfile.ZipFile(dst_pptx, "w", zipfile.ZIP_DEFLATED) as zout:
        for i in items:
            zout.writestr(i, data[i.filename])


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def reformat(src_path, out_path, cfg=DEFAULT_RULES):
    prs = Presentation(src_path)
    report = {"source": src_path, "output": out_path, "slides": []}

    for idx, slide in enumerate(prs.slides):
        kind = slide_kind(slide, idx)
        stats = {"words": 0, "bullets": 0}
        for sh in iter_shapes(slide.shapes):
            role = classify(sh)
            if cfg.snap_colors:
                snap_shape_fill(sh, cfg)
            if sh.has_text_frame and sh.text_frame.text.strip():
                style_text_frame(sh, role, kind, cfg, stats)
                if role == "textbox" and cfg.add_box_borders:
                    add_red_border(sh, cfg)

        dense = stats["words"] > cfg.split_word_cap or stats["bullets"] > cfg.split_bullet_cap
        report["slides"].append({
            "index": idx + 1, "kind": kind,
            "words": stats["words"], "bullets": stats["bullets"],
            "split_candidate": dense,
        })

    if cfg.retint_theme:
        tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False).name
        prs.save(tmp)
        retint_theme(tmp, out_path, cfg)
        os.unlink(tmp)
    else:
        prs.save(out_path)

    return report
