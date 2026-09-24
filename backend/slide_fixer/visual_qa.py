"""Visual QA loop: render -> inspect (vision model) -> fix -> re-check.

Provider-agnostic. A `VisionProvider` receives a rendered slide image plus a
*shape manifest* (each shape's id + bounding box + text + role) and returns a
list of structured issues that reference shape ids, so every fix maps to a
deterministic python-pptx edit instead of a pixel guess.

No network provider is wired yet (no API key). `NullProvider` returns no issues
so the loop is testable; add an OpenAI/Claude/Gemini provider when a key exists.
"""
from abc import ABC, abstractmethod

from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

from . import palette
from .engine import iter_shapes, classify, snap_shape_fill
from .rules import DEFAULT_RULES

RUBRIC = """You are a Stevens Institute brand QA reviewer. Inspect the slide image.
Using the provided shape manifest (ids + bounding boxes + text), report problems:
- overlap: two shapes visibly overlap or text collides
- overflow: text spills outside its box or off the slide
- offpalette: a color that is not a Stevens-allowed color
- contrast: low-contrast text (e.g. light text on light fill)
- misalign: uneven margins / broken alignment
Return JSON: [{"type","shape_ids":[...],"severity":"low|med|high","suggestion":"..."}].
Only report real, visible problems. Reference shapes by their manifest id.
"""


def build_manifest(slide):
    """List of {id, role, left, top, width, height, text} in EMU."""
    items = []
    for sh in iter_shapes(slide.shapes):
        try:
            box = dict(left=int(sh.left or 0), top=int(sh.top or 0),
                       width=int(sh.width or 0), height=int(sh.height or 0))
        except Exception:
            box = dict(left=0, top=0, width=0, height=0)
        text = sh.text_frame.text if sh.has_text_frame else ""
        items.append({
            "id": sh.shape_id,
            "role": classify(sh),
            **box,
            "text": (text or "")[:200],
        })
    return items


class VisionProvider(ABC):
    @abstractmethod
    def inspect(self, image_path, manifest, rubric=RUBRIC):
        """Return list of issue dicts referencing manifest shape ids."""
        raise NotImplementedError


class NullProvider(VisionProvider):
    """No-op provider so the loop runs without an API key."""
    def inspect(self, image_path, manifest, rubric=RUBRIC):
        return []


# --------------------------------------------------------------------------- #
# deterministic fixes (issue -> python-pptx edit)
# --------------------------------------------------------------------------- #
def _shape_by_id(slide, sid):
    for sh in iter_shapes(slide.shapes):
        if sh.shape_id == sid:
            return sh
    return None


def _shrink_font(sh, step=2, floor=11):
    if not sh.has_text_frame:
        return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                new = max(floor, int(r.font.size.pt) - step)
                r.font.size = Pt(new)


def _recolor_black(sh):
    if not sh.has_text_frame:
        return
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor.from_string(palette.BLACK)


def apply_fix(slide, issue):
    """Apply a single issue's fix. Returns True if something changed."""
    changed = False
    for sid in issue.get("shape_ids", []):
        sh = _shape_by_id(slide, sid)
        if sh is None:
            continue
        t = issue.get("type")
        if t in ("overflow", "overlap"):
            _shrink_font(sh)
            changed = True
        elif t == "contrast":
            _recolor_black(sh)
            changed = True
        elif t == "offpalette":
            snap_shape_fill(sh, DEFAULT_RULES)
            changed = True
    return changed


def run_qa(prs, provider=None, render_fn=None, max_passes=3, out_dir=None):
    """Render -> inspect -> fix loop over all slides. Requires a render_fn that
    saves the current prs and returns a list of slide PNG paths."""
    provider = provider or NullProvider()
    history = []
    for _ in range(max_passes):
        pngs = render_fn(prs, out_dir) if render_fn else []
        any_change = False
        for i, slide in enumerate(prs.slides):
            img = pngs[i] if i < len(pngs) else None
            if img is None:
                continue
            manifest = build_manifest(slide)
            issues = provider.inspect(img, manifest)
            for issue in issues:
                any_change |= apply_fix(slide, issue)
            history.append({"slide": i + 1, "issues": issues})
        if not any_change:
            break
    return history
