"""Offline adversarial tests: provider failures must never authorize release."""
import json
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
import pytest
import requests
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from fastapi.testclient import TestClient
from app import grounded, generations, sessions
from app.main import app
from app.ai import providers, layout, pipeline
from app.qa.artifact_coverage import audit
from test_regressions import fixture, text
from test_artifacts import rich_fixture
from rubric_fixtures import passed_checks, role_map, source_choice, repair_evidence


@pytest.fixture
def config(tmp_path, monkeypatch):
    env=tmp_path/'ai.env'
    env.write_text('ANTHROPIC_API_KEY=test-anthropic\nGEMINI_API_KEY=test-gemini\n')
    monkeypatch.setattr(providers,'ENV_FILE',env)
    for name in ['OPENAI_API_KEY','OPENAI_MODEL','OPENAI_REASONING_EFFORT','OPENAI_REVIEW_REASONING_EFFORT','ANTHROPIC_API_KEY','GEMINI_API_KEY','STEVENS_AI_PLANNER','STEVENS_AI_REVIEWER',
                 'STEVENS_AI_MAX_CALLS','STEVENS_OFFLINE','ANTHROPIC_MODEL','GEMINI_MODEL',
                 'STEVENS_AI_PROVIDER','STEVENS_AI_MODEL','STEVENS_AI_PLANNER_MODEL','STEVENS_AI_REVIEWER_MODEL',
                 'AI_GATEWAY_API_KEY','AI_GATEWAY_MODEL','AI_GATEWAY_REASONING_EFFORT','AI_GATEWAY_REVIEW_REASONING_EFFORT']:
        monkeypatch.delenv(name,raising=False)
    # Legacy provider fixtures explicitly request legacy auto selection.
    monkeypatch.setenv('STEVENS_AI_PLANNER','auto')
    monkeypatch.setenv('STEVENS_AI_REVIEWER','auto')
    return env


def identity_plan(objects):
    return {'layout':'Aligned native objects','rationale':'Preserve source content',
            'objects':[{'id':o['id'],**dict(zip(('x','y','w','h'),o['box']))} for o in objects]}


def mocked_provider(role,system,payload,images=(),max_tokens=0):
    assert images and all(Path(path).is_file() for _,path in images)
    value=role_map(payload['objects']) if role=='element_roles' else identity_plan(payload['objects']) if role=='planner' else {'verdict':'passed','summary':'Synthetic review','findings':[],'rubric':passed_checks()}
    if payload.get('stage')=='source_decisions':value=source_choice(payload)
    return {'status':'completed','provider':'mock','model':'offline-test','data':value}


def fake_render(candidate,directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    pages=[]
    for i in range(len(Presentation(candidate).slides)):
        path=directory/f'slide-{i}.png';Image.new('RGB',(160,90),'white').save(path)
        pages.append({'output_slide':i,'png':str(path)})
    return {'status':'passed','findings':[],'pages':pages,'candidate_sha256':generations.sha256(candidate)}


@pytest.fixture
def ai_session(tmp_path,monkeypatch,config):
    s=sessions.create();Path(s.source_path).write_bytes(fixture(tmp_path).read_bytes())
    s.analysis=grounded.analyze(s.source_path)
    for i in range(s.analysis['slide_count']):Image.new('RGB',(160,90),'white').save(s.preview_path(i,'before'))
    monkeypatch.setattr(generations.render_verify,'check',fake_render)
    monkeypatch.setattr(providers,'generate',mocked_provider)
    yield s
    sessions.delete(s.id)


def test_configuration_refresh_and_single_provider(config,monkeypatch):
    assert providers.capabilities()['independent_providers']
    config.write_text('GEMINI_API_KEY=test-gemini\n')
    value=providers.capabilities()
    assert value['configured'] and not value['independent_providers']
    assert value['planner']['provider']=='gemini'
    config.write_text('ANTHROPIC_API_KEY=test-anthropic\n')
    assert providers.capabilities()['reviewer']['provider']=='anthropic'
    monkeypatch.setenv('STEVENS_OFFLINE','1')
    assert not providers.capabilities()['configured']


@pytest.mark.parametrize('code,status',[(401,'authentication_error'),(403,'permission_error'),(404,'model_unavailable'),(429,'rate_limited'),(500,'provider_error')])
def test_provider_http_errors_redact_secrets(config,monkeypatch,code,status):
    monkeypatch.setattr(providers.requests,'post',lambda *a,**k:SimpleNamespace(status_code=code,text='SECRET'))
    result=providers.generate('planner','Test',{})
    assert result['status']==status
    assert 'SECRET' not in json.dumps(result) and 'test-anthropic' not in json.dumps(result)


@pytest.mark.parametrize('provider', ['anthropic','gemini'])
def test_real_request_contract_and_truncation(config,monkeypatch,tmp_path,provider):
    monkeypatch.setenv('STEVENS_AI_PLANNER',provider)
    path=tmp_path/'image.png';Image.new('RGB',(10,10)).save(path)
    returned=({'content':[{'type':'text','text':'{"ok":true}'}],'stop_reason':'end_turn','usage':{'input_tokens':12}}
              if provider=='anthropic' else {'candidates':[{'finishReason':'STOP','content':{'parts':[{'text':'{"ok":true}'}]}}]})
    captured={}
    def post(url,**kw):
        captured.update(url=url,**kw)
        return SimpleNamespace(status_code=200,json=lambda:returned)
    monkeypatch.setattr(providers.requests,'post',post)
    assert providers.generate('planner','System',{'test':'Full content'},[('Source',path)])['data']=={'ok':True}
    assert 'key=' not in captured['url']
    assert captured['timeout']==(15,180)
    assert ('x-api-key' if provider=='anthropic' else 'x-goog-api-key') in captured['headers']
    assert 'image/png' in json.dumps(captured['json'])
    if provider=='anthropic':returned['stop_reason']='max_tokens'
    else:returned['candidates'][0]['finishReason']='MAX_TOKENS'
    assert providers.generate('planner','System',{})['status']=='truncated'


@pytest.mark.parametrize('fault,status',[(requests.Timeout(),'timeout'),(requests.ConnectionError(),'provider_error'),(ValueError(),'invalid_response')])
def test_provider_transport_failures(config,monkeypatch,fault,status):
    def fail(*a,**kw):raise fault
    monkeypatch.setattr(providers.requests,'post',fail)
    assert providers.generate('reviewer','System',{})['status']==status


@pytest.mark.parametrize('damage',['missing','duplicate','invented','offslide','nan','text'])
def test_invalid_plans_rejected(tmp_path,damage):
    source=fixture(tmp_path);candidate=tmp_path/'candidate.pptx';grounded.build_deck(source,candidate)
    p=Presentation(candidate);slide=p.slides[0];value=identity_plan(layout.describe(slide))
    if damage=='missing':value['objects'].pop()
    if damage=='duplicate':value['objects'].append(deepcopy(value['objects'][0]))
    if damage=='invented':value['objects'][0]['id']='fake'
    if damage=='offslide':value['objects'][0]['x']=99
    if damage=='nan':value['objects'][0]['x']=float('nan')
    if damage=='text':value['objects'][0]['text']='Replacement content'
    with pytest.raises(ValueError):layout.validate(layout.LayoutPlan.model_validate(value),slide,p.slide_width/layout.EMU,p.slide_height/layout.EMU)


def test_native_geometry_edit_preserves_relationships(tmp_path):
    source=rich_fixture(tmp_path);candidate=tmp_path/'candidate.pptx';out=tmp_path/'changed.pptx'
    report=grounded.build_deck(source,candidate);p=Presentation(candidate);plans={}
    for i,slide in enumerate(p.slides):
        value=identity_plan(layout.describe(slide))
        for o in value['objects']:
            if i > 0: o['x']+=.003
        plans[i]=layout.LayoutPlan.model_validate(value)
    updated,changed=layout.apply(candidate,out,plans,report)
    assert changed>0 and updated['llm_used']
    assert audit(source,out,updated)['status']=='passed'


def test_overflow_feedback_identifies_the_editable_object(tmp_path):
    from pptx.util import Pt
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(.3))
    shape.name='sss:synthetic/0/slide/2|title'
    shape.text='A title needs enough vertical space'
    for run in shape.text_frame.paragraphs[0].runs:
        run.font.name='Arial';run.font.size=Pt(40)
    candidate=tmp_path/'tight-title.pptx';prs.save(candidate)
    description=layout.describe(slide)[0]
    assert description['text_fit_ratio']>1
    assert description['text_margins']['top']>0
    findings=pipeline.structural(candidate)['findings']
    overflow=[f for f in findings if f['code']=='OVERFLOW']
    assert overflow and overflow[0]['object_ids']==[shape.name]
    shape.height=Inches(2)
    prs.save(candidate)
    assert layout.describe(slide)[0]['text_fit_ratio']<.85
    assert not any(f['code']=='OVERFLOW' for f in pipeline.structural(candidate)['findings'])


def test_group_absolute_edit_keeps_children_and_semantics(tmp_path):
    source=fixture(tmp_path);p=Presentation(source);g=p.slides[0].shapes.add_group_shape()
    sh=g.shapes.add_textbox(Inches(1),Inches(4),Inches(2),Inches(.5));sh.text='Grouped label'
    p.save(source);candidate=tmp_path/'candidate.pptx';out=tmp_path/'changed.pptx'
    report=grounded.build_deck(source,candidate);p=Presentation(candidate);slide=p.slides[0]
    before=layout.describe(slide);value=identity_plan(before)
    group=next(o for o in before if o['kind']=='group')
    for o in value['objects']:
        if o['id']==group['id'] or next(n for n in before if n['id']==o['id'])['parent']==group['id']:o['x']+=.2
    updated,_=layout.apply(candidate,out,{0:layout.LayoutPlan.model_validate(value)},report)
    after={o['id']:o for o in layout.describe(Presentation(out).slides[0])}
    child=next(o for o in before if o['parent']==group['id'])
    assert after[child['id']]['box'][0]==pytest.approx(child['box'][0]+.2,abs=.00001)
    assert audit(source,out,updated)['status']=='passed'


def test_ai_success_requires_all_checks_and_exact_download(ai_session):
    r=generations.build(ai_session,mode='ai',repair_passes=0)
    assert r['checks']['ai_redesign']['status']=='passed',r['findings']
    assert r['checks']['ai_visual_review']['status']=='passed'
    assert r['checks']['artifact_coverage']['status']=='passed'
    assert len(r['ai_pipeline']['calls'])==6
    assert [c['role'] for c in r['ai_pipeline']['calls'][3:]]==['reviewer']*3
    assert not r['ai_pipeline']['element_roles']  # No invented edits after a clean native review.
    assert all(v['candidate_sha256']==r['candidate_sha256'] for v in r['ai_pipeline']['final_reviews'])
    ids=[f['id'] for f in r['findings'] if f['severity']=='review']
    if ids:generations.decide(ai_session,generations.Decision(generation_id=r['generation_id'],candidate_sha256=r['candidate_sha256'],finding_ids=ids,rationale='Synthetic fixture inspection'))
    assert r['state']=='ready'
    response=TestClient(app).get(f'/api/sessions/{ai_session.id}/download',params={'generation_id':r['generation_id']})
    import hashlib
    assert hashlib.sha256(response.content).hexdigest()==r['candidate_sha256']


@pytest.mark.parametrize('fault',['missing_id','timeout','review_invalid','review_timeout','missing_key','budget'])
def test_ai_incomplete_never_releases(ai_session,monkeypatch,config,fault):
    def provider(role,system,payload,images=(),max_tokens=0):
        result=mocked_provider(role,system,payload,images,max_tokens)
        if fault=='missing_id' and role=='planner':result['data']['objects'].pop()
        if fault=='missing_id' and role=='reviewer':
            fs=[{**repair_evidence(),'criterion':'spatial_layout','severity':'blocking','message':'Repair required.'}]
            result['data'].update(verdict='failed',findings=fs,rubric=passed_checks(fs))
        if fault=='timeout' or (fault=='review_timeout' and role=='reviewer'):return {'status':'timeout','message':'Injected timeout'}
        if fault=='review_invalid' and role=='reviewer':result['data']['verdict']='failed'
        return result
    monkeypatch.setattr(providers,'generate',provider)
    if fault=='missing_key':config.write_text('')
    if fault=='budget':monkeypatch.setenv('STEVENS_AI_MAX_CALLS','1')
    r=generations.build(ai_session,mode='ai',repair_passes=1)
    assert r['state'] in ('error','failed'),r
    assert any(f['severity']=='blocking' for f in r['findings'])
    assert TestClient(app).get(f'/api/sessions/{ai_session.id}/download',params={'generation_id':r['generation_id']}).status_code==409
    assert TestClient(app).get(f'/api/sessions/{ai_session.id}/download',params={'generation_id':r['generation_id'],'draft':True}).status_code==409


def test_repair_rollback_keeps_reviewed_bytes(ai_session,monkeypatch):
    review_calls=0
    def provider(role,system,payload,images=(),max_tokens=0):
        nonlocal review_calls
        result=mocked_provider(role,system,payload,images,max_tokens)
        if role=='planner':
            for obj in result['data']['objects']:
                if payload['slide'] > 0: obj['x']+=.003
        elif role=='reviewer':
            review_calls+=1
            count=1 if review_calls<=3 else 2
            findings=[{**repair_evidence(),'criterion':'spatial_layout','severity':'review','object_ids':[],
                'message':f'Spacing needs inspection in region {i}'} for i in range(count)]
            result['data']={'verdict':'needs_review','summary':'Requires spacing review','rubric':passed_checks(findings),'findings':findings}
        return result
    monkeypatch.setattr(providers,'generate',provider)
    r=generations.build(ai_session,mode='ai',repair_passes=1)
    attempts=r['ai_pipeline']['attempts']
    assert [a['accepted'] for a in attempts]==[False],attempts
    assert r['candidate_sha256']!=attempts[0]['candidate_sha256']
    assert all(v['candidate_sha256']==r['candidate_sha256'] for v in r['ai_pipeline']['final_reviews'])
    assert r['state']=='needs_review'


def test_ai_cannot_delete_content_even_if_model_reviewer_passes(ai_session,monkeypatch):
    calls=0
    def provider(role,system,payload,images=(),max_tokens=0):
        nonlocal calls
        result=mocked_provider(role,system,payload,images,max_tokens)
        if role=='reviewer':
            calls+=1
            if calls<=3:
                fs=[{**repair_evidence(),'criterion':'spatial_layout','severity':'review','message':'Spacing repair required.'}]
                result['data'].update(verdict='needs_review',findings=fs,rubric=passed_checks(fs))
        return result
    monkeypatch.setattr(providers,'generate',provider)
    original=layout.apply
    def corrupt(candidate,out,plans,report):
        report,changed=original(candidate,out,plans,report)
        p=Presentation(out);p.slides[0].shapes[1].text='Deleted original content';p.save(out)
        return report,changed
    monkeypatch.setattr(layout,'apply',corrupt)
    r=generations.build(ai_session,mode='ai')
    assert r['state']=='error'
    assert len([c for c in r['ai_pipeline']['calls'] if c['role']=='reviewer'])==3  # Baseline only; corrupt proposal never reaches QA.
    assert audit(ai_session.source_path,r['candidate'],r['report'])['status']=='passed'
    assert r['checks']['ai_visual_review']['status']=='error'


def test_api_mode_validation_and_diagnostics(ai_session,monkeypatch):
    client=TestClient(app)
    assert generations.GenerateRequest().mode=='ai'
    assert client.post(f'/api/sessions/{ai_session.id}/generate',json={'mode':'preserve'}).status_code==422
    assert client.post(f'/api/sessions/{ai_session.id}/generate',json={'mode':'unknown'}).status_code==422
    assert client.post(f'/api/sessions/{ai_session.id}/generate',json={'mode':'ai','repair_passes':10}).status_code==422
    monkeypatch.setattr(providers,'generate',lambda *a,**kw:{'status':'completed','data':{'ok':True},'model':'mock'})
    value=client.post('/api/ai/test',json={}).json()
    assert all(v['status']=='completed' for v in value['results'].values())
    assert 'test-anthropic' not in json.dumps(value)


@pytest.mark.renderer
def test_ai_native_edits_with_real_powerpoint_render(tmp_path,monkeypatch,config):
    from app import rendering
    s=sessions.create()
    try:
        Path(s.source_path).write_bytes(fixture(tmp_path).read_bytes())
        s.analysis=grounded.analyze(s.source_path)
        rendering.render_to_pdf(s.source_path,s.source_pdf)
        rendering.rasterize_pdf(s.source_pdf,lambda i:s.preview_path(i,'before'))
        review_calls=0
        def provider(role,system,payload,images=(),max_tokens=0):
            nonlocal review_calls
            result=mocked_provider(role,system,payload,images,max_tokens)
            if role=='reviewer':
                review_calls+=1
                if review_calls<=3:
                    f={**repair_evidence(),'criterion':'spatial_layout','severity':'review','message':'Synthetic move required.'}
                    result['data'].update(verdict='needs_review',findings=[f],rubric=passed_checks([f]))
            if role=='planner':
                for obj in result['data']['objects']:
                    if payload['slide']>0: obj['x']+=.05
            return result
        monkeypatch.setattr(providers,'generate',provider)
        r=generations.build(s,mode='ai',repair_passes=0)
        assert r['checks']['ai_redesign']['status']=='passed',r['findings']
        assert r['checks']['artifact_coverage']['status']=='passed'
        assert r['checks']['render_verification']['status']=='passed',r['findings']
        assert r['ai_pipeline']['changed_objects']>0
        assert len(r['checks']['render_verification']['pages'])==3
    finally:sessions.delete(s.id)


def test_rejected_slide_does_not_discard_other_slide_repairs(ai_session,monkeypatch):
    baseline={}
    def provider(role,system,payload,images=(),max_tokens=0):
        result=mocked_provider(role,system,payload,images,max_tokens)
        if role=='planner':
            if payload['slide']==1: result['data']['objects'][0]['x']=99
            elif payload['slide']==2:
                for obj in result['data']['objects']:
                    obj['x']+=.03;obj['w']-=.03 # Stay inside the closing template's right edge.
        elif role=='reviewer':
            ident=payload['required_objects'][0]['id']
            index=int(ident.split('/')[1])
            x=payload['required_objects'][0]['box'][0]
            baseline.setdefault(index,x)
            if index==1 or (index==2 and abs(x-baseline[index])<.01):
                fs=[{**repair_evidence(),'criterion':'spatial_layout','severity':'blocking','message':'Required spacing repair.'}]
                result['data'].update(verdict='failed',findings=fs,rubric=passed_checks(fs))
        return result
    monkeypatch.setattr(providers,'generate',provider)
    r=generations.build(ai_session,mode='ai',repair_passes=0)
    assert r['ai_pipeline']['changed_objects']>0
    assert r['ai_pipeline']['rejected_slide_plans']
    assert r['checks']['ai_visual_review']['status']=='failed'
    assert {f['output_slide'] for f in r['checks']['ai_visual_review']['findings']}=={1}
    assert not generations.download_allowed(r)
    assert all(v['candidate_sha256']==r['candidate_sha256'] for v in r['ai_pipeline']['final_reviews'])


def test_requested_picture_rectangle_fits_without_restyling_pdf_heading(tmp_path):
    from pptx.util import Pt
    from slide_engine import template_policy as T
    prs=Presentation();prs.slide_width,prs.slide_height=T.CANVAS
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    heading=slide.shapes.add_textbox(Inches(1),Inches(.5),Inches(8),Inches(1))
    heading.name='pdf-heading';heading.text='Existing readable heading'
    heading.text_frame.paragraphs[0].runs[0].font.size=Pt(28)
    path=tmp_path/'figure.png';Image.new('RGB',(200,100),'blue').save(path)
    picture=slide.shapes.add_picture(str(path),Inches(2),Inches(2),Inches(4),Inches(2));picture.name='figure'
    proposal=identity_plan(layout.describe(slide))
    proposal['objects'][0].update(role='title',font_size=28)
    proposal['objects'][1].update(x=11.8,y=5.8,w=4,h=1)
    safe=layout.constrain(layout.LayoutPlan.model_validate(proposal),slide,[{'object_ids':['figure']}])
    layout.validate(safe,slide,prs.slide_width/layout.EMU,prs.slide_height/layout.EMU)
    assert safe.objects[0].role=='keep' and safe.objects[0].font_size is None
    assert safe.objects[1].w/safe.objects[1].h==pytest.approx(2)
    assert T.contains([safe.objects[1].x,safe.objects[1].y,safe.objects[1].w,safe.objects[1].h],T.CONTENT)


def test_scope_includes_paired_style_and_displaced_following_text():
    from rubric_fixtures import role_map
    from app.ai.element_roles import RoleMap
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    for name,y in [('marker',1),('equation',1.5),('note',3),('examples',3.5),('row',4)]:
        sh=slide.shapes.add_textbox(Inches(1),Inches(y),Inches(4),Inches(.35));sh.name=name;sh.text=name
    objects=layout.describe(slide);roles=role_map(objects)
    roles['elements'][1]['related_ids']=['marker']
    plan=identity_plan(objects)
    plan['objects'][0]['color']=plan['objects'][1]['color']='red'
    plan['objects'][2]['y']=3.5
    plan['objects'][3].update(y=4,font_size=30)
    plan['objects'][4]['y']=4.5
    safe=layout.constrain(layout.LayoutPlan.model_validate(plan),slide,
                          [{'object_ids':['equation','note']}],RoleMap.model_validate(roles))
    edits={e.id:e for e in safe.objects}
    assert edits['marker'].color==edits['equation'].color=='red'
    assert edits['examples'].y==4 and edits['row'].y==4.5
    assert edits['examples'].font_size is None  # Neighbor reflow does not authorize restyling.


def test_passing_slides_not_replanned_for_old_advisories(ai_session,monkeypatch):
    planned=[]
    def provider(role,system,payload,*args,**kwargs):
        result=mocked_provider(role,system,payload,*args,**kwargs)
        if role=='reviewer' and not planned and '/1/' in payload['required_objects'][0]['id']:
            f={**repair_evidence(),'criterion':'spatial_layout','severity':'review','message':'Move body right.'}
            result['data'].update(verdict='needs_review',findings=[f],rubric=passed_checks([f]))
        if role=='planner':
            planned.append(payload['slide'])
            result['data']['objects'][0]['x']+=.03
        return result
    monkeypatch.setattr(providers,'generate',provider)
    r=generations.build(ai_session,mode='ai',repair_passes=2)
    assert planned==[1]
    assert r['checks']['ai_visual_review']['status']=='passed'


@pytest.mark.parametrize('changed_render',[False,True])
def test_review_reuse_requires_identical_render_and_native_evidence(ai_session,monkeypatch,changed_render):
    reviewed=[];planned=[]
    def provider(role,system,payload,*args,**kwargs):
        result=mocked_provider(role,system,payload,*args,**kwargs)
        if role=='reviewer':
            index=int(payload['required_objects'][0]['id'].split('/')[1]);reviewed.append(index)
            if index==1 and not planned:
                f={**repair_evidence(),'criterion':'spatial_layout','severity':'review','message':'Move body right.'}
                result['data'].update(verdict='needs_review',findings=[f],rubric=passed_checks([f]))
        if role=='planner':
            planned.append(payload['slide']);result['data']['objects'][0]['x']+=.03
        return result
    def render(candidate,directory):
        result=fake_render(candidate,directory)
        if changed_render and Path(candidate).name.startswith('ai-attempt-'):
            Image.new('RGB',(160,90),'gray').save(Path(directory)/'slide-2.png')
        return result
    monkeypatch.setattr(providers,'generate',provider)
    monkeypatch.setattr(generations.render_verify,'check',render)
    r=generations.build(ai_session,mode='ai',repair_passes=0)
    assert reviewed==([0,1,2,1,2] if changed_render else [0,1,2,1])
    receipts=r['ai_pipeline']['final_reviews']
    assert len(receipts)==3 and all(v['candidate_sha256']==r['candidate_sha256'] for v in receipts)
    assert all(v.get('evidence_sha256') for v in receipts)
    assert 'reused_from_candidate_sha256' in receipts[0]
    assert 'reused_from_candidate_sha256' not in receipts[1]
    assert ('reused_from_candidate_sha256' not in receipts[2])==changed_render
