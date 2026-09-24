from copy import deepcopy
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pptx.oxml.ns import qn
from app import grounded
from app.qa.artifact_coverage import audit
from slide_engine.inventory import inspect
from slide_engine.preserve import CoverageError
from test_regressions import text, fixture


def rich_fixture(tmp_path):
    path = fixture(tmp_path)
    p = Presentation(path)
    s = p.slides[1]
    s.notes_slide.notes_text_frame.text = 'Keep these speaker notes: 42 is not 24.'
    link = text(s, 'External reference', 5)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://example.edu/reference'
    internal = text(s, 'Go to closing', 6)
    internal.click_action.target_slide = p.slides[2]
    image = tmp_path / 'picture.png'
    Image.new('RGB', (80,40), 'orange').save(image)
    s.shapes.add_picture(str(image), Inches(8), Inches(1), width=Inches(1))
    data = CategoryChartData(); data.categories = ['A','B']; data.add_series('Scores',[12,34])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(8), Inches(3), Inches(1.5), Inches(1.5), data)
    p.save(path)
    return path


@pytest.mark.parametrize('mutation', ['text','table','image','chart','link','notes','duplicate','hidden'])
def test_independent_mutations(tmp_path, mutation):
    source = rich_fixture(tmp_path)
    candidate = tmp_path / 'out.pptx'
    report = grounded.build_deck(source,candidate)
    assert audit(source,candidate,report)['status'] == 'passed'
    p = Presentation(candidate)
    if mutation == 'text':
        p.slides[0].shapes[1].text_frame.paragraphs[1].runs[0].text = 'Altered paragraph'
    elif mutation == 'table':
        next(s for s in p.slides[1].shapes if s.has_table).table.cell(0,0).text = 'Changed'
    elif mutation == 'image':
        sh = next(s for s in p.slides[1].shapes if s._element.tag == qn('p:pic'))
        sh._element.getparent().remove(sh._element)
    elif mutation == 'chart':
        sh = next(s for s in p.slides[1].shapes if s.has_chart)
        sh.chart.part._element.xpath('.//c:numCache/c:pt/c:v')[0].text = '999'
    elif mutation == 'link':
        sh = next(s for s in p.slides[1].shapes if s.has_text_frame and s.text == 'External reference')
        sh.text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://example.edu/wrong'
    elif mutation == 'notes':
        p.slides[1].notes_slide.notes_text_frame.text = 'Changed notes'
    elif mutation == 'duplicate':
        el = deepcopy(p.slides[0].shapes[1]._element)
        p.slides[0].shapes._spTree.insert_element_before(el,'p:extLst')
    elif mutation == 'hidden':
        p.slides[0].shapes[1]._element.xpath('.//p:cNvPr')[0].set('hidden','1')
    p.save(candidate)
    assert audit(source,candidate,report)['status'] == 'failed'


def test_split_keeps_notes_and_internal_links(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = 'Dense content'
    s.placeholders[1].text = '\n'.join(f'Paragraph number {i}' for i in range(14))
    s.notes_slide.notes_text_frame.text = 'Original teaching notes'
    other = p.slides.add_slide(p.slide_layouts[6]); sh = text(other,'Back to dense',1)
    sh.click_action.target_slide = s
    src=tmp_path/'source.pptx'; out=tmp_path/'out.pptx'; p.save(src)
    report=grounded.build_deck(src,out,revisions={'0':{'tags':['split']}})
    assert len(report['source_to_output_slides']['0']) == 2
    assert audit(src,out,report)['status'] == 'passed'
    result=Presentation(out)
    assert result.slides[-1].shapes[0].click_action.target_slide == result.slides[0]


def test_unsupported_object_cannot_disappear(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement
    src=fixture(tmp_path); p=Presentation(src)
    p.slides[0]._element.append(OxmlElement('p:timing')); p.save(src)
    assert inspect(src).unsupported
    with pytest.raises(CoverageError):
        grounded.build_deck(src,tmp_path/'out.pptx')


def test_same_text_repeated_is_not_deduplicated(tmp_path):
    src=fixture(tmp_path); p=Presentation(src)
    p.slides[0].shapes[1].text='Repeat this\nRepeat this\nRepeat this'
    p.save(src); out=tmp_path/'out.pptx'
    report=grounded.build_deck(src,out)
    assert audit(src,out,report)['status']=='passed'
    p=Presentation(out); p.slides[0].shapes[1].text='Repeat this\nRepeat this';p.save(out)
    assert audit(src,out,report)['status']=='failed'


@pytest.mark.parametrize('mutation', ['preset', 'adjustment', 'custom_path'])
def test_geometry_ignores_unused_namespaces_but_detects_edits(tmp_path, mutation):
    from lxml import etree
    from zipfile import ZipFile, ZIP_DEFLATED
    from pptx.enum.shapes import MSO_SHAPE
    from slide_engine.inventory import canonical
    a='http://schemas.openxmlformats.org/drawingml/2006/main'
    p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6])
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1))
    shape.adjustments[0]=.2
    shape.text='Preserve this geometry'
    if mutation=='custom_path':
        geom=shape._element.xpath('./p:spPr/a:prstGeom')[0]
        custom=etree.fromstring(f'<a:custGeom xmlns:a="{a}"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/><a:pathLst><a:path w="100" h="100"><a:moveTo><a:pt x="0" y="0"/></a:moveTo><a:lnTo><a:pt x="100" y="100"/></a:lnTo></a:path></a:pathLst></a:custGeom>')
        geom.getparent().replace(geom,custom)
    source=tmp_path/'minimal-namespaces.pptx';p.save(source)
    # This deck has no relationships in its slide XML: remove unused declarations,
    # reproducing slides emitted by tools other than PowerPoint/python-pptx.
    with ZipFile(source) as archive: parts={n:archive.read(n) for n in archive.namelist()}
    element=etree.fromstring(parts['ppt/slides/slide1.xml'])
    etree.cleanup_namespaces(element)
    parts['ppt/slides/slide1.xml']=etree.tostring(element,xml_declaration=True,encoding='UTF-8')
    with ZipFile(source,'w',ZIP_DEFLATED) as archive:
        for name,data in parts.items():archive.writestr(name,data)
    output=tmp_path/'output.pptx';report=grounded.build_deck(source,output)
    assert audit(source,output,report)['status']=='passed'
    original=Presentation(source).slides[0].shapes[0]._element.xpath('./p:spPr/a:prstGeom | ./p:spPr/a:custGeom')[0]
    result=Presentation(output)
    copied=result.slides[0].shapes[0]._element.xpath('./p:spPr/a:prstGeom | ./p:spPr/a:custGeom')[0]
    assert etree.tostring(original,method='c14n')!=etree.tostring(copied,method='c14n')
    assert canonical(original)==canonical(copied)
    if mutation=='preset':copied.set('prst','ellipse')
    elif mutation=='adjustment':copied.xpath('.//a:gd')[0].set('fmla','val 9999')
    else:copied.xpath('.//a:lnTo/a:pt')[0].set('x','50')
    result.save(output)
    assert any(f['code']=='ALTERED_OR_DUPLICATED' for f in audit(source,output,report)['findings'])


def test_table_split_preserves_rows_and_repeats_header(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[5]);s.shapes.title.text='Long table'
    table=s.shapes.add_table(15,2,Inches(1),Inches(2),Inches(7),Inches(4)).table
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells):cell.text=f'Row {i}, column {j}'
    src=tmp_path/'source.pptx';out=tmp_path/'out.pptx';p.save(src)
    report=grounded.build_deck(src,out,revisions={'0':{'tags':['split']}})
    assert len(report['source_to_output_slides']['0'])==2
    assert audit(src,out,report)['status']=='passed'


def test_run_fragmentation_passes_but_emphasis_change_fails(tmp_path):
    src=fixture(tmp_path);out=tmp_path/'out.pptx';report=grounded.build_deck(src,out)
    p=Presentation(out);para=p.slides[0].shapes[1].text_frame.paragraphs[0]
    value=para.runs[0].text
    para.runs[0].text=value[:5];para.add_run().text=value[5:];p.save(out)
    assert audit(src,out,report)['status']=='passed'
    p=Presentation(out);p.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].font.bold=True;p.save(out)
    assert audit(src,out,report)['status']=='failed'
