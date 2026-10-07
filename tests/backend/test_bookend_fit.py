"""Bookends are limited by rendered space, with correction before composing a deck."""
from copy import deepcopy

import pytest
from pptx import Presentation
from pptx.util import Pt

from app import sessions
from app.authoring import composer, service
from app.authoring.models import CreationRequest, DeckSpec, Outline, OutlineSlide, SlideSpec, VisualPlan
from slide_engine import templates


LONG = [
    'Connect partial differential equations with physical models, mathematical structure, and observed data. '
    'Use characteristic curves and analytical tools to explain solution behavior and interpret conclusions.',
]
MANY = ['Models and evidence', 'Equations and data', 'Questions and discussion']
OVERFLOW = [f'Step {i}: Explain the detailed method with all its assumptions, examples, proof steps and physical interpretation.'
            for i in range(8)]


@pytest.fixture(autouse=True)
def cleanup_created_jobs():
    before = set(sessions._sessions)
    yield
    for sid in set(sessions._sessions) - before:
        sessions.delete(sid)


def approved_session(template='stevens'):
    sess = service.create(CreationRequest(topic='PDEs', audience='Students', template_id=template))
    outline = Outline(title='PDEs', rationale='An introduction, a lesson and a closing.', slides=[
        OutlineSlide(id=kind, kind=kind, title='Thank you!' if kind == 'closing' else 'PDEs',
                     points=['Understand the idea.'], visual=VisualPlan(
                         kind='text_only', description='Explain in words.', reason='Text is sufficient.'))
        for kind in ('opening','content','closing')])
    service.save_outline(sess, outline, 0)
    service.approve(sess, 1)
    return sess, outline


def mock_render(monkeypatch):
    monkeypatch.setattr(service.render_verify, 'check', lambda *a: {'status':'passed','findings':[]})


@pytest.mark.parametrize('template', ['stevens','cpe'])
@pytest.mark.parametrize('bullets', [LONG, MANY])
def test_actual_fit_accepts_more_than_180_characters_or_two_points(template, bullets, tmp_path):
    assert len(bullets)>2 or sum(map(len,bullets))>180
    deck=DeckSpec(slides=[SlideSpec(id=kind,title='Thank you!' if kind=='closing' else 'Introduction to PDEs',bullets=bullets)
                          for kind in ('opening','closing')])
    with templates.use(template):
        for slide, kind in zip(deck.slides, ('opening','closing')):
            composer.validate_bookend(slide, kind)
        path=tmp_path/'bookends.pptx'
        manifest=composer.compose(deck,path,tmp_path/'assets',kinds=['opening','closing'])
        assert composer.audit(path,manifest)['status']=='passed'
        for slide in Presentation(path).slides:
            body=next(s for s in slide.shapes if s.name=='authored-body')
            assert body.text=='\n'.join(bullets)
            assert all(Pt(16)<=p.font.size<=Pt(20) for p in body.text_frame.paragraphs)


@pytest.mark.parametrize('template', ['stevens','cpe'])
@pytest.mark.parametrize('kind', ['opening','closing'])
def test_short_character_count_can_still_overflow_actual_height(template,kind):
    slide=SlideSpec(id=kind,title='PDEs',bullets=['\n'.join(['Short line']*14)])
    assert sum(map(len,slide.bullets))<180
    with templates.use(template), pytest.raises(ValueError,match='speaker notes'):
        composer.validate_bookend(slide,kind)


@pytest.mark.parametrize('kind',['opening','closing'])
def test_generated_overflow_is_corrected_before_authoring_other_slides(monkeypatch,kind):
    sess,outline=approved_session()
    requests=[]
    def provider(role,instruction,payload,**kwargs):
        requests.append(deepcopy(payload))
        item=payload['slide']
        bullets=OVERFLOW if item['kind']==kind and 'repair_instruction' not in payload else MANY
        return {'status':'completed','data':SlideSpec(id=item['id'],title=item['title'],bullets=bullets,notes='Keep these speaker notes.').model_dump()}
    monkeypatch.setattr(service.providers,'generate',provider)
    mock_render(monkeypatch)
    with sessions.job(sess):
        result=service.generate(sess)
    expected=['opening','opening','content','closing'] if kind=='opening' else ['opening','content','closing','closing']
    assert [p['slide']['id'] for p in requests]==expected
    correction=next(p for p in requests if 'repair_instruction' in p)
    assert 'speaker notes' in correction['repair_instruction']
    assert correction['previous_response']['bullets']==OVERFLOW
    fixed=next(s for s in result['deck']['slides'] if s['id']==kind)
    assert fixed['bullets']==MANY and 'Keep these speaker notes.' in fixed['notes']
    assert all(line in fixed['notes'] for line in OVERFLOW)
    assert result['generation']['checks']['artifact_coverage']['status']=='passed'
    prs=Presentation(sess.generation['candidate'])
    index=0 if kind=='opening' else 2
    assert all(line in prs.slides[index].notes_slide.notes_text_frame.text for line in OVERFLOW)


def test_retry_stays_bounded_and_does_not_compose_unreadable_output(monkeypatch):
    sess,outline=approved_session()
    calls=[]
    def provider(*args,**kwargs):
        calls.append(args[2])
        return {'status':'completed','data':SlideSpec(id='opening',title='PDEs',bullets=OVERFLOW).model_dump()}
    monkeypatch.setattr(service.providers,'generate',provider)
    with sessions.job(sess),pytest.raises(ValueError,match='for slide opening'):
        service.generate(sess)
    assert len(calls)==2 and sess.generation is None and sess.creation['deck'] is None


def test_manual_text_that_fits_is_preserved_and_overflow_is_not_silently_rewritten(monkeypatch):
    sess,outline=approved_session()
    deck=DeckSpec(slides=[SlideSpec(id=s.id,title=s.title,bullets=LONG) for s in outline.slides])
    monkeypatch.setattr(service.providers,'generate',lambda *a,**k:pytest.fail('Manual text must not be rewritten by AI.'))
    mock_render(monkeypatch)
    with sessions.job(sess):
        service.generate(sess,deck,content_edit=True)
    assert sess.creation['deck']==deck.model_dump()
    deck.slides[-1].bullets=OVERFLOW
    with sessions.job(sess),pytest.raises(ValueError,match='Closing slide closing'):
        service.generate(sess,deck,content_edit=True)


def test_visual_and_identity_contracts_still_apply():
    sess,outline=approved_session()
    with pytest.raises(ValueError,match='approved outline ID and title'):
        service.validate_authored_slide(outline.slides[0],SlideSpec(id='changed',title='PDEs'))
    with pytest.raises(ValueError,match='text only'):
        composer.validate_bookend(SlideSpec(id='opening',title='PDEs',equation='x=1'),'opening')


def test_qa_repair_of_a_bookend_uses_the_same_fit_correction(monkeypatch):
    sess,outline=approved_session()
    deck=DeckSpec(slides=[SlideSpec(id=s.id,title=s.title,bullets=MANY) for s in outline.slides])
    reviews=[]
    def qa(*args,**kwargs):
        reviews.append(1)
        return {'output_qa_visual':{'status':'failed' if len(reviews)==1 else 'passed','findings':[
            {'code':'OUTPUT_VISUAL','severity':'blocking','output_slide':2,
             'message':'Clarify the takeaway.', 'required_correction':'Clarify the takeaway.'}]
            if len(reviews)==1 else []}}
    calls=[]
    def provider(role,instruction,payload,**kwargs):
        calls.append(deepcopy(payload))
        return {'status':'completed','data':SlideSpec(id='closing',title='Thank you!',
                 bullets=OVERFLOW if len(calls)==1 else MANY).model_dump()}
    monkeypatch.setattr(service.output_qa,'run',qa)
    monkeypatch.setattr(service.providers,'generate',provider)
    mock_render(monkeypatch)
    with sessions.job(sess):
        result=service.generate(sess,deck)
    assert len(calls)==2 and len(reviews)==2
    assert 'speaker notes' in calls[-1]['repair_instruction']
    assert all(line in result['deck']['slides'][-1]['notes'] for line in OVERFLOW)
    assert result['generation']['checks']['artifact_coverage']['status']=='passed'


def test_preserved_details_do_not_bypass_the_notes_limit():
    previous=SlideSpec(id='closing',title='Thank you!',bullets=OVERFLOW)
    corrected=SlideSpec(id='closing',title='Thank you!',bullets=MANY,notes='N'*7990)
    with pytest.raises(ValueError):
        service.preserve_bookend_details(previous.model_dump(),corrected)


@pytest.mark.parametrize('previous',[None,[],{'bullets':None},{'bullets':'invalid string'}])
def test_schema_correction_does_not_try_to_preserve_malformed_bullets(previous):
    corrected=SlideSpec(id='closing',title='Thank you!',bullets=MANY)
    assert service.preserve_bookend_details(previous,corrected)==corrected
