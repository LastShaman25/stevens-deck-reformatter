"""Adversarial coverage and release cases independent of the normal workflow."""
import json
from copy import deepcopy
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches
from app import grounded, generations, benchmark, api, rendering
from app.qa import artifact_coverage, render_verify
from slide_engine.inventory import inspect, validate_placements
from test_regressions import fixture
from test_release import sess, fake_render


@pytest.mark.parametrize('kind',[XL_CHART_TYPE.LINE,XL_CHART_TYPE.PIE])
def test_editable_chart_and_workbook_preserved(tmp_path,kind):
    src=fixture(tmp_path);p=Presentation(src)
    data=CategoryChartData();data.categories=['First','Second'];data.add_series('Values',[4,9])
    p.slides[1].shapes.add_chart(kind,Inches(8),Inches(1),Inches(2),Inches(2),data)
    p.save(src);out=tmp_path/'out.pptx';report=grounded.build_deck(src,out)
    assert artifact_coverage.audit(src,out,report)['status']=='passed'
    chart=next(sh.chart for sh in Presentation(out).slides[1].shapes if sh.has_chart)
    assert tuple(chart.series[0].values)==(4,9)
    assert chart.part.chart_workbook.xlsx_part.blob


def test_paragraph_emphasis_does_not_leak(tmp_path):
    src=fixture(tmp_path);p=Presentation(src)
    paragraphs=p.slides[0].shapes[1].text_frame.paragraphs
    paragraphs[0].font.bold=True
    p.save(src);out=tmp_path/'out.pptx';report=grounded.build_deck(src,out)
    actual=Presentation(out).slides[0].shapes[1].text_frame.paragraphs
    assert actual[0].runs[0].font.bold is True
    assert actual[1].runs[0].font.bold is False
    assert artifact_coverage.audit(src,out,report)['status']=='passed'


def test_inherited_placeholder_emphasis_preserved(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Inherited heading'
    layout_title=next(sh for sh in s.slide_layout.placeholders if sh.placeholder_format.idx==s.shapes.title.placeholder_format.idx)
    layout_title.text_frame.paragraphs[0].font.italic=True
    src=tmp_path/'source.pptx';out=tmp_path/'out.pptx';p.save(src)
    report=grounded.build_deck(src,out)
    title=next(sh for sh in Presentation(out).slides[0].shapes if sh.name.endswith('|title'))
    assert title.text_frame.paragraphs[0].runs[0].font.italic is True
    # Reading python-pptx font.color can create an empty solidFill: never append a second.
    from zipfile import ZipFile
    from lxml import etree
    with ZipFile(out) as package:
        xml=etree.fromstring(package.read('ppt/slides/slide1.xml'))
        assert not xml.xpath('//*[local-name()="rPr"][count(*[local-name()="solidFill"]) > 1]')
    assert artifact_coverage.audit(src,out,report)['status']=='passed'


def test_plan_rejects_cross_slide_and_invented_ids(tmp_path):
    src=fixture(tmp_path);out=tmp_path/'out.pptx';r=grounded.build_deck(src,out)
    placements=deepcopy(r['placements']);placements[0]['source_slide']=999
    placements.append({'source_id':'invented','source_slide':0,'output_slide':0})
    codes={f['code'] for f in validate_placements(inspect(src),placements)['findings']}
    assert {'INVALID_PLACEMENT_SLIDE','INVENTED_PLACEMENT'}<=codes


def test_notes_link_and_image_crop_mutations(tmp_path):
    from test_artifacts import rich_fixture
    src=rich_fixture(tmp_path);p=Presentation(src)
    p.slides[1].notes_slide.notes_text_frame.paragraphs[0].runs[0].hyperlink.address='https://example.edu/notes'
    p.save(src);out=tmp_path/'out.pptx';r=grounded.build_deck(src,out)
    assert artifact_coverage.audit(src,out,r)['status']=='passed'
    p=Presentation(out)
    p.slides[1].notes_slide.notes_text_frame.paragraphs[0].runs[0].hyperlink.address='https://example.edu/changed'
    p.save(out)
    assert artifact_coverage.audit(src,out,r)['status']=='failed'
    r=grounded.build_deck(src,out);p=Presentation(out)
    picture=next(sh for sh in p.slides[1].shapes if sh.shape_type==13);picture.crop_left=.4;p.save(out)
    assert artifact_coverage.audit(src,out,r)['status']=='failed'


def test_optional_ai_rerun_requires_new_review(sess,monkeypatch):
    monkeypatch.setattr(generations.render_verify,'check',fake_render)
    r=generations.build(sess)
    monkeypatch.setattr(api.optional_review,'review_slide_result',lambda *a:{'status':'completed','findings':[{'note':'Inspect this visual'}]})
    api.run_ai(sess,0)
    decision=generations.Decision(generation_id=r['generation_id'],candidate_sha256=r['candidate_sha256'],finding_ids=['ai:0:0'],rationale='Inspected old response')
    generations.decide(sess,decision)
    api.run_ai(sess,0)
    assert not any('ai:0:0' in d['finding_ids'] for d in r['human_decisions'])
    assert r['state']=='needs_review'


def test_benchmark_is_exact_and_idempotent(sess,tmp_path,monkeypatch):
    monkeypatch.setattr(generations.render_verify,'check',fake_render)
    r=generations.build(sess)
    ids=[f['id'] for f in r['findings'] if f['severity']=='review']
    if ids:generations.decide(sess,generations.Decision(generation_id=r['generation_id'],candidate_sha256=r['candidate_sha256'],finding_ids=ids,rationale='Synthetic inspection'))
    store=tmp_path/'benchmarks';log=tmp_path/'benchmarks.jsonl'
    monkeypatch.setattr(benchmark,'_STORE',str(store));monkeypatch.setattr(benchmark,'_LOG',str(log))
    first=benchmark.capture(sess);assert benchmark.capture(sess)==first
    assert len(log.read_text().splitlines())==1
    saved=json.loads(log.read_text())
    assert saved['candidate_sha256']==generations.sha256(store/first['deck'])
    assert saved['patterns']['layout_histogram']=={'1_Title Slide':1, 'Title Only':2}


def test_render_page_mismatch_and_stale_directory_block(tmp_path,monkeypatch):
    import fitz
    src=fixture(tmp_path)
    def missing_pages(candidate,pdf):
        doc=fitz.open();doc.new_page();doc.save(pdf);doc.close()
        Path(pdf).with_suffix('.renderer.json').write_text(json.dumps({'renderer':'injected test','version':'test'}))
    monkeypatch.setattr(rendering,'render_to_pdf',missing_pages)
    result=render_verify.check(src,tmp_path/'render')
    assert result['status']=='failed'
    assert any(f['code']=='RENDER_PAGE_COUNT' for f in result['findings'])
    with pytest.raises(FileExistsError):render_verify.check(src,tmp_path/'render')


def test_template_and_policy_versions_invalidate(sess,monkeypatch):
    monkeypatch.setattr(generations.render_verify,'check',fake_render)
    r=generations.build(sess)
    r['template_sha256']='old'
    with pytest.raises(ValueError,match='STALE_POLICY_OR_TEMPLATE'):generations.verify_identity(sess,r,ready=False)
    r['template_sha256']=generations.sha256(grounded.TEMPLATE_PATH);r['policy_version']='old'
    with pytest.raises(ValueError,match='STALE_POLICY_OR_TEMPLATE'):generations.verify_identity(sess,r,ready=False)


def test_renderer_timeout_and_missing_output_are_errors(tmp_path,monkeypatch):
    import subprocess
    monkeypatch.setattr(rendering,'renderer_info',lambda:{'name':'LibreOffice','executable':'fake-renderer'})
    def timeout(*args,**kwargs):raise subprocess.TimeoutExpired('fake-renderer',1)
    monkeypatch.setattr(rendering.subprocess,'run',timeout)
    with pytest.raises(RuntimeError,match='timeout'):rendering.pptx_to_pdf(tmp_path/'source.pptx',tmp_path/'timeout')
    monkeypatch.setattr(rendering.subprocess,'run',lambda *a,**k:subprocess.CompletedProcess('fake',0,b'',b''))
    with pytest.raises(RuntimeError,match='did not create'):rendering.pptx_to_pdf(tmp_path/'source.pptx',tmp_path/'missing')


def test_merged_table_preserved_and_merge_mutation_detected(tmp_path):
    src=fixture(tmp_path);p=Presentation(src);table=next(sh.table for sh in p.slides[1].shapes if sh.has_table)
    table.cell(0,0).merge(table.cell(0,1));p.save(src)
    out=tmp_path/'out.pptx';report=grounded.build_deck(src,out)
    assert artifact_coverage.audit(src,out,report)['status']=='passed'
    p=Presentation(out);table=next(sh.table for sh in p.slides[1].shapes if sh.has_table);table.cell(0,0).split();p.save(out)
    assert artifact_coverage.audit(src,out,report)['status']=='failed'


def test_renderer_cannot_omit_a_raster_picture(tmp_path,monkeypatch):
    import fitz
    from PIL import Image
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    picture=tmp_path/'image.png';Image.new('RGB',(300,300),'red').save(picture)
    s.shapes.add_picture(str(picture),Inches(2),Inches(2),Inches(2),Inches(2))
    candidate=tmp_path/'candidate.pptx';p.save(candidate)
    def blank(candidate,pdf):
        doc=fitz.open();doc.new_page(width=720,height=540);doc.save(pdf);doc.close()
        Path(pdf).with_suffix('.renderer.json').write_text(json.dumps({'renderer':'injected','version':'test'}))
    monkeypatch.setattr(rendering,'render_to_pdf',blank)
    result=render_verify.check(candidate,tmp_path/'render')
    assert result['status']=='failed'
    assert any(f['code']=='RENDER_IMAGE_UNVERIFIED' and f['severity']=='blocking' for f in result['findings'])


def test_parallel_previews_do_not_conflict(sess,tmp_path):
    from app import sessions
    from fastapi.testclient import TestClient
    from app.main import app
    Path(sess.preview_path(0,'before')).write_bytes(b'png')
    with sessions.read_job(sess):
        response=TestClient(app).get(f'/api/sessions/{sess.id}/slides/0/preview?variant=before')
        assert response.status_code==200
        assert response.content==b'png'
        assert sessions.delete(sess.id) is False
        assert Path(sess.dir).exists()
    assert not Path(sess.dir).exists()
