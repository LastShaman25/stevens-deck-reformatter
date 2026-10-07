"""Reject unusable authored math before composition and repair the same slide."""
from copy import deepcopy
from types import SimpleNamespace

import numpy as np
import pytest
from PIL import Image
from matplotlib.mathtext import MathTextParser

from app.authoring import graphics, service
from app.authoring.models import PlotSpec, SlideSpec


@pytest.mark.parametrize('expression,reason', [
    ('y = (x - c)/2', 'without'), ('(x-c)/2', 'Undefined plot parameter'),
    ('x = 2', 'without'), ('__import__("os")', 'Unsupported function syntax'),
    ('x.__class__', 'Unsupported function syntax'), ('sqrt(-1)', 'no finite values'),
])
def test_plot_preflight_rejects_unrenderable_content(expression, reason):
    with pytest.raises(ValueError, match=reason):
        PlotSpec(functions=[expression])


@pytest.mark.parametrize('plot,reason', [
    ({'functions':[]}, 'No functions'),
    ({'functions':['sin(x)'], 'log_x':True}, 'positive domain'),
    ({'functions':['-1'], 'log_y':True}, 'no finite values'),
    ({'kind':'scatter', 'x':[1,2], 'y':[3]}, 'lengths'),
    ({'kind':'histogram'}, 'empty'),
    ({'kind':'heatmap', 'matrix':[[]]}, 'nonempty'),
])
def test_all_plot_kinds_fail_during_schema_validation(plot, reason):
    with pytest.raises(ValueError, match=reason): PlotSpec(**plot)


@pytest.mark.parametrize('expression', ['1/x', 'log(x)', 'sqrt(x)', 'sin(x)', '(x-1)/2'])
def test_valid_functions_keep_supported_partial_domains(tmp_path, expression):
    spec=PlotSpec(functions=[expression])
    graphics.render_plot(spec, tmp_path/'plot.png')
    with Image.open(tmp_path/'plot.png') as im:
        assert im.width >= 1000 and im.height >= 600


@pytest.mark.parametrize('bad_function', ['y = (x-c)/2', '(x-c)/2'])
def test_model_gets_one_targeted_correction_before_acceptance(monkeypatch, bad_function):
    bad={'id':'s6', 'title':'Characteristics', 'plot':{'functions':[bad_function]}}
    corrected=deepcopy(bad); corrected['plot']['functions']=['(x-1)/2']
    corrected['notes']='Illustrative parameter c=1; y=(x-c)/2.'
    responses=[bad, corrected]; requests=[]
    def provider(role, system, payload, **kwargs):
        requests.append(deepcopy(payload))
        return {'status':'completed', 'data':responses.pop(0)}
    monkeypatch.setattr(service.providers, 'generate', provider)
    result=service.call(SimpleNamespace(ensure_active=lambda:None), 'author', 'Write the approved slide.',
                        {'slide':{'id':'s6'}}, SlideSpec)
    assert result.plot.functions == ['(x-1)/2'] and len(requests)==2
    assert requests[1]['previous_response']==bad
    assert 'plot:' in requests[1]['repair_instruction'] and 'parameter' in requests[1]['repair_instruction']
    np.testing.assert_allclose(graphics.evaluate(result.plot.functions[0], np.array([1.,3.])), [0.,1.])


def test_failed_correction_is_bounded_and_explains_the_slide(monkeypatch):
    calls=[]
    def provider(*args, **kwargs):
        calls.append(args)
        return {'status':'completed', 'data':{'id':'s6', 'title':'Characteristics',
                                              'plot':{'functions':['x+c']}}}
    monkeypatch.setattr(service.providers, 'generate', provider)
    with pytest.raises(ValueError, match="slide s6: .*Undefined plot parameter 'c'"):
        service.call(SimpleNamespace(ensure_active=lambda:None), 'author', 'Write the slide.',
                     {'slide':{'id':'s6'}}, SlideSpec)
    assert len(calls)==2


def test_multiline_equations_are_actually_parsed_as_math(tmp_path, monkeypatch):
    source=r'-\partial_t^2 u + 2\partial_x^2 u + u = t'+'\n'+r'-\partial_t^2 u + (1+\cos u)\partial_x^3 u = 0'
    calls=[]; original=MathTextParser.parse
    def parsed(self, text, *args, **kwargs):
        calls.append(text)
        return original(self, text, *args, **kwargs)
    monkeypatch.setattr(MathTextParser, 'parse', parsed)
    slide=SlideSpec(id='s2', title='PDE examples', equation=source)
    calls.clear()  # Observe the actual drawing, not schema validation alone.
    graphics.render_equation(slide.equation, tmp_path/'equation.png')
    for line in source.splitlines():
        assert calls.count('$'+line+'$') >= 2  # Preflight AND renderer use math mode.
    assert all('\n' not in text for text in calls)
    assert slide.equation == source
    with Image.open(tmp_path/'equation.png') as im: assert im.width>500 and im.height>100


def test_common_math_aliases_keep_exact_authored_specification(tmp_path):
    source=r'1\le p<\infty,\quad x\ge 0,\quad x\ne 1'
    slide=SlideSpec(id='norm',title='Conditions',equation=source)
    assert graphics.equation_lines(source)==[r'1\leq p<\infty,\quad x\geq 0,\quad x\neq 1']
    assert graphics.equation_lines(r'\left(x\right)\neq 1')==[r'\left(x\right)\neq 1']
    graphics.render_equation(slide.equation,tmp_path/'math.png')
    assert slide.equation==source


@pytest.mark.parametrize('expression', [r'\unsupported{x}', '$x$', '\n', '\\left(x\n\\right)'])
def test_invalid_equations_are_rejected_before_composition(expression):
    with pytest.raises(ValueError): SlideSpec(id='bad', title='Bad math', equation=expression)


def test_narrow_body_fit_preserves_wording_and_safe_region(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from app.authoring import composer,text_fit
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
    body=slide.shapes.add_textbox(Inches(.75),Inches(2),Inches(4),Inches(4.4))
    lines=['Model selection draws on theory and experimental data; there is no universal recipe.',
           'Given a PDE and data, ask: Do solutions exist? Are they unique? What are their qualitative properties? Can singularities form?',
           'Study how solutions respond to changes in data, and derive useful norms and quantitative estimates.',
           'Explicit formulas are often unavailable, so analysis helps us understand and estimate solutions.']
    composer.text(body,lines,20)
    before=(body.left,body.top,body.width,body.height)
    text_fit.fit_content_body(body)
    path=tmp_path/'fit.pptx'; prs.save(path)
    reopened=Presentation(path).slides[0].shapes[0]
    assert reopened.text=='\n'.join(lines)
    assert (reopened.left,reopened.top,reopened.width,reopened.height)==before
    assert all(Pt(16)<=p.font.size<=Pt(20) for p in reopened.text_frame.paragraphs)
    assert reopened.text_frame.paragraphs[-1].space_after==0
