"""Release checkpoints A/B/C/C-math/D/Q: real auth and QA contracts, synthetic data."""
import copy
import json
import time
from pathlib import Path
import fitz
import numpy as np
import pytest
from PIL import Image
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches
from app import auth, sessions, generations
from app.main import app
from app.ai import output_qa, providers
from app.authoring import service, composer, graphics
from app.authoring.models import CreationRequest, Outline, OutlineSlide, DeckSpec, SlideSpec, ChartSpec, PlotSpec, Series


@pytest.fixture(autouse=True)
def isolation(tmp_path, monkeypatch):
    monkeypatch.setenv('STEVENS_AUTH_DB', str(tmp_path/'accounts.sqlite3'))
    monkeypatch.setenv('STEVENS_AUTH_MODE', 'invitation')
    monkeypatch.setenv('STEVENS_PUBLIC_URL', 'http://localhost:8000')
    root = tmp_path/'workspaces'; root.mkdir()
    monkeypatch.setattr(sessions, '_ROOT', str(root))
    monkeypatch.setattr(sessions, '_sessions', {})
    auth._attempts.clear()
    yield
    for sid in list(sessions._sessions): sessions.delete(sid)


def login(code='admin'):
    client = TestClient(app)
    result = client.post('/api/auth/code', json={'code':code})
    assert result.status_code == 200, result.text
    client.headers['X-CSRF-Token'] = result.json()['csrf']
    return client


def new_job(client):
    result = client.post('/api/jobs', json={'topic':'Sine functions', 'audience':'Students', 'length_preference':'auto'})
    assert result.status_code == 200, result.text
    return result.json()['id']


def test_code_login_admin_members_and_cross_user_boundaries():
    anonymous = TestClient(app)
    assert anonymous.get('/api/jobs/unknown').status_code == 401
    admin = login()
    created = admin.post('/api/admin/codes', json={'name':'Student', 'role':'member'}).json()
    member = login(created['code'])
    sid = new_job(admin)
    for path in (f'/api/jobs/{sid}', f'/api/jobs/{sid}/source/1', f'/api/jobs/{sid}/preview/0?generation_id=x',
                 f'/api/jobs/{sid}/download?generation_id=x', f'/api/sessions/{sid}'):
        assert member.get(path).status_code == 404
    for method, tail in [('post','/cancel'),('post','/generate'),('post','/outline/approve'),('put','/outline'),('delete','')]:
        assert getattr(member,method)(f'/api/jobs/{sid}{tail}').status_code == 404
    assert member.get('/api/admin/users').status_code == 403
    assert member.post('/api/admin/codes',json={'name':'Nope'}).status_code == 403
    assert admin.get(f'/api/jobs/{sid}').status_code == 200
    assert admin.get(f'/api/jobs/{sid}').headers['cache-control'] == 'no-store'


def test_csrf_revocation_last_admin_and_code_hashing():
    admin = login()
    uid = admin.get('/api/auth/status').json()['user']['id']
    assert admin.patch(f'/api/admin/users/{uid}',json={'role':'member','active':True}).status_code == 409
    assert admin.post('/api/jobs',json={'audience':'X'},headers={'X-CSRF-Token':'wrong'}).status_code == 403
    assert admin.post('/api/jobs',json={'audience':'X'},headers={'Origin':'https://untrusted.example'}).status_code == 403
    created = admin.post('/api/admin/codes',json={'name':'Member'}).json()
    member = login(created['code']); sid = new_job(member)
    directory = Path(sessions.get(sid).dir)
    assert admin.patch('/api/admin/users/'+created['user']['id'],json={'role':'member','active':False}).status_code == 200
    assert member.get(f'/api/jobs/{sid}').status_code == 401
    assert not directory.exists()
    with auth.database() as db:
        assert created['code'] not in str([tuple(r) for r in db.execute('SELECT * FROM codes')])
        assert not db.execute("SELECT name FROM sqlite_master WHERE name LIKE '%deck%'").fetchall()


def test_logout_and_code_rotation_revoke_login():
    admin = login(); code = admin.post('/api/admin/codes',json={'name':'Member'}).json()
    member = login(code['code']); sid = new_job(member)
    changed = admin.post('/api/admin/users/'+code['user']['id']+'/rotate-code',json={})
    assert changed.status_code == 200
    assert member.get(f'/api/jobs/{sid}').status_code == 401
    assert not sessions._sessions
    member = login(changed.json()['code']); sid = new_job(member)
    assert member.post('/api/auth/logout',json={}).status_code == 200
    assert sessions.get(sid) is None


def test_account_removal_requires_cleanup_and_last_admin_survives():
    admin=login(); uid=admin.get('/api/auth/status').json()['user']['id']
    assert admin.delete('/api/admin/users/'+uid).status_code==409
    added=admin.post('/api/admin/codes',json={'name':'Remove me'}).json()
    member=login(added['code']); sid=new_job(member)
    assert admin.delete('/api/admin/users/'+added['user']['id']).json()['deleted']
    assert member.get('/api/jobs/'+sid).status_code==401
    assert sessions.get(sid) is None


def test_admin_can_rotate_own_code_without_losing_the_new_code_screen():
    admin=login(); uid=admin.get('/api/auth/status').json()['user']['id']
    rotated=admin.post(f'/api/admin/users/{uid}/rotate-code',json={})
    assert rotated.status_code==200
    admin.headers['X-CSRF-Token']=rotated.json()['csrf']
    assert admin.get('/api/admin/users').status_code==200
    assert TestClient(app).post('/api/auth/code',json={'code':'admin'}).status_code==401
    replacement=login(rotated.json()['code'])
    assert replacement.get('/api/admin/users').status_code==200


def test_native_table_cells_and_scatter_x_values_are_verified(tmp_path):
    from app.authoring.models import TableSpec
    spec=DeckSpec(slides=[SlideSpec(id='t',title='Data',table=TableSpec(headers=['A','B'],rows=[['1','2']])),
        SlideSpec(id='c',title='Coordinates',chart=ChartSpec(kind='scatter',categories=['1','2'],series=[Series(name='y',values=[3,4])]))])
    candidate=tmp_path/'native.pptx';manifest=composer.compose(spec,candidate,tmp_path/'assets')
    assert composer.audit(candidate,manifest)['status']=='passed'
    prs=Presentation(candidate);table=next(s.table for s in prs.slides[0].shapes if s.has_table);table.cell(1,1).text='999';prs.save(candidate)
    assert composer.audit(candidate,manifest)['status']=='failed'


def test_cleanup_pin_retry_and_api_revocation(monkeypatch):
    admin = login(); sid = new_job(admin); sess = sessions.get(sid)
    with sessions.job(sess):
        assert admin.post(f'/api/jobs/{sid}/cancel',json={}).json()['deleted'] is False
        assert admin.get(f'/api/jobs/{sid}').status_code == 404
        assert Path(sess.dir).exists()
        with pytest.raises(ValueError): sess.ensure_active()
    assert not Path(sess.dir).exists()
    assert admin.post(f'/api/jobs/{sid}/cancel',json={}).json()['deleted'] is True
    sid = new_job(admin); sess = sessions.get(sid)
    original = sessions.shutil.rmtree
    monkeypatch.setattr(sessions.shutil,'rmtree',lambda *a,**k:(_ for _ in ()).throw(PermissionError('Locked')))
    assert sessions.delete(sid) is False
    assert sess.lifecycle == 'purge_failed' and Path(sess.dir).exists()
    monkeypatch.setattr(sessions.shutil,'rmtree',original); sessions.sweep()
    assert not Path(sess.dir).exists() and sid not in sessions._sessions


def test_polling_does_not_extend_lifetime_and_restart_purges():
    admin = login(); sid = new_job(admin); sess = sessions.get(sid)
    touched = sess.touched
    admin.get(f'/api/jobs/{sid}'); assert sess.touched == touched
    sess.expires = time.time()-1
    assert admin.get(f'/api/jobs/{sid}').status_code == 404
    assert not Path(sess.dir).exists()
    orphan = Path(sessions._ROOT,'old-worker'); orphan.mkdir(); (orphan/'private.txt').write_text('synthetic')
    sessions.cleanup_orphans(startup=True)
    assert not orphan.exists()


def test_readonly_processing_artifacts_are_deleted():
    import os, stat
    sess=sessions.create();child=Path(sess.dir,'readonly.txt');child.write_text('synthetic')
    os.chmod(child,stat.S_IREAD)
    assert sessions.delete(sess.id)
    assert not Path(sess.dir).exists()


def test_cancellation_waits_for_both_writer_and_reader_leases():
    sess=sessions.create()
    with sessions.read_job(sess):
        with sessions.job(sess):
            assert sessions.delete(sess.id) is False
        assert Path(sess.dir).exists()
    assert not Path(sess.dir).exists()


def simple_outline(count=2):
    return Outline(title='Functions', rationale='Enough space for each concept.', slides=[
        OutlineSlide(id=f's{i}', title=f'Function {i}', points=['Explain the idea']) for i in range(count)])


def test_adaptive_outline_approval_revision_and_pdf_reference_validation():
    sess = service.create(CreationRequest(topic='Topic', audience='Students'))
    outline = simple_outline()
    service.save_outline(sess,outline,0); service.approve(sess,1)
    first = sess.creation['approved_hash']
    updated = simple_outline(4)
    service.save_outline(sess,updated,1)
    assert sess.creation['approved_hash'] is None and len(sess.creation['outline']['slides']) == 4
    with pytest.raises(ValueError): service.approve(sess,1)
    service.approve(sess,2); assert sess.creation['approved_hash'] != first
    with pytest.raises(ValueError): CreationRequest(topic='X',audience='Y',slide_count=10)
    updated.slides[0].source_pages = [90]
    with pytest.raises(ValueError): service.save_outline(sess,updated,2)


def test_pdf_extraction_provenance_citations_and_rejection():
    sess = service.create(CreationRequest(audience='Students',source='pdf'))
    doc = fitz.open(); page = doc.new_page(); page.insert_text((72,72),'Measured revenue was 42 units in the controlled synthetic experiment.')
    data = doc.tobytes(); doc.close()
    pages = service.ingest_pdf(sess,data)
    assert pages[0]['page'] == 1 and '42 units' in pages[0]['text'] and Path(pages[0]['image']).exists()
    from app.authoring.models import Citation
    deck = DeckSpec(slides=[SlideSpec(id='s0',title='Evidence',citations=[Citation(page=1,quote='42 units')])])
    assert service.validate_citations(sess,deck)['status'] == 'passed'
    deck.slides[0].citations[0].quote = '84 units'
    assert service.validate_citations(sess,deck)['status'] == 'failed'
    with pytest.raises(ValueError): service.ingest_pdf(sess,b'not pdf')


@pytest.mark.parametrize('expression',['__import__("os")','x.__class__','[x for x in x]','open("file")','sin(x, x)','2**(2**10000)'])
def test_math_input_never_executes_code(expression):
    with pytest.raises((ValueError, SyntaxError)): graphics.evaluate(expression,np.array([1.,2.]))


def test_function_values_domains_and_math_render(tmp_path):
    x = np.array([0.,np.pi/2,np.pi])
    np.testing.assert_allclose(graphics.evaluate('sin(x)',x),[0,1,0],atol=1e-12)
    assert np.isnan(graphics.evaluate('sqrt(x)',np.array([-1.])))[0]
    for expr in ('1/x','log(x)','sqrt(x)','sin(x)'):
        graphics.render_plot(PlotSpec(functions=[expr]), tmp_path/(expr.replace('/','_')+'.png'))
    graphics.render_equation(r'\int_0^\pi \sin(x)\,dx = 2', tmp_path/'math.png')
    with Image.open(tmp_path/'math.png') as im: assert im.width > 300


def graphics_deck():
    return DeckSpec(slides=[
        SlideSpec(id='chart',title='Measured data',bullets=['Synthetic example'],chart=ChartSpec(kind='column',categories=['A','B'],series=[Series(name='Units',values=[2,4])])),
        SlideSpec(id='plot',title='Sine function',bullets=['A periodic function'],plot=PlotSpec(functions=['sin(x)'])),
        SlideSpec(id='math',title='Integral',bullets=['Area over half a period'],equation=r'\int_0^\pi \sin(x)\,dx = 2')])


def test_native_composition_detects_changed_chart_text_and_assets(tmp_path):
    candidate = tmp_path/'candidate.pptx'; manifest = composer.compose(graphics_deck(),candidate,tmp_path/'assets')
    assert composer.audit(candidate,manifest)['status'] == 'passed'
    prs = Presentation(candidate)
    prs.slides[0].shapes[0].text = 'Changed'
    prs.save(candidate)
    assert composer.audit(candidate,manifest)['status'] == 'failed'


def qa_fixture(tmp_path,count=12):
    prs=Presentation(); pages=[]
    for i in range(count):
        slide=prs.slides.add_slide(prs.slide_layouts[6]); slide.shapes.add_textbox(0,0, Inches(5), Inches(1)).text=f'Slide {i+1}'
        image=tmp_path/f'slide-{i}.png'; Image.new('RGB',(320,180),(i*15%255,230,255)).save(image)
        pages.append({'output_slide':i,'success':True,'png':str(image)})
    candidate=tmp_path/'candidate.pptx'; prs.save(candidate)
    return {'candidate':str(candidate),'checks':{'render_verification':{'candidate_sha256':generations.sha256(candidate),'pages':pages}}}


def test_ordered_qa_covers_all_12_slides_and_global_synthesis(tmp_path,monkeypatch):
    sess=sessions.create(); record=qa_fixture(tmp_path); requests=[]
    def provider(role,system,payload,images=(),**kwargs):
        requests.append((copy.deepcopy(payload),images))
        assert [int(Path(path).stem.split('-')[1])+1 for _,path in images] == payload['expected']
        return {'status':'completed','data':{'reviewed':payload['expected'],'summary':'Synthetic reviewer','findings':[]}}
    monkeypatch.setattr(providers,'generate',provider)
    checks=output_qa.run(sess,record)
    assert all(c['status']=='passed' for c in checks.values())
    assert [p['expected'] for p,_ in requests] == [list(range(1,6)),list(range(5,11)),[10,11,12],list(range(1,13))]
    assert len(record['output_manifest']) == 12


@pytest.mark.parametrize('fault',['missing','swapped','stale','hidden','partial_response','timeout','unknown_id'])
def test_qa_faults_block_verification(tmp_path,monkeypatch,fault):
    sess=sessions.create(); record=qa_fixture(tmp_path)
    pages=record['checks']['render_verification']['pages']
    if fault=='missing': pages.pop()
    if fault=='swapped': pages[1],pages[10]=pages[10],pages[1]
    if fault=='stale': record['checks']['render_verification']['candidate_sha256']='wrong'
    if fault=='hidden':
        prs=Presentation(record['candidate']);prs.slides[0]._element.set('show','0');prs.save(record['candidate'])
        record['checks']['render_verification']['candidate_sha256']=generations.sha256(record['candidate'])
    def provider(role,system,payload,**kwargs):
        return {'status':'timeout'} if fault=='timeout' else {'status':'completed','data':{
            'reviewed':payload['expected'][:-1] if fault=='partial_response' else payload['expected'],
            'summary':'Mock','findings':[{'slides':[999],'category':'accuracy','severity':'blocking','accuracy':'contradicted','message':'Invalid reference','evidence':''}] if fault=='unknown_id' else []}}
    monkeypatch.setattr(providers,'generate',provider)
    checks=output_qa.run(sess,record)
    assert checks['output_qa_coverage']['status']=='error'


@pytest.mark.renderer
def test_real_renderer_native_chart_plot_and_math(tmp_path):
    from app.qa.render_verify import check
    candidate=tmp_path/'native.pptx'; manifest=composer.compose(graphics_deck(),candidate,tmp_path/'assets')
    assert composer.audit(candidate,manifest)['status']=='passed'
    result=check(candidate,tmp_path/'render')
    assert result['status'] in ('passed','needs_review'), result['findings']
    assert len(result['pages'])==3


def test_authoring_api_approval_release_identity_and_cleanup(monkeypatch):
    admin=login();sid=new_job(admin)
    outline=simple_outline(1)
    assert admin.put(f'/api/jobs/{sid}/outline',json={'revision':0,'outline':outline.model_dump()}).status_code==200
    assert admin.post(f'/api/jobs/{sid}/outline/approve',json={'revision':1}).status_code==200
    def render(candidate,directory):
        directory=Path(directory);directory.mkdir(parents=True)
        Image.new('RGB',(400,220),'white').save(directory/'slide-0.png')
        return {'status':'passed','findings':[], 'candidate_sha256':generations.sha256(candidate),
                'pages':[{'output_slide':0,'success':True,'png':str(directory/'slide-0.png')}]}
    monkeypatch.setattr(service.render_verify,'check',render)
    monkeypatch.setattr(providers,'generate',lambda role,system,payload,**kw:{'status':'completed',
        'data':{'reviewed':payload['expected'],'summary':'Synthetic contract check','findings':[]}})
    deck=DeckSpec(slides=[SlideSpec(id='s0',title='Function 0',bullets=['One verified statement.'])])
    result=admin.put(f'/api/jobs/{sid}/content',json={'revision':1,'deck':deck.model_dump()})
    assert result.status_code==200,result.text
    record=result.json()['generation'];assert record['state']=='ready',record['findings']
    response=admin.get(f"/api/jobs/{sid}/download?generation_id={record['generation_id']}")
    assert response.status_code==200
    assert __import__('hashlib').sha256(response.content).hexdigest()==record['candidate_sha256']
    sess=sessions.get(sid);Path(sess.generation['output_manifest'][0]['image']).write_bytes(b'changed')
    assert admin.get(f"/api/jobs/{sid}/download?generation_id={record['generation_id']}").status_code==409
    directory=Path(sess.dir)
    assert admin.post(f'/api/jobs/{sid}/finalize',json={}).json()['deleted']
    assert not directory.exists()


def test_bounded_visual_repair_rerenders_and_reviews_final_candidate(monkeypatch):
    sess=service.create(CreationRequest(topic='Explain x',audience='Students'))
    service.save_outline(sess,simple_outline(1),0);service.approve(sess,1)
    renders=[]
    def render(candidate,directory):
        renders.append(str(candidate));Path(directory).mkdir(parents=True)
        png=Path(directory,'slide-0.png');Image.new('RGB',(400,220),'white').save(png)
        return {'status':'passed','findings':[], 'candidate_sha256':generations.sha256(candidate),
                'pages':[{'output_slide':0,'success':True,'png':str(png)}]}
    calls=[]
    def provider(role,system,payload,**kw):
        calls.append(role)
        if role=='author':
            return {'status':'completed','data':SlideSpec(id='s0',title='Function 0',bullets=['A clear explanation.']).model_dump()}
        findings=[{'slides':[1],'category':'visual','severity':'blocking','accuracy':'not_applicable',
                   'message':'Synthetic first-pass legibility defect','evidence':'Synthetic fixture'}] if len(renders)==1 else []
        return {'status':'completed','data':{'reviewed':payload['expected'],'summary':'Synthetic contract test','findings':findings}}
    monkeypatch.setattr(service.render_verify,'check',render);monkeypatch.setattr(providers,'generate',provider)
    result=service.generate(sess)
    assert result['generation']['state']=='ready'
    assert len(renders)==2 and renders[0]!=renders[1]
    assert calls.count('output_qa')==4 and calls.count('author')==2
    assert len(sess.generation['repair_history'])==1
