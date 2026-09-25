from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import grounded
from app.ai import source_decisions as D, layout
from app.qa.artifact_coverage import audit
from slide_engine import inventory, logo_assets as L
from slide_engine.preserve import marker
from test_source_decisions import decision

REGION=[805,40,975,355]


def fixture(tmp_path):
    p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
    s=p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='Source title must survive'
    with ZipFile(grounded.TEMPLATE_PATH) as z:
        image=Image.open(BytesIO(z.read('ppt/media/image2.png'))).convert('RGBA')
        image.thumbnail((1920,1080));stream=BytesIO();image.save(stream,format='PNG')
    picture=s.shapes.add_picture(stream,0,0,p.slide_width,p.slide_height)
    s.shapes._spTree.insert(2,picture._element)
    source=tmp_path/'source.pptx';p.save(source)
    prs,digest,objects,value=decision(source)
    role=next(e for e in value['elements'] if e['id']==next(o['id'] for o in objects if o['kind']=='picture'))
    # Keep the picture role last in this fixture for mutation tests.
    value['elements'].remove(role);value['elements'].append(role)
    role.update(role='background',confidence='high',contains_logo=True,content_bearing=False,artwork_action='extract_logo')
    shape=next(s for s in prs.slides[0].shapes if s.shape_type==13)
    blob,evidence=L.extract(shape,REGION)
    review=dict(complete_mark=True,background_only_removed=True,clean_edges_and_clear_space=True,
                source_pixels_faithful=True,template_identity_matches=False,reason='Synthetic independent reviewer fixture, not a live review.')
    value['logo_extractions']=[dict(source_id=role['id'],region=REGION,reason='Keep complete original mark; discard obsolete cover background.',
        verification=dict(source_image_sha256=evidence['source_image_sha256'],crop_sha256=evidence['crop_sha256'],review=review))]
    return source,prs,digest,value,blob,evidence


def test_verified_crop_replaces_background_with_exact_source_logo_pixels(tmp_path):
    source,prs,digest,value,blob,evidence=fixture(tmp_path)
    assert evidence['approved_template_match'] is not None
    candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    output=Presentation(candidate);picture=next(s for s in output.slides[0].shapes if s.name.endswith('|extracted|logo'))
    assert picture.image.blob==blob
    assert picture.image.blob!=next(s for s in prs.slides[0].shapes if s.shape_type==13).image.blob
    assert 'Source title must survive' in '\n'.join(s.text for s in output.slides[0].shapes if s.has_text_frame)
    assert audit(source,candidate,report)['status']=='passed'
    assert len([s for s in output.slides[0].shapes if s.shape_type==13])==1


def test_image_matching_proposes_a_valid_region_without_certifying_semantics(tmp_path):
    _,p,_,_,_,_=fixture(tmp_path)
    shape=next(s for s in p.slides[0].shapes if s.shape_type==13)
    suggestions=L.suggested_regions(L.raster(shape))
    assert suggestions and suggestions[0]['reference_asset']=='ppt/media/image2.png'
    assert L.extract(shape,suggestions[0]['region'])[1]['approved_template_match']
    assert L.suggested_regions(Image.new('RGBA',(1920,1080),'red'))==[]


@pytest.mark.parametrize('damage',['pixels','background_restored','crop','stretch','upscale','hidden','duplicate','missing'])
def test_independent_artifact_gate_rejects_logo_tampering(tmp_path,damage):
    source,prs,_,value,blob,_=fixture(tmp_path)
    candidate=tmp_path/'candidate.pptx';report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    p=Presentation(candidate);slide=p.slides[0];s=next(s for s in slide.shapes if s.name.endswith('|extracted|logo'))
    if damage in ('pixels','background_restored'):
        if damage=='background_restored':changed=next(s for s in prs.slides[0].shapes if s.shape_type==13).image.blob
        else:
            im=Image.open(BytesIO(blob));im.putpixel((15,15),(0,0,0,255));b=BytesIO();im.save(b,format='PNG');changed=b.getvalue()
        s.part.related_part(s._element.blipFill.blip.rEmbed)._blob=changed
    elif damage=='crop':s.crop_left=.1
    elif damage=='stretch':s.width=int(s.width*1.1)
    elif damage=='upscale':s.width*=8;s.height*=8
    elif damage=='hidden':s._element.xpath('.//p:cNvPr')[0].set('hidden','1')
    elif damage=='duplicate':slide.shapes._spTree.append(deepcopy(s._element))
    else:slide.shapes._spTree.remove(s._element)
    # Forged placement claims cannot waive independent pixel/aspect checks.
    for item in report['placements']:
        if item['source_id']==marker(s)[4:]:item['bounds']=[s.left,s.top,s.width,s.height]
    p.save(candidate)
    assert audit(source,candidate,report)['status']=='failed'


@pytest.mark.parametrize('damage',['no_review','hash','review_rejected','meaningful_content','no_logo','duplicate_request','rotated','cropped','stretched_source','whole_background','low_resolution','textured'])
def test_uncertain_or_unreviewed_extraction_cannot_be_composed(tmp_path,damage):
    _,p,digest,value,_,_=fixture(tmp_path)
    e=value['logo_extractions'][0];s=next(s for s in p.slides[0].shapes if s.shape_type==13)
    if damage=='no_review':e['verification']=None
    elif damage=='hash':e['verification']['crop_sha256']='forged'
    elif damage=='review_rejected':e['verification']['review']['complete_mark']=False
    elif damage=='meaningful_content':value['elements'][-1]['content_bearing']=True
    elif damage=='no_logo':value['elements'][-1]['contains_logo']=False
    elif damage=='duplicate_request':value['logo_extractions'].append(deepcopy(e))
    elif damage=='rotated':s.rotation=15
    elif damage=='cropped':s.crop_left=.1
    elif damage=='stretched_source':s.width=int(s.width*.8)
    elif damage=='whole_background':e['region']=[0,0,1000,1000]
    else:
        im=Image.open(BytesIO(s.image.blob)).convert('RGBA')
        if damage=='low_resolution':im.thumbnail((160,90))
        else:
            # Texture across the requested region's boundary defeats clean isolation.
            for y in range(40,390):
                for x in range(1500,1890):
                    if (x+y)%2:im.putpixel((x,y),(0,255,0,255))
        b=BytesIO();im.save(b,format='PNG');s.part.related_part(s._element.blipFill.blip.rEmbed)._blob=b.getvalue()
        # Reload: python-pptx caches image metadata.
        stream=BytesIO();p.save(stream);stream.seek(0);p=Presentation(stream)
    with pytest.raises(ValueError):D.validate(value,p,0,digest)


def test_separate_native_logo_is_reused_without_extraction_or_recoloring(tmp_path):
    source,p,_,value,_,_=fixture(tmp_path)
    value['logo_extractions']=[]
    value['elements'][-1].update(role='logo',content_bearing=True,artwork_action='retain')
    candidate=tmp_path/'native.pptx';report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    s=next(s for s in Presentation(candidate).slides[0].shapes if s.name.endswith('|logo'))
    assert s.image.blob==next(s for s in p.slides[0].shapes if s.shape_type==13).image.blob
    assert audit(source,candidate,report)['status']=='passed'


def test_identical_template_logo_is_reused_without_adding_a_duplicate(tmp_path):
    source,_,_,value,_,_=fixture(tmp_path)
    e=value['logo_extractions'][0];e['placement']='template_logo'
    e['verification']['review']['template_identity_matches']=True
    candidate=tmp_path/'reuse.pptx';report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    p=Presentation(candidate)
    assert not any(s.shape_type==13 for s in p.slides[0].shapes)
    assert audit(source,candidate,report)['status']=='passed'
    sh=next(s for origin,_,s in inventory.source_objects(p.slides[0]) if origin=='layout' and s.shape_id==6)
    sh.left+=Inches(.2);p.save(candidate)
    assert any(f['code']=='REUSED_TEMPLATE_LOGO_ALTERED' for f in audit(source,candidate,report)['findings'])


def test_template_reuse_requires_matching_variant_and_independent_identity_review(tmp_path):
    _,p,digest,value,blob,_=fixture(tmp_path)
    e=value['logo_extractions'][0];e['placement']='template_logo'
    with pytest.raises(ValueError,match='independent review'):D.validate(value,p,0,digest)
    # A recolored variant is not substituted merely because its silhouette matches.
    crop=Image.open(BytesIO(blob));pixels=crop.load()
    for y in range(crop.height):
        for x in range(crop.width):
            r,g,b,a=pixels[x,y]
            if min(r,g,b)>200:pixels[x,y]=(0,0,0,a)
    stream=BytesIO();crop.save(stream,format='PNG')
    with pytest.raises(ValueError,match='No identical approved'):L.template_logo(stream.getvalue(),True)


def test_cover_title_contrast_is_replaced_but_other_dependencies_remain_protected(tmp_path):
    _,p,digest,value,_,_=fixture(tmp_path)
    title=value['elements'][0];background=value['elements'][-1]
    title['related_ids']=[background['id']];background['related_ids']=[title['id']]
    D.validate(value,p,0,digest)  # Native cover text moves onto the known red template field.
    p.slides[0].shapes.add_textbox(0,0,Inches(2),Inches(1)).text='Required figure caption'
    from rubric_fixtures import role_map
    objects=D.source_objects(p,0,digest);obj=objects[-1]
    caption=role_map([obj])['elements'][0];caption.update(role='caption',related_ids=[background['id']])
    value['elements'].append(caption)
    with pytest.raises(ValueError,match='supporting retained content'):D.validate(value,p,0,digest)


@pytest.mark.parametrize('placement',['source_crop','template_logo'])
def test_source_stage_requires_a_separate_review_and_passes_raw_image(tmp_path,placement):
    source,_,_,value,_,_=fixture(tmp_path)
    value['logo_extractions'][0]['placement']=placement
    value['logo_extractions'][0]['verification']['review']['template_identity_matches']=placement=='template_logo'
    receipt=value['logo_extractions'][0].pop('verification');seen=[]
    preview=tmp_path/'before.png';Image.new('RGB',(100,60),'white').save(preview)
    sess=SimpleNamespace(source_path=source,revisions={},ensure_active=lambda:None,preview_path=lambda *args:str(preview))
    def provider(role,system,payload,images,max_tokens):
        seen.append((role,payload,images))
        if payload['stage']=='source_decisions':
            assert any('FULL RAW IMAGE' in label for label,_ in images)
            assert 'verification' not in payload['schema']['$defs']['LogoExtraction']['properties']
            return {'status':'completed','data':deepcopy(value)}
        assert role=='reviewer'
        assert [label for label,_ in images[:2]]==['ORIGINAL slide','FULL raw background image']
        if placement=='source_crop':assert images[2][0]=='EXTRACTED logo to verify'
        else:
            assert images[2][0].startswith('EQUAL-SCALE COMPARISON')
            assert Image.open(images[2][1]).size==(800,440)
        return {'status':'completed','data':receipt['review']}
    decisions,calls=D.run(sess,lambda **kwargs:None,generate=provider)
    assert len(seen)==2
    assert decisions['0']['logo_extractions'][0]['verification']==receipt
    assert calls[-1]['role']=='logo_extraction_review'


def test_rejected_crop_cannot_switch_to_plain_removal(tmp_path):
    source,_,_,value,_,_=fixture(tmp_path)
    receipt=value['logo_extractions'][0].pop('verification');receipt['review']['complete_mark']=False
    preview=tmp_path/'before.png';Image.new('RGB',(100,60),'white').save(preview)
    sess=SimpleNamespace(source_path=source,revisions={},ensure_active=lambda:None,preview_path=lambda *args:str(preview))
    def provider(role,system,payload,images,max_tokens):
        if payload['stage']=='logo_extraction_review':return {'status':'completed','data':receipt['review']}
        response=deepcopy(value)
        if 'validation_error' in payload:
            response['logo_extractions']=[];e=response['elements'][-1]
            e.update(contains_logo=False,artwork_action='remove');response['remove_ids']=[e['id']]
        return {'status':'completed','data':response}
    with pytest.raises(ValueError,match='cannot switch'):D.run(sess,lambda **kwargs:None,generate=provider)


def test_disputed_template_identity_returns_actionable_feedback_and_preserves_source_crop(tmp_path):
    source,_,_,value,_,_=fixture(tmp_path)
    receipt=value['logo_extractions'][0].pop('verification')
    value['logo_extractions'][0]['placement']='template_logo'
    receipt['review']['reason']='Synthetic reviewer cannot confirm the template variant.'
    preview=tmp_path/'before.png';Image.new('RGB',(100,60),'white').save(preview)
    sess=SimpleNamespace(source_path=source,revisions={},ensure_active=lambda:None,preview_path=lambda *args:str(preview))
    attempts=[]
    def provider(role,system,payload,images,max_tokens):
        if payload['stage']=='logo_extraction_review':return {'status':'completed','data':receipt['review']}
        response=deepcopy(value);attempts.append(payload.get('validation_error'))
        if 'validation_error' in payload:
            assert 'cannot confirm the template variant' in payload['validation_error']
            assert 'source_crop' in payload['validation_error']
            response['logo_extractions'][0]['placement']='source_crop'
        return {'status':'completed','data':response}
    decisions,_=D.run(sess,lambda **kwargs:None,generate=provider)
    assert len(attempts)==2 and decisions['0']['logo_extractions'][0]['placement']=='source_crop'


@pytest.mark.renderer
@pytest.mark.parametrize('placement',['source_crop','template_logo'])
def test_extracted_logo_deck_renders_in_powerpoint(tmp_path,placement):
    from app.qa.render_verify import check
    source,_,_,value,_,_=fixture(tmp_path)
    value['logo_extractions'][0]['placement']=placement
    value['logo_extractions'][0]['verification']['review']['template_identity_matches']=placement=='template_logo'
    candidate=tmp_path/'candidate.pptx';report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    assert audit(source,candidate,report)['status']=='passed'
    result=check(candidate,tmp_path/'render')
    assert result['status']=='passed'
    assert len(result['pages'])==1 and result['pages'][0]['success']
