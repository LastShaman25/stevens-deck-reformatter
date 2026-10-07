"""Clipped image placements must not disappear from graphic region inventory."""
from io import BytesIO

import fitz
import pytest
from PIL import Image, ImageChops, ImageDraw
from pptx import Presentation

from app import pdf_import, pdf_regions


@pytest.mark.parametrize('explicit_clip',[False,True])
def test_offpage_image_bounds_preserve_visible_plot_labels_and_arrow_overlays(tmp_path,explicit_clip):
    plot=Image.new('RGB',(1200,600),'white');draw=ImageDraw.Draw(plot)
    for inset in (10,60,110,170):
        draw.ellipse((inset,120+inset//3,1200-inset,550-inset//3),outline='black',width=4)
    draw.text((850,35),'Local Minimum',fill='red',stroke_width=1)
    blob=BytesIO();plot.save(blob,format='PNG')
    source=tmp_path/'clipped.pdf'
    with fitz.open() as doc:
        doc.new_page(width=720,height=405).insert_text((30,60),'Opening',fontsize=20)
        page=doc.new_page(width=720,height=405)
        page.insert_text((30,60),'Inspect the complete contour plot',fontsize=20)
        page.insert_image((-30,80,750,370),stream=blob.getvalue(),keep_proportion=False)
        if explicit_clip:
            image_stream=page.get_contents()[-1]
            # A second PDF clip hides peripheral image pixels. The importer
            # must preserve that mask rather than reveal the full source image.
            doc.update_stream(image_stream,b'q 40 55 640 270 re W n\n'+doc.xref_stream(image_stream)+b'\nQ')
        page.draw_line((150,150),(185,325),color=(.1,.3,.9),width=2)
        page.draw_line((185,325),(170,310),color=(.1,.3,.9),width=2)
        # This is the source of the historical loss: regular text extraction
        # omits the oversized image, although the renderer shows it on the page.
        assert not [b for b in page.get_text('dict')['blocks'] if b['type']==1]
        assert len(page.get_image_info())==1
        doc.save(source)
    original=source.read_bytes()
    imported=tmp_path/'source.pptx'
    pdf_import.convert(source,imported,lambda i:tmp_path/f'before-{i}.png')
    assert source.read_bytes()==original
    prs=Presentation(imported)
    composed=Image.new('RGBA',(1440,810),'white')
    scale=min(prs.slide_width/914400/720,prs.slide_height/914400/405)
    dx=(prs.slide_width/914400-720*scale)/2
    dy=(prs.slide_height/914400-405*scale)/2
    for shape in prs.slides[1].shapes:
        if shape.shape_type!=13: continue
        image=Image.open(BytesIO(shape.image.blob)).convert('RGBA')
        x=round((shape.left/914400-dx)/scale*2)
        y=round((shape.top/914400-dy)/scale*2)
        assert image.size==(round(shape.width/914400/scale*2),round(shape.height/914400/scale*2))
        composed.alpha_composite(image,(x,y))
    with fitz.open(source) as doc:
        page=doc[1]
        for block in page.get_text('dict')['blocks']:
            if block['type']==0:
                for line in block['lines']: page.add_redact_annot(line['bbox'],fill=False)
        page.apply_redactions(images=0,graphics=0,text=0)
        expected=Image.open(BytesIO(page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).tobytes('png'))).convert('RGB')
    # Proves complete curves, raster label, layer order and original clipping,
    # including the spaces between separate native arrow paths.
    assert ImageChops.difference(composed.convert('RGB'),expected).getbbox() is None


@pytest.mark.parametrize('semantic_pixel',[False,True])
def test_repeated_perimeter_shading_requires_uniform_neutral_rows(tmp_path,semantic_pixel):
    strip=Image.new('RGB',(720,5))
    draw=ImageDraw.Draw(strip)
    for row,gray in enumerate((100,130,160,190,220)):
        draw.line((0,row,719,row),fill=(gray,gray,gray))
    if semantic_pixel: strip.putpixel((350,2),(0,0,0))
    blob=BytesIO();strip.save(blob,format='PNG')
    with fitz.open() as doc:
        for i in range(4):
            page=doc.new_page(width=720,height=405)
            page.insert_image((0,-1,720,4),stream=blob.getvalue())
            page.insert_text((30,60),f'Body {i}',fontsize=20)
        found=pdf_regions.repeated_edge_shadows(doc)
        assert bool(found) is not semantic_pixel
        if found: assert set(found)==set(range(4))


@pytest.mark.parametrize('case',['mask','same-bounds'])
def test_perimeter_shadow_detection_preserves_mask_only_content_and_ambiguous_aliases(case):
    image=Image.new('RGB',(720,5),(120,120,120))
    if case=='mask':
        image=image.convert('RGBA');image.putpixel((350,2),(120,120,120,0))
    blob=BytesIO();image.save(blob,format='PNG')
    other=Image.new('RGB',(720,5),(200,200,200));other.putpixel((350,2),(0,0,0))
    other_blob=BytesIO();other.save(other_blob,format='PNG')
    with fitz.open() as doc:
        for i in range(4):
            page=doc.new_page(width=720,height=405)
            page.insert_image((0,-1,720,4),stream=blob.getvalue())
            if case=='same-bounds': page.insert_image((0,-1,720,4),stream=other_blob.getvalue())
        assert not pdf_regions.repeated_edge_shadows(doc)
