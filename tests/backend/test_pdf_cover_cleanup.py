from io import BytesIO

import fitz
import pytest
from PIL import Image
from pptx import Presentation

from app import grounded, pdf_import
from app.qa.artifact_coverage import audit
from slide_engine import templates


def test_pdf_plot_labels_keep_source_contrast_and_semantic_colors(tmp_path):
    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    page.insert_text((45, 100), 'Optimization', fontsize=32)
    page = doc.new_page(width=720, height=405)
    image = BytesIO()
    Image.new('RGB', (300, 180), '#303030').save(image, format='PNG')
    page.insert_image((90, 100, 390, 280), stream=image.getvalue())
    page.insert_text((120, 145), 'AdaGrad', fontsize=18, color=(.9, .9, .9))
    page.insert_text((120, 185), 'Gradient Descent', fontsize=18, color=(0, .8, 1))
    page.insert_text((120, 225), 'Threshold', fontsize=18, color=(1, 0, 0))
    source_pdf = tmp_path / 'source.pdf'
    doc.save(source_pdf)
    doc.close()
    imported = tmp_path / 'source.pptx'
    pdf_import.convert(source_pdf, imported, lambda i: str(tmp_path / f'{i}.png'))
    before = Presentation(imported)
    expected = {s.text: str(s.text_frame.paragraphs[0].runs[0].font.color.rgb)
                for s in before.slides[1].shapes if s.has_text_frame}
    assert set(expected) == {'AdaGrad', 'Gradient Descent', 'Threshold'}
    with templates.use('cpe'):
        candidate = tmp_path / 'candidate.pptx'
        report = grounded.build_deck(imported, candidate)
        actual = {s.text: str(s.text_frame.paragraphs[0].runs[0].font.color.rgb)
                  for s in Presentation(candidate).slides[1].shapes if s.has_text_frame}
        assert actual == expected
        assert audit(imported, candidate, report)['status'] == 'passed'


def layered_cover():
    from PIL import ImageDraw
    doc = fitz.open()
    page = doc.new_page(width=720, height=405)
    background = BytesIO()
    Image.new('RGB', (720, 405), '#155597').save(background, format='PNG')
    page.insert_image(page.rect, stream=background.getvalue())
    logo = Image.new('RGBA', (200, 50), (0, 0, 0, 0))
    draw = ImageDraw.Draw(logo)
    draw.rectangle((5, 5, 194, 44), outline='white', width=4)
    draw.text((15, 15), 'REQUIRED SOURCE LOGO', fill='white')
    image = BytesIO()
    logo.save(image, format='PNG')
    page.insert_image((25, 30, 225, 80), stream=image.getvalue())
    page.draw_rect((25, 230, 490, 290), color=None, fill=(.1, .1, .1), fill_opacity=.5)
    page.insert_text((40, 275), 'Optimization', fontsize=32, color=(1, 1, 1))
    page.insert_text((40, 345), 'Course details', fontsize=18, color=(1, 1, 1))
    return doc


def test_cover_layers_isolate_full_logo_from_background_without_source_changes():
    from app.pdf_cover import graphics
    doc = layered_cover()
    page = doc[0]
    original = page.get_pixmap().tobytes('png')
    layers = graphics(page, page.get_text('dict')['blocks'], page.get_drawings(), [])
    assert len(layers) == 3
    photo, logo, panel = [Image.open(BytesIO(layer['blob'])).convert('RGBA') for layer in layers]
    assert photo.getpixel((100, 100))[:3] == (21, 85, 151)
    assert logo.getchannel('A').getextrema() == (0, 255)
    assert all(abs(r-g) <= 1 and abs(g-b) <= 1 for r, g, b, a in logo.getdata() if a > 0)
    assert 120 <= panel.getchannel('A').getextrema()[1] <= 130
    assert page.get_pixmap().tobytes('png') == original
    page.draw_line((10, 100), (350, 150), color=(1, 0, 0))
    assert graphics(page, page.get_text('dict')['blocks'], page.get_drawings(), []) is None
    doc.close()


def test_pdf_import_separates_cover_after_repeated_line_primitive_frame_cleanup(tmp_path):
    from types import SimpleNamespace
    from app.ai.cover_preparation import needed
    from app.pdf_cover import graphics

    doc = layered_cover()
    for i in range(2):
        page = doc.new_page(width=720, height=405)
        page.insert_text((40, 80), f'Content page {i+1}', fontsize=22)
    for page in doc:
        # Office PDFs encode this rectangle with move/line/close operators,
        # although MuPDF exposes it as an 're' drawing. Exercise actual import.
        page.draw_polyline([(0, .378), (720, .378), (720, 404.622), (0, 404.622)],
                           closePath=True, color=(.702, .106, .106), width=6)
    source = tmp_path/'source.pdf'
    doc.save(source)
    doc.close()
    imported = tmp_path/'source.pptx'
    with templates.use('cpe'):
        receipt = pdf_import.convert(source, imported, lambda i: str(tmp_path/f'{i}.png'))
    assert all(page['excluded_chrome']['removed_page_frame_paths'] == 1
               for page in receipt['page_evidence'])
    assert len(receipt['page_evidence'][0]['independent_cover_graphics']) == 3
    assert needed(SimpleNamespace(pdf_import=receipt))
    pictures = [shape for shape in Presentation(imported).slides[0].shapes if shape.shape_type == 13]
    assert len(pictures) == 3 and all('independent cover' in picture.name for picture in pictures)
    logo = Image.open(BytesIO(pictures[1].image.blob)).convert('RGBA')
    assert logo.getchannel('A').getextrema() == (0, 255)
    assert all(abs(r-g) <= 1 and abs(g-b) <= 1 for r, g, b, a in logo.getdata() if a > 0)


def test_cover_glyph_fallback_does_not_duplicate_independent_layers(tmp_path, monkeypatch):
    from app import pdf_cover
    doc = layered_cover()
    source = tmp_path/'source.pdf'
    doc.save(source)
    doc.close()
    # Emulate an embedded PDF font whose title maps to XML-illegal glyphs.
    monkeypatch.setattr(pdf_import, 'needs_glyph_image',
                        lambda line: any('Optimization' in span['text'] for span in line['spans']))
    monkeypatch.setattr(pdf_cover, 'graphics',
                        lambda *args: (_ for _ in ()).throw(AssertionError('Must use composited fallback')))
    receipt = pdf_import.convert(source, tmp_path/'source.pptx', lambda i: str(tmp_path/f'{i}.png'))
    page = receipt['page_evidence'][0]
    assert page['rasterized_text_lines'] == 1
    assert page['independent_cover_graphics'] == []


@pytest.mark.parametrize('remove_photo', [True, False])
def test_native_cover_preparation_reviews_only_cover_and_retains_full_logo(tmp_path, monkeypatch, remove_photo):
    from types import SimpleNamespace
    from app.ai import cover_preparation, providers
    from app.pdf_cover import graphics
    from pptx.util import Inches, Pt
    from rubric_fixtures import source_choice
    from test_ai_pipeline import fake_render
    from slide_engine import template_policy as T

    doc = layered_cover()
    page = doc[0]
    layers = graphics(page, page.get_text('dict')['blocks'], page.get_drawings(), [])
    prs = Presentation()
    prs.slide_width, prs.slide_height = T.CANVAS
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide._element.cSld.set('name', 'sss:pdf-import')
    for i, layer in enumerate(layers):
        rect = layer['bounds']
        slide.shapes.add_picture(BytesIO(layer['blob']), *[Inches(v/54) for v in
                                 (rect.x0, rect.y0, rect.width, rect.height)])
    for y, text, size in ((5, 'Optimization', 32), (6, 'Course details', 18)):
        shape = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(8), Inches(.6))
        shape.text = text
        shape.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    # Content pages must not trigger another source-classification call.
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Required body'
    source = tmp_path/'source.pptx'
    prs.save(source)
    page.get_pixmap().save(tmp_path/'before-0.png')
    sess = SimpleNamespace(source_path=source, dir=str(tmp_path), revisions={},
                           pdf_import={'page_evidence': [{'independent_cover_graphics': [l['evidence'] for l in layers]}]},
                           ensure_active=lambda: None, preview_path=lambda i, stage: str(tmp_path/f'before-{i}.png'))
    calls = []

    def generate(role, system, payload, images, max_tokens):
        calls.append(payload['source_slide'])
        assert len([label for label, _ in images if label.startswith('FULL RAW IMAGE')]) == 3
        result = source_choice(payload)
        pictures = [o for o in payload['objects'] if o['kind'] == 'picture']
        for element in result['elements']:
            if element['id'] == pictures[1]['id']:
                element.update(role='logo', confidence='high', contains_logo=True,
                               content_bearing=True, artwork_action='retain')
            elif element['id'] == pictures[0]['id'] and not remove_photo:
                element.update(role='image', confidence='high', contains_logo=False,
                               content_bearing=True, artwork_action='retain')
            elif element['id'] in {pictures[0]['id'], pictures[2]['id']}:
                element.update(role='panel' if element['id']==pictures[2]['id'] else 'background', confidence='high', contains_logo=False,
                               content_bearing=False, artwork_action='remove')
                result['remove_ids'].append(element['id'])
            elif element['role'] in ('title', 'subtitle'):
                # Source contrast dependencies are replaced by the approved
                # cover field when these native text objects are transferred.
                element['related_ids'] = [pictures[2]['id']]
        return {'status': 'completed', 'data': result}

    monkeypatch.setattr(providers, 'role_config', lambda role: {'configured': True, 'provider': 'mock', 'model': 'fixture'})
    monkeypatch.setattr(providers, 'generate', generate)
    monkeypatch.setattr(cover_preparation.render_verify, 'check', fake_render)
    with templates.use('cpe'):
        decisions, details = cover_preparation.run(sess, tmp_path, lambda **kwargs: None)
        assert calls == [0] and details['status'] == 'completed', details
        assert (tmp_path/'cover-call-001.json').is_file()
        candidate = tmp_path/'candidate.pptx'
        report = grounded.build_deck(source, candidate, source_decisions=decisions)
        out = Presentation(candidate)
        pictures = [s for s in out.slides[0].shapes if s.shape_type == 13]
        expected = [layers[1]['blob']] if remove_photo else [layers[0]['blob'], layers[1]['blob']]
        assert [picture.image.blob for picture in pictures] == expected
        if remove_photo:
            assert pictures[0].width > Inches(6)  # Do not shrink a logo against the removed canvas.
        assert all(T.contains(tuple(v/914400 for v in (p.left, p.top, p.width, p.height)),
                              T.COVER_SUPPORT) for p in pictures)
        assert audit(source, candidate, report)['status'] == 'passed'
        assert T.check(candidate)['status'] == 'passed'
    doc.close()


def test_cover_preparation_offline_retains_artwork_without_calling_ai(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from app.ai import cover_preparation, providers
    sess = SimpleNamespace(pdf_import={'page_evidence': [{'independent_cover_graphics': [{}]}]})
    monkeypatch.setattr(providers, 'role_config', lambda role: {'configured': False})
    monkeypatch.setattr(providers, 'generate', lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError('Offline AI call')))
    decisions, details = cover_preparation.run(sess, tmp_path, lambda **kwargs: None)
    assert decisions == {} and details['status'] == 'not_configured'
