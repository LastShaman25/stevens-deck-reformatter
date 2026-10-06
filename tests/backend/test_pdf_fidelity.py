from copy import deepcopy
from pathlib import Path
import hashlib
import fitz
import pytest
from pptx import Presentation
from pptx.util import Inches
from app import pdf_import, pdf_regions, grounded
from app.ai.output_qa import consolidate_findings
from slide_engine import templates


def test_small_pdf_code_grows_without_colliding_with_next_line():
    from pptx.util import Pt
    from slide_engine.pdf_text_fit import enlarge_small_text
    prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6])
    code=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(.12))
    code.text="(x+dx)'*(x+dx) - x'x"
    code.text_frame.paragraphs[0].runs[0].font.size=Pt(7)
    note=slide.shapes.add_textbox(Inches(1),Inches(1.14),Inches(5),Inches(.3))
    note.text='Note: preserve this explanation.'
    note.text_frame.paragraphs[0].runs[0].font.size=Pt(18)
    enlarge_small_text(slide,[],(.7,.4,11.7,6.05))
    assert code.text=="(x+dx)'*(x+dx) - x'x"
    assert code.text_frame.paragraphs[0].runs[0].font.size.pt==11
    assert code.top+code.height < note.top
    assert note.text=='Note: preserve this explanation.'


def test_region_closure_keeps_equation_neighbors_once():
    lines = [{'bbox':(10,10,40,25)}, {'bbox':(35,20,60,35)}, {'bbox':(100,10,140,25)}]
    regions = pdf_regions.fidelity_regions(lines, [(55,30,70,40)], [lines[0],lines[1]])
    assert [tuple(r) for r in regions] == [(10,10,70,40)]
    assert not regions[0].intersects(fitz.Rect(lines[2]['bbox']))


def test_exact_path_cleanup_keeps_overlapping_artwork_and_literal_text():
    from app.pdf_paths import remove_paths
    doc=fitz.open(); page=doc.new_page(width=600,height=400)
    page.draw_rect((20,20,580,395),fill=(1,1,1),color=None)
    page.draw_rect((450,370,460,380),color=(.5,.5,.8))
    page.insert_text((30,380),'Literal commands: (q 450 20 10 10 re S Q)',fontsize=12)
    before=page.get_text(); drawings=page.get_drawings(); target=[drawings[1]]
    assert remove_paths(page,target)==1
    assert page.get_text()==before
    after=page.get_drawings()
    assert len(after)==1 and after[0]['items']==drawings[0]['items']
    doc.close()


def test_repeated_footer_does_not_discard_unique_footnote(tmp_path):
    doc = fitz.open()
    for i in range(4):
        page = doc.new_page(width=600,height=400)
        page.insert_text((30,60),'Running presentation title',fontsize=20)
        page.insert_text((30,390),'Running presentation title',fontsize=8)
        page.insert_text((250,390),f'Unique citation {i}',fontsize=8)
        page.insert_text((30,380),'Copyright 2022 Example Institute',fontsize=8)
    pdf=tmp_path/'source.pdf'; doc.save(pdf); doc.close()
    receipt = pdf_import.convert(pdf,tmp_path/'out.pptx',lambda i:str(tmp_path/f'{i}.png'))
    for i, slide in enumerate(Presentation(tmp_path/'out.pptx').slides):
        text = '\n'.join(s.text for s in slide.shapes if s.has_text_frame)
        assert text.count('Running presentation title') == 1
        assert f'Unique citation {i}' in text
        assert 'Copyright 2022 Example Institute' in text
        assert receipt['page_evidence'][i]['excluded_chrome']['navigation_links'] == 0


def test_title_backdrop_preserves_overlapping_source_picture():
    from io import BytesIO
    from PIL import Image
    doc = fitz.open(); page = doc.new_page(width=600,height=400)
    line = {'bbox':(100,80,400,120),'spans':[{'text':'Title','size':24}]}
    drawing = {'rect':(80,60,420,140),'fill':(0,0,.7)}
    blob=BytesIO(); Image.new('RGB',(150,80),(20,100,180)).save(blob,format='PNG')
    picture={'bbox':(80,60,170,100),'width':150,'height':80,'image':blob.getvalue()}
    assert pdf_regions.title_backdrop(page,[line],[line],[drawing],[picture]) == ([],[])
    assert pdf_regions.title_backdrop(page,[line],[line],[drawing],[]) == ([drawing],[])
    doc.close()


def test_qa_repeated_observations_keep_strongest_verdict_and_evidence():
    item = dict(criterion='content_presence',slides=[1],object_ids=['title'],region='cover',
                severity='review',accuracy='supported',message='Title repeats.',evidence='First evidence',
                required_correction='Remove the extra author and title text, retaining the title and date once each.',
                acceptance_condition='No repetition',category='accuracy')
    repeat = {**item,'severity':'blocking','accuracy':'contradicted','message':'Cover has duplicate title.',
              'region':'cover title area',
              'evidence':'Second evidence',
              'required_correction':'Remove extra author and title text, retaining the cover title and date once each.'}
    other = {**item,'required_correction':'Restore the missing word in the title.', 'message':'Word omitted.'}
    original=deepcopy([item,repeat,other])
    result=consolidate_findings([item,repeat,other])
    assert len(result)==2
    assert result[0]['severity']=='blocking' and result[0]['accuracy']=='contradicted'
    assert len(result[0]['observations'])==2
    assert [item,repeat,other]==original
    assert len(consolidate_findings([item,{**repeat,'slides':[2]}]))==2
    a={**item,'message':'The title is legible, but its line wrapping separates a proxy from the phrase.',
       'required_correction':'Adjust wrapping while preserving readability.'}
    b={**item,'message':'The title is legible, but its line wrapping isolates a proxy on the middle line.',
       'required_correction':'Keep the short phrase with the surrounding words.'}
    assert len(consolidate_findings([a,b]))==1
    assert len(consolidate_findings([a,other]))==2


@pytest.mark.parametrize('template',['stevens','cpe'])
def test_actual_columbia_cover_and_source_font_pixels(tmp_path,template):
    source = Path(__file__).resolve().parents[2]/'test_asset/Columbia_talk.pdf'
    if not source.exists(): pytest.skip('User regression asset is not distributed with the repository')
    imported=tmp_path/'source.pptx'; candidate=tmp_path/'candidate.pptx'
    with templates.use(template):
        receipt=pdf_import.convert(source,imported,lambda i:str(tmp_path/f'{i}.png'))
        prs=Presentation(imported)
        cover=prs.slides[0]
        assert len(cover.shapes)==3
        assert cover.shapes[0].text == 'Gambling under unknown probabilities as a proxy\nfor real world decisions under uncertainty'
        assert [s.text for s in cover.shapes][1:]==['David Aldous','6 October 2022']
        assert receipt['page_evidence'][0]['excluded_chrome']['title_shadow_images']==5
        # Recompute original source pixels for every preserved region, rather
        # than accepting an import receipt's own claim of fidelity.
        with fitz.open(source) as pdf:
            for i, page in enumerate(receipt['page_evidence']):
                original_page = pdf[i]
                source_drawings = original_page.get_drawings()
                pictures=[s for s in prs.slides[i].shapes if 'source-font region' in s.name]
                assert len(pictures)==page['source_font_regions']
                cleaned=fitz.open(); cleaned.insert_pdf(pdf,from_page=i,to_page=i)
                clean_page=cleaned[0]
                from app.pdf_paths import remove_paths, path_bounds
                nav=[{'rect':box} for box in page['excluded_chrome']['navigation_path_bounds']]
                assert page['excluded_chrome']['removed_native_paths'] == len(nav)
                remove_paths(clean_page,nav)
                assert clean_page.get_text() == original_page.get_text()
                # Every original non-navigation drawing survives. This protects
                # article artwork crossing the footer against area redaction.
                remaining=clean_page.get_drawings()
                for drawing in source_drawings:
                    if any(all(abs(a-b)<.025 for a,b in zip(path_bounds(drawing),d['rect'])) for d in nav): continue
                    assert any(drawing['items']==d['items'] for d in remaining)
                for picture, evidence in zip(pictures,page['source_font_evidence']):
                    assert evidence['bounds'][3] <= page['visible_content_bounds'][3] + 1/3 + .001
                    pix=clean_page.get_pixmap(matrix=fitz.Matrix(3,3),clip=fitz.Rect(evidence['bounds']),alpha=False)
                    assert picture.image.blob==pix.tobytes('png')
                    assert hashlib.sha256(picture.image.blob).hexdigest()==evidence['image_sha256']
                cleaned.close()
        grounded.build_deck(str(imported),str(candidate),require_closing=True)
        output=Presentation(candidate)
        assert len(output.slides)==32
        title=next(s for s in output.slides[0].shapes if s.name.endswith('|title'))
        details=[s for s in output.slides[0].shapes if s.has_text_frame and not s.name.endswith('|title')]
        assert title.top+title.height <= min(s.top for s in details)
        sizes={r.font.size.pt for p in title.text_frame.paragraphs for r in p.runs}
        assert 20 <= min(sizes) <= max(sizes) < 40
        assert title.width > Inches(7)
