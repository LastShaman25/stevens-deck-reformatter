"""Repeated source decoration must not join diagrams into an old-slide inset."""
from io import BytesIO
import hashlib

import fitz
import pytest
from PIL import Image
from pptx import Presentation
from app import pdf_import, pdf_regions
from app.pdf_paths import remove_paths


def source_deck(path):
    with fitz.open() as doc:
        for i in range(4):
            page=doc.new_page(width=720,height=405)
            page.draw_rect(page.rect,color=None,fill=(1,1,1))
            # An intervening content path keeps MuPDF from coalescing the two
            # independent same-bounds fill/stroke operations in its inventory.
            page.draw_circle((140,150),30,color=(0,0,0))
            # Identical bounds to the white fill: only the stroke may disappear.
            page.draw_rect(page.rect,color=(.7,.1,.1),width=4)
            if i!=2: page.insert_text((30,65),f'Heading {i+1}',fontsize=24)
            # Meaningful diagram outlines/arrows must survive even when red.
            page.draw_rect((75,100,590,320),color=(.7,.1,.1),width=3)
            page.draw_line((170,150),(410,220),color=(1,0,0),width=2)
            picture=BytesIO(); Image.new('RGB',(70,70),(40,80,120)).save(picture,format='PNG')
            page.insert_image((440,180,490,230),stream=picture.getvalue())
        doc.save(path)


def test_perimeter_cleanup_keeps_same_bounds_fill_diagrams_images_and_text(tmp_path):
    source=tmp_path/'frames.pdf'; source_deck(source)
    before=source.read_bytes()
    with fitz.open(source) as doc:
        keys=pdf_regions.repeated_frames(doc)
        assert len(keys)==1
        for page in doc:
            original=page.get_drawings(); text=page.get_text(); images=page.get_images()
            border=[d for d in original if pdf_regions.frame_key(d,page.rect) in keys]
            assert len(border)==1
            assert remove_paths(page,border,stroke_only=True)==1
            remaining=page.get_drawings()
            assert len(remaining)==len(original)-1
            assert [d['items'] for d in remaining]==[d['items'] for d in original if d not in border]
            assert page.get_text()==text and page.get_images()==images
    assert source.read_bytes()==before


@pytest.mark.parametrize('font_fallback',[False,True])
def test_import_removes_frames_before_graphic_and_font_merging(tmp_path,monkeypatch,font_fallback):
    source=tmp_path/'frames.pdf'; source_deck(source)
    if font_fallback: monkeypatch.setattr(pdf_regions,'needs_source_font',lambda line:True)
    imported=tmp_path/'imported.pptx'
    receipt=pdf_import.convert(source,imported,lambda i:tmp_path/f'{i}.png')
    prs=Presentation(imported)
    assert len(prs.slides)==4
    for page in receipt['page_evidence']:
        assert page['excluded_chrome']['page_frames']==1
        assert page['excluded_chrome']['removed_page_frame_paths']==1
    for i in (0,1,3):
        # The old perimeter no longer creates a canvas-sized content picture.
        assert all(s.width < prs.slide_width*.95 for s in prs.slides[i].shapes)
    # Textless page still keeps every original pixel except the identified frame.
    with fitz.open(source) as doc:
        page=doc[2]
        border=[d for d in page.get_drawings() if pdf_regions.frame_key(d,page.rect) is not None]
        remove_paths(page,border,stroke_only=True)
        expected=page.get_pixmap(matrix=fitz.Matrix(3,3),alpha=False).tobytes('png')
    assert prs.slides[2].shapes[0].image.blob==expected
    assert receipt['page_evidence'][2]['page_image_evidence']['image_sha256']==hashlib.sha256(expected).hexdigest()


def test_nonrepeated_or_content_panel_borders_are_not_decorative(tmp_path):
    with fitz.open() as doc:
        for i in range(4):
            page=doc.new_page(width=720,height=405)
            page.draw_rect((70,90,650,360),color=(1,0,0),width=3)
            if i==0: page.draw_rect(page.rect,color=(1,0,0),width=3)
        assert not pdf_regions.repeated_frames(doc)
    filled={'items':[('re',fitz.Rect(0,0,720,405),1)],'rect':(0,0,720,405),
            'color':(1,0,0),'stroke_opacity':1,'width':3,'fill':(.1,.2,.3),'fill_opacity':1}
    assert pdf_regions.frame_key(filled,fitz.Rect(0,0,720,405)) is None


@pytest.mark.parametrize('content_path',[
    b'0 0 m 720 405 l S',                                     # Diagonal.
    b'0 0 m 720 0 l 720 405 l 0 405 l S',                    # Open three sides.
    b'0 202.5 m 360 0 l 720 202.5 l 360 405 l h S',           # Closed diamond.
    b'0 0 m 720 405 l 720 0 l 0 405 l h S',                  # Crossing bowtie.
    b'0 0 720 405 re 80 80 m 140 140 l S',                  # Compound content.
])
def test_embedded_frame_does_not_match_meaningful_root_path_with_same_bounds(content_path):
    with fitz.open() as embedded, fitz.open() as doc:
        inner=embedded.new_page(width=720,height=405)
        inner.draw_rect(inner.rect,color=(1,0,0),width=4)
        page=doc.new_page(width=720,height=405)
        page.show_pdf_page(page.rect,embedded,0)
        # This content has identical bounds to the old frame, but is neither
        # an independent rectangle nor part of its embedded content stream.
        # Keep the embedded Form reference and append root-level content.
        first=page.get_contents()[0]
        doc.update_stream(first,doc.xref_stream(first)+b'\nq 0 0 0 RG 2 w '+content_path+b' Q')
        before=page.get_drawings()
        frames=[d for d in before if pdf_regions.frame_key(d,page.rect) is not None]
        assert len(frames)==1
        assert remove_paths(page,frames,stroke_only=True)==0
        assert page.get_drawings()==before


@pytest.mark.parametrize('rectangle',[
    b'0 0 m 720 0 l 720 405 l 0 405 l h S',
    b'0 0 m 720 0 l 720 405 l 0 405 l s',
])
def test_polygon_encoded_perimeter_is_removed_like_native_rectangle(rectangle):
    with fitz.open() as doc:
        page=doc.new_page(width=720,height=405)
        page.insert_text((40,60),'Retain this source text')
        first=page.get_contents()[0]
        doc.update_stream(first,doc.xref_stream(first)+b'\nq 1 0 0 RG 4 w '+rectangle+b' Q')
        before=page.get_drawings(); text=page.get_text()
        frames=[d for d in before if pdf_regions.frame_key(d,page.rect) is not None]
        assert len(frames)==1
        assert remove_paths(page,frames,stroke_only=True)==1
        assert page.get_drawings()==[]
        assert page.get_text()==text


def test_embedded_frame_never_removes_same_bounds_filled_panel_outline():
    with fitz.open() as embedded, fitz.open() as doc:
        inner=embedded.new_page(width=720,height=405)
        inner.draw_rect(inner.rect,color=(1,0,0),width=4)
        page=doc.new_page(width=720,height=405)
        page.show_pdf_page(page.rect,embedded,0)
        page.draw_rect(page.rect,color=(0,0,0),fill=(.8,.8,.8),width=2)
        before=page.get_drawings()
        frames=[d for d in before if pdf_regions.frame_key(d,page.rect) is not None]
        assert len(frames)==1 and len(before)==2
        assert remove_paths(page,frames,stroke_only=True)==0
        assert page.get_drawings()==before


def test_glyph_fallback_cannot_reintroduce_removed_perimeter(tmp_path,monkeypatch):
    source=tmp_path/'glyph-frames.pdf'
    with fitz.open() as doc:
        for i in range(4):
            page=doc.new_page(width=720,height=405)
            page.draw_rect(page.rect,color=None,fill=(1,1,1))
            page.draw_circle((140,150),30,color=(0,0,0))
            page.draw_rect(page.rect,color=(.7,.1,.1),width=4)
            # The glyph raster region overlaps the actual perimeter stroke.
            page.insert_text((1,20),f'Heading {i+1}',fontsize=14)
        doc.save(source)
    monkeypatch.setattr(pdf_import,'needs_glyph_image',lambda line:True)
    monkeypatch.setattr(pdf_regions,'needs_source_font',lambda line:False)
    imported=tmp_path/'glyph-imported.pptx'
    receipt=pdf_import.convert(source,imported,lambda i:tmp_path/f'glyph-{i}.png')
    prs=Presentation(imported)
    with fitz.open(source) as doc:
        for i,page in enumerate(doc):
            line=pdf_regions.lines(page.get_text('dict')['blocks'])[0]
            frames=[d for d in page.get_drawings() if pdf_regions.frame_key(d,page.rect) is not None]
            rect=fitz.Rect(line['bbox']) & page.rect
            original=page.get_pixmap(matrix=fitz.Matrix(3,3),clip=rect,alpha=True).tobytes('png')
            assert remove_paths(page,frames,stroke_only=True)==1
            expected=page.get_pixmap(matrix=fitz.Matrix(3,3),clip=rect,alpha=True).tobytes('png')
            assert expected!=original  # The negative case really includes frame pixels.
            glyph=next(s for s in prs.slides[i].shapes if 'preserved text glyphs' in s.name)
            assert glyph.image.blob==expected
            assert receipt['page_evidence'][i]['excluded_chrome']['removed_page_frame_paths']==1
