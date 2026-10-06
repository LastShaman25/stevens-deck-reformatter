"""Visible PDF content must survive even when the PDF has no text layer."""
from io import BytesIO
import hashlib

import fitz
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from fastapi.testclient import TestClient
from app import pdf_import, grounded, sessions
from app.main import app
from app.qa.artifact_coverage import audit
from slide_engine import templates, template_policy as T


def image_page(doc, kind):
    page=doc.new_page(width=720,height=405)
    if kind=='raster':
        image=Image.new('RGB',(1440,810),'white')
        draw=ImageDraw.Draw(image)
        draw.ellipse((100,200,300,400),outline='black',width=5)
        draw.ellipse((700,200,900,400),outline='black',width=5)
        draw.line((300,300,700,300),fill='red',width=8)
        draw.text((150,275),'x1',fill='black')
        draw.text((740,275),'z1',fill='black')
        blob=BytesIO(); image.save(blob,format='PNG')
        page.insert_image(page.rect,stream=blob.getvalue())
    elif kind=='vector':
        # Outlined labels are drawing paths, just like outlined PDF font glyphs.
        page.draw_circle((160,160),40,color=(0,0,0))
        page.draw_circle((420,160),40,color=(0,0,0))
        page.draw_line((200,160),(380,160),color=(1,0,0),width=3)
        page.draw_line((150,150),(170,170),color=(0,0,0),width=2)
        page.draw_line((170,150),(150,170),color=(0,0,0),width=2)
    return page


@pytest.mark.parametrize('template',['stevens','cpe'])
@pytest.mark.parametrize('kind',['raster','vector','blank'])
def test_textless_page_pixels_links_and_order_survive_reformatting(tmp_path,template,kind):
    source=tmp_path/'mixed.pdf'; imported=tmp_path/'imported.pptx'; output=tmp_path/'output.pptx'
    with fitz.open() as doc:
        doc.new_page(width=720,height=405).insert_text((50,80),'Editable opening',fontsize=24)
        image_page(doc,kind)
        doc.new_page(width=720,height=405).insert_text((50,80),'Text after the figure',fontsize=24)
        doc[1].insert_link({'kind':fitz.LINK_GOTO,'from':fitz.Rect(100,100,130,120),'page':2})
        doc[1].insert_link({'kind':fitz.LINK_URI,'from':fitz.Rect(300,100,330,120),'uri':'https://example.com/'})
        doc.save(source)
    before=source.read_bytes()
    with templates.use(template):
        receipt=pdf_import.convert(source,imported,lambda i:tmp_path/f'preview-{i}.png')
        prs=Presentation(imported)
        assert len(prs.slides)==3
        assert any(s.has_text_frame and s.text=='Editable opening' for s in prs.slides[0].shapes)
        assert any(s.has_text_frame and s.text=='Text after the figure' for s in prs.slides[2].shapes)
        pictures=[s for s in prs.slides[1].shapes if s.shape_type==13]
        assert len(pictures)==1
        with fitz.open(source) as doc:
            assert not doc[1].get_text().strip()
            expected=doc[1].get_pixmap(matrix=fitz.Matrix(3,3),alpha=False).tobytes('png')
        assert pictures[0].image.blob==expected
        evidence=receipt['page_evidence'][1]
        assert evidence['page_image_evidence']['image_sha256']==hashlib.sha256(expected).hexdigest()
        assert evidence['page_image_regions']==1
        assert 'No OCR was performed' in evidence['preservation']
        report=grounded.build_deck(imported,output,require_closing=True)
        candidate=Presentation(output)
        assert len(candidate.slides)==4
        picture=next(s for s in candidate.slides[1].shapes if s.shape_type==13)
        assert picture.image.blob==expected
        assert abs(picture.width/picture.height-720/405)<.001
        assert T.contains([v/T.EMU for v in (picture.left,picture.top,picture.width,picture.height)],T.CONTENT)
        assert any(s.click_action.target_slide==candidate.slides[2] for s in candidate.slides[1].shapes)
        assert any(s.click_action.hyperlink.address=='https://example.com/' for s in candidate.slides[1].shapes)
        assert audit(imported,output,report)['status']=='passed'
    assert source.read_bytes()==before


def test_api_accepts_mixed_deck_with_textless_page_30():
    with fitz.open() as doc:
        for i in range(29):
            doc.new_page(width=720,height=405).insert_text((50,80),f'Page {i+1}',fontsize=24)
        image_page(doc,'vector')
        pdf=doc.tobytes()
    response=TestClient(app).post('/api/sessions',files={'file':('diagram-deck.pdf',pdf,'application/pdf')})
    assert response.status_code==200,response.text
    session=sessions.get(response.json()['session_id'])
    try:
        assert response.json()['slide_count']==30
        assert session.pdf_import['page_evidence'][29]['page_image_regions']==1
    finally: sessions.delete(session.id)
