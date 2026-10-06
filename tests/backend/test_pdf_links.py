"""PDF navigation must survive import, including Beamer named destinations."""
import fitz
import pytest
from pptx import Presentation

from app.pdf_import import convert


def linked_pdf(path, actions):
    with fitz.open() as doc:
        for i in range(2):
            doc.new_page(width=720, height=405).insert_text((30, 60), f'Page {i+1}')
        # Use actual PDF named GoTo destinations, as emitted by LaTeX Beamer.
        doc.xref_set_key(doc.pdf_catalog(), 'Names',
                        f'<< /Dests << /Names [(Navigation2) [{doc[1].xref} 0 R /XYZ 0 405 null]] >> >>')
        refs=[]
        for action in actions:
            xref=doc.get_new_xref()
            action=action.replace('TARGET_PAGE', str(doc[1].xref))
            doc.update_object(xref, f'<< /Type /Annot /Subtype /Link /Rect [30 300 100 330] /A << {action} >> >>')
            refs.append(f'{xref} 0 R')
        doc.xref_set_key(doc[0].xref, 'Annots', f'[{" ".join(refs)}]')
        doc.save(path)


def test_named_direct_and_web_links_survive_saved_presentation(tmp_path):
    source=tmp_path/'links.pdf'
    output=tmp_path/'links.pptx'
    linked_pdf(source, [
        '/S /GoTo /D (Navigation2)',
        '/S /GoTo /D [TARGET_PAGE 0 R /XYZ 0 405 null]',
        '/S /Named /N /NextPage',
        '/S /URI /URI (https://example.com/)',
    ])
    with fitz.open(source) as doc:
        assert [link['kind'] for link in doc[0].get_links()]==[fitz.LINK_NAMED, fitz.LINK_GOTO, fitz.LINK_GOTO, fitz.LINK_URI]
    evidence=convert(source, output, lambda i: str(tmp_path/f'{i}.png'))
    prs=Presentation(output)
    internal=[s for s in prs.slides[0].shapes if s.name.startswith('pdf-internal-link:')]
    assert len(internal)==3
    assert all(s.click_action.target_slide==prs.slides[1] for s in internal)
    assert [s.click_action.hyperlink.address for s in prs.slides[0].shapes
            if not s.name.startswith('pdf-internal-link:') and s.click_action.hyperlink.address]==['https://example.com/']
    assert evidence['page_evidence'][0]['links']==4


@pytest.mark.parametrize('action', [
    '/S /GoTo /D (MissingDestination)',
])
def test_unresolved_named_destinations_fail_explicitly(tmp_path, action):
    source=tmp_path/'unresolved.pdf'
    linked_pdf(source, [action])
    with pytest.raises(ValueError, match='internal link whose target page could not be resolved'):
        convert(source, tmp_path/'out.pptx', lambda i: str(tmp_path/f'{i}.png'))


@pytest.mark.parametrize('action', [
    '/S /Launch /F (other.pdf)',
    '/S /GoToR /F (other.pdf) /D [0 /Fit]',
])
def test_external_file_actions_remain_unsupported(tmp_path, action):
    source=tmp_path/'external.pdf'
    linked_pdf(source, [action])
    with pytest.raises(ValueError, match='unsupported link type'):
        convert(source, tmp_path/'out.pptx', lambda i: str(tmp_path/f'{i}.png'))
