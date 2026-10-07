"""PDF line boxes must fit the metrics of the standardized native font."""
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from slide_engine.pdf_text_fit import fit_native_lines, _line_width


def subtitle():
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    shape=slide.shapes.add_textbox(Inches(1.06),Inches(4.48),Inches(11.0762),Inches(.5549))
    tf=shape.text_frame;tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    parts=[('Different Learning Rate ',True),('for ',False),('each element ',True),
           ('of the',False),(' Model Weights!',True)]
    for text,bold in parts:
        run=tf.paragraphs[0].add_run();run.text=text
        run.font.name='Arial';run.font.size=Pt(30.32)
        run.font.bold=bold;run.font.italic=True
        run.font.color.rgb=RGBColor.from_string('000000')
    return prs,slide,shape


def test_wider_replacement_font_fits_full_line_without_changing_copy_or_emphasis():
    _,slide,shape=subtitle()
    text=shape.text;bounds=(shape.left,shape.top,shape.width,shape.height)
    runs=list(shape.text_frame.paragraphs[0].runs)
    styling=[(r.text,r.font.bold,r.font.italic,str(r.font.color.rgb)) for r in runs]
    assert _line_width(runs)/4>shape.width/12700
    fit_native_lines(slide)
    assert shape.text==text and (shape.left,shape.top,shape.width,shape.height)==bounds
    assert [(r.text,r.font.bold,r.font.italic,str(r.font.color.rgb)) for r in runs]==styling
    assert all(11<=r.font.size.pt<30.32 for r in runs)
    assert _line_width(runs)/4<shape.width/12700*.98
    sizes=[r.font.size for r in runs]
    fit_native_lines(slide)
    assert [r.font.size for r in runs]==sizes


@pytest.mark.parametrize('case',['wrapped','code','logo','unknown-font','tiny-box','multi-paragraph','formula-baseline'])
def test_uncertain_or_unreadably_small_fitting_is_left_to_the_verifier(case):
    _,slide,shape=subtitle()
    if case=='wrapped': shape.text_frame.word_wrap=True
    elif case in ('code','logo'): shape.name='source|'+case
    elif case=='unknown-font': shape.text_frame.paragraphs[0].runs[0].font.name='UninstalledFont'
    elif case=='tiny-box': shape.width=Inches(1)
    elif case=='multi-paragraph': shape.text_frame.add_paragraph().text='Retain the second paragraph.'
    elif case=='formula-baseline': shape.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr().set('baseline','30000')
    before=shape._element.xml
    fit_native_lines(slide)
    assert shape._element.xml==before
