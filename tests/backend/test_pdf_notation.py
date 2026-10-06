"""Regression for the actual PDF failure: code points survived, run baselines did not."""
import fitz
import pytest
from pptx import Presentation
from pptx.util import Inches,Pt
from app.pdf_import import convert
from app.ai.layout import resize_runs


def test_opening_retains_date_field_background_without_sample_text():
    from app.grounded import TEMPLATE_PATH
    from slide_engine import template_policy as T
    prs=Presentation(TEMPLATE_PATH)
    T.retain_opening_artwork(prs);T.retain_opening_artwork(prs)
    layout=next(l for l in prs.slide_layouts if l.name==T.OPENING_LAYOUT)
    masks=[s for s in layout.shapes if s.name.startswith('template-placeholder-artwork-')]
    assert len(masks)==1
    source=next(s for s in layout.shapes if s.is_placeholder and s.placeholder_format.idx==10)
    assert (masks[0].left,masks[0].top,masks[0].width,masks[0].height)==(source.left,source.top,source.width,source.height)
    assert not masks[0].is_placeholder
    assert not masks[0]._element.xpath('./p:txBody')
    assert masks[0]._element.xpath('./p:spPr/a:solidFill')[0].xml==source._element.xpath('./p:spPr/a:solidFill')[0].xml


def test_import_preserves_superscript_and_subscript_origins(tmp_path):
    doc=fitz.open();page=doc.new_page(width=720,height=405)
    page.insert_text((30,60),'x',fontsize=18)
    page.insert_text((39,54),'2',fontsize=12)
    page.insert_text((46,60),' + y',fontsize=18)
    page.insert_text((75,65),'0',fontsize=12)
    pdf=tmp_path/'notation.pdf';out=tmp_path/'notation.pptx';doc.save(pdf);doc.close()
    convert(pdf,out,lambda i:str(tmp_path/f'{i}.png'))
    runs=[r for sh in Presentation(out).slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    exponent=next(r for r in runs if r.text=='2')
    subscript=next(r for r in runs if r.text=='0')
    assert int(exponent._r.rPr.get('baseline','0'))>0
    assert int(subscript._r.rPr.get('baseline','0'))<0


def test_non_unicode_math_glyphs_preserved_as_source_pixels(tmp_path, monkeypatch):
    source=tmp_path/'glyphs.pdf';output=tmp_path/'glyphs.pptx'
    with fitz.open() as doc:
        page=doc.new_page(width=720,height=405)
        page.insert_text((30,60),'Editable explanation',fontsize=18)
        page.insert_text((30,100),'(x)',fontsize=18)
        doc.save(source)
    original=fitz.Page.get_text
    def text_with_font_controls(page, option='text', *args, **kwargs):
        result=original(page,option,*args,**kwargs)
        if option=='dict':
            for block in result['blocks']:
                for line in block.get('lines',[]):
                    for span in line['spans']:
                        if span['text']=='(x)': span['text']='\x00x\x01'
        return result
    monkeypatch.setattr(fitz.Page,'get_text',text_with_font_controls)
    evidence=convert(source,output,lambda i:str(tmp_path/f'{i}.png'))
    slide=Presentation(output).slides[0]
    assert ''.join(s.text for s in slide.shapes if s.has_text_frame)=='Editable explanation'
    pictures=[s for s in slide.shapes if s.name.startswith('PDF page 1 preserved text glyphs')]
    assert len(pictures)==1
    # The fallback must contain the original rendered equation, not a blank or
    # replacement character. Compare exact PNG bytes with the source region.
    with fitz.open(source) as doc:
        line=original(doc[0],'dict')['blocks'][1]['lines'][0]
        expected=doc[0].get_pixmap(matrix=fitz.Matrix(3,3),clip=fitz.Rect(line['bbox']),alpha=True)
    assert pictures[0].image.blob==expected.tobytes('png')
    assert evidence['page_evidence'][0]['rasterized_text_lines']==1


def test_resizing_math_retains_small_runs_and_baselines():
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    tf=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(3),Inches(1)).text_frame
    a=tf.paragraphs[0].add_run();a.text='X';a.font.size=Pt(18)
    b=tf.paragraphs[0].add_run();b.text='2';b.font.size=Pt(12);b._r.get_or_add_rPr().set('baseline','50000')
    resize_runs(tf,24)
    assert a.font.size.pt==24 and b.font.size.pt==16
    assert b._r.rPr.get('baseline')=='50000'


def test_graphic_white_padding_does_not_become_giant_logo_box(tmp_path):
    doc=fitz.open();page=doc.new_page(width=720,height=405)
    page.insert_text((20,30),'A body line',fontsize=18)
    # Uniform outer padding touches the artwork and previously merged it into
    # a near page-wide picture, making legal logo placement impossible.
    page.draw_rect((50,70,650,300),fill=(1,1,1),color=None)
    page.draw_rect((150,170,200,200),fill=(.8,0,0),color=None)
    pdf=tmp_path/'logo.pdf';out=tmp_path/'logo.pptx';doc.save(pdf);doc.close()
    convert(pdf,out,lambda i:str(tmp_path/f'{i}.png'))
    pictures=[sh for sh in Presentation(out).slides[0].shapes if sh.shape_type==13]
    assert len(pictures)==1
    assert pictures[0].width/Inches(1)<1
    assert pictures[0].left/Inches(1)==pytest.approx(150/54,abs=.03)


def test_transparent_path_does_not_merge_unrelated_graphics(tmp_path):
    doc=fitz.open();page=doc.new_page(width=720,height=405)
    page.insert_text((20,30),'Two independent figures',fontsize=18)
    page.draw_rect((10,60,690,350),fill=(0,0,0),fill_opacity=0,color=None)
    page.draw_rect((50,100,150,200),fill=(.8,0,0),color=None)
    page.draw_rect((350,100,550,250),fill=(0,0,.8),color=None)
    pdf=tmp_path/'independent.pdf';out=tmp_path/'independent.pptx';doc.save(pdf);doc.close()
    convert(pdf,out,lambda i:str(tmp_path/f'{i}.png'))
    pictures=[sh for sh in Presentation(out).slides[0].shapes if sh.shape_type==13]
    assert len(pictures)==2
    assert sorted(p.left/Inches(1) for p in pictures)==pytest.approx([50/54,350/54],abs=.03)


@pytest.mark.renderer
def test_shifted_pdf_glyph_size_survives_powerpoint_render(tmp_path):
    import json
    from app.rendering import render_to_pdf
    doc=fitz.open();page=doc.new_page(width=720,height=405)
    page.insert_text((30,60),'x',fontsize=18)
    page.insert_text((39,54),'2',fontsize=12)
    page.insert_text((80,60),'y',fontsize=18)
    page.insert_text((89,66),'0',fontsize=12)
    source=tmp_path/'shifted.pdf';ppt=tmp_path/'shifted.pptx';output=tmp_path/'rendered.pdf'
    doc.save(source);doc.close()
    convert(source,ppt,lambda i:str(tmp_path/f'preview-{i}.png'))
    render_to_pdf(str(ppt),str(output))
    with fitz.open(output) as rendered:
        spans=[s for b in rendered[0].get_text('dict')['blocks'] for line in b.get('lines',[]) for s in line['spans']]
        base=next(s for s in spans if s['text'].strip()=='x')
        exponent=next(s for s in spans if s['text'].strip()=='2')
        subscript=next(s for s in spans if s['text'].strip()=='0')
        # OOXML baseline runs are scaled by the renderer: PowerPoint uses 2/3,
        # LibreOffice uses 58%. Check each engine's native size, so an accidental
        # second reduction (the original import bug) still fails on either.
        renderer=json.loads(output.with_suffix('.renderer.json').read_text())['renderer']
        expected=.58 if renderer=='LibreOffice' else 2/3
        assert exponent['size']/base['size']==pytest.approx(expected,abs=.03)
        assert subscript['size']/base['size']==pytest.approx(expected,abs=.03)
        assert exponent['origin'][1]<base['origin'][1]<subscript['origin'][1]
