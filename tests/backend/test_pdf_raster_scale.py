"""Native PDF composition must not magnify raster pixels to fill empty space."""
from io import BytesIO

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
import pytest

from app import grounded
from app.qa.artifact_coverage import audit
from slide_engine import templates


def source_deck(path, dense=False, with_image=True):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    cover = prs.slides.add_slide(prs.slide_layouts[6])
    cover.shapes.add_textbox(Inches(.8), Inches(2), Inches(6), Inches(1)).text = 'Optimization'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide._element.cSld.set('name', 'sss:pdf-import')
    if with_image:
        image = Image.new('RGB', (120, 30), 'white')
        ImageDraw.Draw(image).text((3, 5), 'SOURCE MARK', fill='#b31b1b')
        png = BytesIO()
        image.save(png, format='PNG')
        picture = slide.shapes.add_picture(BytesIO(png.getvalue()), Inches(.8), Inches(.7),
                                          Inches(10.5 if dense else 1), Inches(5.7 if dense else .25))
        picture.name = 'PDF preserved source artwork'
    for text, y in [('Gradient Descent', 6.45 if dense else 1.4),
                    ('Keep the original source wording.', 6.8 if dense else 2)]:
        shape = slide.shapes.add_textbox(Inches(.8), Inches(y), Inches(4), Inches(.45))
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        run = tf.paragraphs[0].add_run()
        run.text = text
        run.font.name = 'Arial'
        run.font.size = Pt(18)
    prs.save(path)


@pytest.mark.parametrize('dense', [False, True])
def test_pdf_raster_group_does_not_upscale_and_dense_groups_still_fit(tmp_path, dense):
    source, output = tmp_path/'source.pptx', tmp_path/'output.pptx'
    source_deck(source, dense=dense)
    original = list(Presentation(source).slides[1].shapes)
    with templates.use('cpe'):
        report = grounded.build_deck(source, output)
        assert audit(source, output, report)['status'] == 'passed'
        result = list(Presentation(output).slides[1].shapes)
    # All original pixels and words survive, with one common geometry scale.
    assert len(result) == len(original)
    factor = result[0].width/original[0].width
    assert 0 < factor <= 1
    assert factor < 1 if dense else factor == pytest.approx(1)
    assert result[0].image.blob == original[0].image.blob
    for before, after in zip(original, result):
        assert after.width/before.width == pytest.approx(factor, abs=.00001)
        assert after.height/before.height == pytest.approx(factor, abs=.00001)
        assert (after.left-result[0].left) == pytest.approx((before.left-original[0].left)*factor, abs=2)
        assert (after.top-result[0].top) == pytest.approx((before.top-original[0].top)*factor, abs=2)
        if before.has_text_frame:
            assert after.text == before.text
            assert after.text_frame.paragraphs[0].runs[0].font.size.pt >= 11


def test_sparse_native_text_only_pdf_can_still_grow_for_readability(tmp_path):
    source, output = tmp_path/'source.pptx', tmp_path/'output.pptx'
    source_deck(source, with_image=False)
    with templates.use('cpe'):
        grounded.build_deck(source, output)
    before = Presentation(source).slides[1].shapes[0]
    after = Presentation(output).slides[1].shapes[0]
    assert after.width > before.width
    assert after.text == before.text
    assert 18 < after.text_frame.paragraphs[0].runs[0].font.size.pt <= 40
