from pathlib import Path
import fitz
import pytest
from pptx import Presentation
from fastapi.testclient import TestClient
from app import api, sessions, generations, pdf_import
from app.main import app
from app.ai import providers, output_qa
from test_ai_pipeline import ai_session, config
from test_openai import openai_config, reply, completed
from slide_engine import inventory, template_policy as T
from app.ai import source_decisions
from app import grounded
from rubric_fixtures import source_choice


def pdf_bytes():
    doc=fitz.open()
    for title in ('Synthetic title','Section content'):
        page=doc.new_page(width=960,height=540)
        page.insert_text((60,90),title,fontsize=32)
        page.insert_text((60,150),'Preserve exact wording and 42 units.',fontsize=18)
        page.draw_rect((500,200,850,420),color=(.6,.1,.2),fill=(.9,.9,.9))
        page.draw_line((530,350),(810,240),color=(.7,.1,.1),width=3)
        page.insert_text((550,380),'Plot label',fontsize=14)
    return doc.tobytes()


def test_pdf_upload_has_editable_text_graphics_and_original_previews():
    client=TestClient(app)
    response=client.post('/api/sessions',files={'file':('slides.pdf',pdf_bytes(),'application/pdf')})
    assert response.status_code==200,response.text
    data=response.json();s=sessions.get(data['session_id'])
    try:
        assert data['slide_count']==2
        p=Presentation(s.source_path)
        assert p.slides[0].shapes[0].shape_type==13  # actual graphic, never whole-page thumbnail
        assert any(x.has_text_frame and 'Synthetic title' in x.text for x in p.slides[0].shapes)
        assert Path(s.preview_path(0,'before')).is_file()
        assert s.pdf_import['pages']==2
        assert s.pdf_import['source_pdf_sha256']==generations.sha256(s.source_pdf)
    finally: sessions.delete(s.id)
    assert not Path(s.dir).exists()


def test_blank_pdf_page_is_preserved_and_reported_as_image():
    doc=fitz.open();doc.new_page()
    client=TestClient(app)
    r=client.post('/api/sessions',files={'file':('scan.pdf',doc.tobytes(),'application/pdf')})
    assert r.status_code==200,r.text
    session=sessions.get(r.json()['session_id'])
    try:
        evidence=session.pdf_import['page_evidence'][0]
        assert evidence['page_image_regions']==1
        assert 'not individually editable' in evidence['preservation']
        assert len(Presentation(session.source_path).slides)==1
    finally: sessions.delete(session.id)


def test_failed_source_preparation_still_runs_independent_qa(ai_session,monkeypatch):
    monkeypatch.setattr(providers,'generate',lambda *a,**kw:{'status':'timeout','message':'AI request timed out.'})
    calls=[]
    def review(*a,**kw):
        calls.append(True)
        return {n:{'status':'passed','findings':[]} for n in output_qa.CHECKS+('output_qa_visual',)}
    monkeypatch.setattr(output_qa,'run',review)
    r=generations.build(ai_session,mode='ai')
    assert r['checks']['ai_redesign']['status']=='error'
    assert calls and all(r['checks'][n]['status']=='passed' for n in output_qa.CHECKS)
    assert generations.download_allowed(r)
    assert not generations.checks_satisfied(r)


def test_timeout_retries_once_with_bounded_timeout(openai_config,monkeypatch):
    import requests
    calls=[]
    def post(url,**kw):
        calls.append(kw['timeout'])
        if len(calls)==1: raise requests.Timeout()
        return reply(completed({'ok':True}))
    monkeypatch.setattr(providers.requests,'post',post)
    r=providers.generate('planner','JSON',{})
    assert r['status']=='completed' and r['request_attempts']==2
    assert all(30<=timeout[1]<=300 for timeout in calls)


def test_pdf_download_has_same_release_gate_and_detects_tampering(ai_session,monkeypatch,tmp_path):
    r=generations.build(ai_session)
    export=tmp_path/'candidate.pdf';export.write_bytes(pdf_bytes())
    r['pdf_export']={'path':str(export),'sha256':generations.sha256(export),'candidate_sha256':r['candidate_sha256']}
    client=TestClient(app);url=f'/api/sessions/{ai_session.id}/download'
    query={'generation_id':r['generation_id'],'format':'pdf'}
    r['state']='error'
    r['processing_complete']=False
    assert client.get(url,params=query).status_code==409
    assert client.get(url,params={**query,'draft':True}).status_code==409
    ids=[f['id'] for f in r['findings'] if f['severity']=='review']
    if ids:
        generations.decide(ai_session,generations.Decision(generation_id=r['generation_id'],candidate_sha256=r['candidate_sha256'],finding_ids=ids,rationale='Reviewed synthetic test fixture.'))
    r['state']='error'  # QA errors do not block an intact, completed PDF.
    r['processing_complete']=True
    assert client.get(url,params=query).content==export.read_bytes()
    export.write_bytes(b'changed')
    assert client.get(url,params=query).status_code==409


def test_cover_photo_is_not_an_old_slide_inset(tmp_path):
    from io import BytesIO
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    from app.qa.artifact_coverage import audit
    p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
    slide=p.slides.add_slide(p.slide_layouts[6]);slide.background.fill.solid()
    slide.background.fill.fore_color.rgb=RGBColor(255,255,255)
    slide.shapes.add_textbox(Inches(.7),Inches(1),Inches(5),Inches(1)).text='Synthetic native cover'
    slide.shapes.add_textbox(Inches(.7),Inches(2.3),Inches(5),Inches(1)).text='Preserved supporting text'
    template=Presentation(grounded.TEMPLATE_PATH)
    photo=next(s for s in template.slide_layouts[1].shapes if s.shape_type==13)
    pic=slide.shapes.add_picture(BytesIO(photo.image.blob),Inches(7),Inches(1),width=Inches(4))
    source=tmp_path/'cover-source.pptx';p.save(source)
    digest=inventory.sha256(source);objects=source_decisions.source_objects(p,0,digest)
    choice=source_choice({'objects':objects,'source_slide':0})
    role=choice['elements'][-1];role.update(role='image',content_bearing=True,confidence='high',contains_logo=False)
    candidate=tmp_path/'cover-candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions={'0':choice})
    result=Presentation(candidate).slides[0]
    assert T.is_cover(result)
    assert result.slide_layout.name==T.OPENING_LAYOUT
    assert len([s for s in result.slide_layout.shapes if s.shape_type==13])==1  # Red artwork, no campus photo.
    assert not any(s.name==T.BACKGROUND_NAME for s in result.shapes)
    retained=next(s for s in result.shapes if s.shape_type==13)
    assert retained.image.blob==pic.image.blob
    assert abs(retained.width/retained.height-pic.width/pic.height)<.001
    assert abs(retained.width/Inches(T.COVER_SUPPORT[2])-1)<.001 or abs(retained.height/Inches(T.COVER_SUPPORT[3])-1)<.001
    assert T.check(candidate)['status']=='passed'
    assert audit(source,candidate,report)['status']=='passed'


def test_white_pdf_background_is_not_imported_as_a_cover_picture(tmp_path):
    doc=fitz.open();page=doc.new_page(width=960,height=540)
    page.draw_rect((10,10,930,510),color=(1,1,1),fill=(1,1,1))
    page.insert_text((70,90),'Cover title',fontsize=32)
    page.insert_text((70,515),'Keep this meaningful bottom note.',fontsize=11)
    source=tmp_path/'cover.pdf';doc.save(source);doc.close()
    imported=tmp_path/'cover.pptx'
    pdf_import.convert(source,imported,lambda i:tmp_path/f'original-{i}.png')
    p=Presentation(imported)
    assert (p.slide_width,p.slide_height)==T.CANVAS
    assert not any(s.shape_type==13 for s in p.slides[0].shapes)
    candidate=tmp_path/'candidate.pptx';report=grounded.build_deck(imported,candidate)
    slide=Presentation(candidate).slides[0]
    assert not any(s.name==T.BACKGROUND_NAME for s in slide.shapes)
    note=next(s for s in slide.shapes if s.has_text_frame and 'meaningful bottom note' in s.text)
    assert T.contains([v/T.EMU for v in (note.left,note.top,note.width,note.height)],T.COVER_DETAILS)
    assert audit_pdf_cover(imported,candidate,report)=='passed'
    out=Presentation(candidate);out.slides[0].shapes[-1].left=0;out.save(candidate)
    assert any(f['code']=='COVER_TEXT_POSITION' for f in T.check(candidate)['findings'])


def audit_pdf_cover(source,candidate,report):
    from app.qa.artifact_coverage import audit
    return audit(source,candidate,report)['status']


def test_frontend_document_is_not_cached_across_restarts():
    response=TestClient(app).get('/')
    assert response.status_code==200
    assert response.headers['cache-control']=='no-store'
