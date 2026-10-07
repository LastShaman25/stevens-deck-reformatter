"""PDF glyph positions, not invented spaces, govern native bullet alignment."""
from copy import deepcopy

import fitz
import pytest
from pptx import Presentation
from pptx.util import Inches

from app import grounded, pdf_import, pdf_regions
from app.pdf_text_spacing import apply_bullet_tab, bullet_tabs, scale_tabs
from app.qa.artifact_coverage import audit
from slide_engine import templates


def positioned_line(text='\u25cfSmaller', gap=18, direction=(1, 0)):
    # Matches the report's same-span bullet with an 18pt visible glyph gap.
    chars = [{'c': text[0], 'origin': (32, 95), 'bbox': (32, 70, 47, 102)}]
    for i, c in enumerate(text[1:]):
        x = 47+gap+i*12
        chars.append({'c': c, 'origin': (x, 95), 'bbox': (x, 70, x+12, 102)})
    line = {'bbox': (32, 70, chars[-1]['bbox'][2], 102), 'dir': direction,
            'spans': [{'text': text, 'size': 25}]}
    raw = {**line, 'spans': [{'chars': chars, 'size': 25}]}
    class Page:
        def get_text(self, kind, **kwargs):
            assert kind == 'rawdict'
            return {'blocks': [{'type': 0, 'lines': [raw]}]}
    return Page(), line


def test_tab_restores_source_word_origin_without_changing_words_and_scales():
    page, line = positioned_line()
    offset = bullet_tabs(page, [line])[id(line)]
    prs = Presentation()
    paragraph = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(
        0, 0, Inches(6), Inches(1)).text_frame.paragraphs[0]
    original = deepcopy(line['spans'])
    spans = apply_bullet_tab(paragraph, line['spans'], offset, 1/72)
    for span in spans:
        paragraph.add_run().text = span['text']
    assert line['spans'] == original
    assert paragraph.text == '\u25cf\tSmaller'
    assert pdf_import.compact(paragraph.text) == pdf_import.compact(line['spans'][0]['text'])
    assert int(paragraph._p.xpath('./a:pPr/a:tabLst/a:tab')[0].get('pos'))/12700 == pytest.approx(33)
    scale_tabs(paragraph, .65)
    assert int(paragraph._p.xpath('./a:pPr/a:tabLst/a:tab')[0].get('pos'))/12700 == pytest.approx(33*.65)


@pytest.mark.parametrize('text,gap,direction', [
    ('\u25cfSmaller', 1, (1, 0)),  # The source really has no visible gap.
    ('\u25cf Smaller', 18, (1, 0)),  # Existing whitespace is already preserved.
    ('\u00b7Smaller', 18, (1, 0)),  # A math middle dot is not a list marker.
    ('\u25cfSmaller', 18, (0, 1)),  # x coordinates cannot measure rotated lines.
])
def test_uncertain_or_already_spaced_lines_are_unchanged(text, gap, direction):
    page, line = positioned_line(text, gap, direction)
    assert bullet_tabs(page, [line]) == {}


def test_import_and_native_rebuild_preserve_bullet_spacing_and_link_underline(tmp_path, monkeypatch):
    doc = fitz.open()
    cover = doc.new_page(width=720, height=405)
    cover.insert_text((40, 80), 'Optimization', fontsize=30)
    page = doc.new_page(width=720, height=405)
    # Portable embedded sans font; no dependency on a machine's Arial install.
    page.insert_font(fontname='Test', fontbuffer=fitz.Font('helv').buffer)
    page.insert_text((32, 95), '\u2022Smaller learning rate', fontname='Test', fontsize=25)
    content = page.get_contents()[0]
    stream = doc.xref_stream(content)
    stream = stream.replace(b'[<0074', b'[<0074> -720 <')
    doc.update_stream(content, stream)
    # Reopen so MuPDF uses the actual stream advances.
    payload = doc.tobytes()
    doc.close()
    doc = fitz.open(stream=payload, filetype='pdf')
    page = doc[1]
    line = page.get_text('dict')['blocks'][0]['lines'][0]
    rect = fitz.Rect(line['bbox'])
    rule = fitz.Rect(rect.x0, 97, rect.x1, 98)
    page.draw_rect(rule, color=None, fill=(0, 0, 0))
    page.insert_link({'kind': fitz.LINK_URI, 'from': rect, 'uri': 'https://example.org/course'})
    source = tmp_path/'source.pdf'
    doc.save(source)
    doc.close()
    get_text = fitz.Page.get_text
    def exporter_text(self, option='text', *args, **kwargs):
        # Some real exporters expose this gap without a whitespace character.
        # Use MuPDF's supported mode to produce that same extraction condition;
        # all glyph coordinates, underline geometry and PDF pixels remain real.
        if option in ('dict', 'rawdict'):
            default = fitz.TEXTFLAGS_DICT if option == 'dict' else fitz.TEXTFLAGS_RAWDICT
            kwargs['flags'] = kwargs.get('flags', default) | fitz.TEXT_INHIBIT_SPACES
        return get_text(self, option, *args, **kwargs)
    monkeypatch.setattr(fitz.Page, 'get_text', exporter_text)
    source_font_required = pdf_regions.needs_source_font
    # Exercise the native-text branch with this known portable sans fixture.
    # Real Arial/Helvetica PDFs already enter that branch without an override.
    monkeypatch.setattr(pdf_regions, 'needs_source_font', lambda line:
                        False if all(s['font'] == 'NimbusSans-Regular' for s in line['spans'])
                        else source_font_required(line))
    imported = tmp_path/'source.pptx'
    candidate = tmp_path/'candidate.pptx'
    with templates.use('cpe'):
        receipt = pdf_import.convert(source, imported, lambda i: str(tmp_path/f'before-{i}.png'))
        assert receipt['page_evidence'][1]['native_bullet_tabs'] == 1
        assert receipt['page_evidence'][1]['native_link_underlines'] == 1
        report = grounded.build_deck(imported, candidate)
        assert audit(imported, candidate, report)['status'] == 'passed'
    native = next(s for s in Presentation(imported).slides[1].shapes
                  if s.has_text_frame and 'Smaller' in s.text)
    output = next(s for s in Presentation(candidate).slides[1].shapes
                  if s.has_text_frame and 'Smaller' in s.text)
    assert native.text == output.text == '\u2022\tSmaller learning rate'
    assert native.text_frame.paragraphs[0].runs[0].font.underline is True
    assert output.text_frame.paragraphs[0].runs[0].font.underline is True
    def tab(shape):
        return int(shape.text_frame.paragraphs[0]._p.xpath('./a:pPr/a:tabLst/a:tab')[0].get('pos'))
    assert tab(output)/tab(native) == pytest.approx(output.width/native.width, abs=.0001)
