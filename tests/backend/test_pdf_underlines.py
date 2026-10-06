import fitz
import pytest
from pptx import Presentation
from app import pdf_import


def source(tmp_path, offset=1.4, fraction=1.0, with_link=True):
    doc=fitz.open(); page=doc.new_page(width=600,height=400)
    page.insert_text((50,80),'https://example.com/terms',fontname='helv',fontsize=12,color=(0,0,1))
    span=page.get_text('dict')['blocks'][0]['lines'][0]['spans'][0]
    rect=fitz.Rect(span['bbox']); baseline=span['origin'][1]
    rule=fitz.Rect(rect.x0,baseline+offset,rect.x0+rect.width*fraction,baseline+offset+.7)
    page.draw_rect(rule,color=None,fill=(0,0,1))
    if with_link:page.insert_link({'kind':fitz.LINK_URI,'from':rect | rule,'uri':'https://example.com/terms'})
    path=tmp_path/'source.pdf';doc.save(path);doc.close()
    return path


def test_pdf_link_rule_becomes_native_underline_without_duplicate_picture(tmp_path):
    original=source(tmp_path)
    receipt=pdf_import.convert(original,tmp_path/'out.pptx',lambda i:str(tmp_path/f'original-{i}.png'))
    slide=Presentation(tmp_path/'out.pptx').slides[0]
    run=next(s.text_frame.paragraphs[0].runs[0] for s in slide.shapes if s.has_text_frame and s.text)
    assert run.font.underline is True
    assert run.text=='https://example.com/terms'
    assert not any(s.shape_type==13 for s in slide.shapes)
    assert any(s.click_action.hyperlink.address=='https://example.com/terms' for s in slide.shapes)
    assert receipt['page_evidence'][0]['native_link_underlines']==1


@pytest.mark.parametrize('options',[{'offset':-4},{'fraction':.5},{'with_link':False}])
def test_ambiguous_or_nonlink_rules_remain_graphics(tmp_path,options):
    original=source(tmp_path,**options)
    receipt=pdf_import.convert(original,tmp_path/'out.pptx',lambda i:str(tmp_path/f'original-{i}.png'))
    slide=Presentation(tmp_path/'out.pptx').slides[0]
    assert any(s.shape_type==13 for s in slide.shapes)
    assert receipt['page_evidence'][0]['native_link_underlines']==0
