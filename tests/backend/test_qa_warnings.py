import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from app import generations
from app.ai import rubric, pipeline, providers, layout
from rubric_fixtures import passed_checks, repair_evidence
from test_ai_pipeline import mocked_provider, identity_plan, ai_session, config


def test_warning_only_review_completes_without_cosmetic_repair(ai_session,monkeypatch):
    planned=[]
    def provider(role,system,payload,*args,**kwargs):
        result=mocked_provider(role,system,payload,*args,**kwargs)
        if role=='planner': planned.append(payload['slide'])
        if role=='reviewer':
            finding={**repair_evidence(),'criterion':'spatial_layout','severity':'warning',
                     'message':'Optional extra caption spacing; all text remains readable.'}
            result['data'].update(verdict='passed',findings=[finding],rubric=passed_checks([finding]))
        return result
    monkeypatch.setattr(providers,'generate',provider)
    record=generations.build(ai_session,mode='ai')
    assert not planned
    assert record['checks']['ai_visual_review']['status']=='passed'
    assert record['checks']['ai_visual_review']['findings']
    assert record['state']=='ready' and generations.download_allowed(record)
    assert not record.get('output_qa_repairs')


@pytest.mark.parametrize('severity',['blocking','review',None])
def test_passed_label_cannot_hide_material_or_unclassified_qa_finding(severity):
    record={'mode':'ai','state':'ready','checks':{},'findings':[],'human_decisions':[]}
    record['checks']={n:{'status':'passed','findings':[]} for n in generations.required_checks(record)}
    record['checks']['ai_visual_review']['findings']=[{'severity':severity}]
    assert not generations.download_allowed(record)


def test_code_box_alignment_is_not_recentered_and_text_style_is_preserved():
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    for i,width in enumerate([4,2,3]):
        shape=slide.shapes.add_textbox(Inches(2),Inches(1+i),Inches(width),Inches(.3))
        shape.name=f'line{i}|code';shape.text='    x = 1' if i else 'if True:'
        shape.text_frame.paragraphs[0].runs[0].font.name='Consolas'
        shape.text_frame.paragraphs[0].runs[0].font.size=Pt(12)
    plan=identity_plan(layout.describe(slide))
    for obj in plan['objects']: obj.update(x=1,w=4,role='body',font_size=24,color='red')
    safe=layout.constrain(layout.LayoutPlan.model_validate(plan),slide,
                          [{'object_ids':[s.name for s in slide.shapes]}])
    assert all(o.x==1 and o.w==4 for o in safe.objects)
    assert all(o.role=='keep' and o.font_size is None and o.color=='keep' for o in safe.objects)
    assert slide.shapes[1].text=='    x = 1'


def test_warning_rubric_status_must_match_evidence():
    f=pipeline.VisualFinding(**repair_evidence(),criterion='spatial_layout',severity='warning',message='Cosmetic gap.')
    checks=[rubric.CriterionResult(**c) for c in passed_checks([f.model_dump()])]
    rubric.validate_checks(checks,[f])
    assert rubric.finding_status([f])=='passed'
    next(c for c in checks if c.criterion=='spatial_layout').status='passed'
    with pytest.raises(ValueError): rubric.validate_checks(checks,[f])
