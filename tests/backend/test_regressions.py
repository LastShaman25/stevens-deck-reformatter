from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from app import grounded, sessions
from app.qa import brand_lint
from slide_engine import ir, planner, verify


def text(slide, value, y=2):
    shape = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(7), Inches(1))
    shape.text = value
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(20)
    return shape


def fixture(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    text(s, 'Opening', .3)
    text(s, 'First paragraph\nSecond paragraph\nThird paragraph')
    s = p.slides.add_slide(p.slide_layouts[6])
    text(s, 'Results', .3)
    t = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            cell.text = f'Cell {i} {j}'
    text(s, 'Footnote that must survive', 4)
    s = p.slides.add_slide(p.slide_layouts[6])
    text(s, 'Thank you', .3)
    text(s, 'Jane Example\njane@example.edu')
    path = tmp_path / 'source.pptx'
    p.save(path)
    return path


def all_text(path):
    return '\n'.join(sh.text for s in Presentation(path).slides for sh in s.shapes if sh.has_text_frame)


def test_cover_closing_and_table_context(tmp_path):
    src = fixture(tmp_path)
    out = tmp_path / 'out.pptx'
    grounded.build_deck(str(src), str(out), use_vision=False, use_llm=False)
    actual = all_text(out)
    for expected in ['First paragraph', 'Second paragraph', 'Third paragraph',
                     'Footnote that must survive', 'Jane Example', 'jane@example.edu']:
        assert expected in actual


def test_flagged_is_not_coverage(tmp_path):
    deck = ir.extract(fixture(tmp_path))
    plans = planner.plan_deterministic(deck)
    aid = plans[0].blocks[0].atom_id
    plans[0].blocks = [b for b in plans[0].blocks if b.atom_id != aid]
    plans[0].flagged.append(aid)
    assert not verify.coverage(deck, plans).ok


def test_vertical_offslide_is_detected(tmp_path):
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    text(s, 'This is below the slide', 8)
    path = tmp_path / 'offslide.pptx'
    p.save(path)
    assert brand_lint.lint(path)['total_fail'] > 0


def test_expired_lookup_does_not_revive():
    s = sessions.create()
    s.touched -= sessions.TTL_SECONDS + 1
    try:
        assert sessions.get(s.id) is None
    finally:
        sessions.delete(s.id)
