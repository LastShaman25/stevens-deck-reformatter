from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches
from app import grounded, generations
from app.qa import artifact_coverage
from app.ai import output_qa
from slide_engine import bookends, template_policy as T
from test_ai_pipeline import ai_session, config, mocked_provider


def source(tmp_path,thanks=False):
    prs=Presentation()
    for text in ('Course title','Thank you!' if thanks else 'References that must survive'):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text=text
    path=tmp_path/'source.pptx';prs.save(path);return path


@pytest.mark.parametrize('thanks',[False,True])
def test_required_closing_preserves_source_and_does_not_duplicate_existing(tmp_path,thanks):
    path=source(tmp_path,thanks);before=path.read_bytes();out=tmp_path/'out.pptx'
    report=grounded.build_deck(path,out,require_closing=True)
    prs=Presentation(out)
    assert len(prs.slides)==(2 if thanks else 3)
    assert T.is_closing(prs.slides[-1]) and bookends.has_thanks(prs.slides[-1])
    assert report['source_to_output_slides']=={'0':[0],'1':[1]}
    assert artifact_coverage.audit(path,out,report)['status']=='passed'
    bookends.ensure(out,report)
    assert len(Presentation(out).slides)==len(prs.slides)
    assert path.read_bytes()==before
    if not thanks:assert 'References that must survive' in '\n'.join(s.text for s in prs.slides[1].shapes if s.has_text_frame)


@pytest.mark.parametrize('mutation',['remove','change','extra','move'])
def test_added_closing_cannot_hide_missing_changed_or_invented_content(tmp_path,mutation):
    path=source(tmp_path);out=tmp_path/'out.pptx';report=grounded.build_deck(path,out,require_closing=True)
    prs=Presentation(out)
    if mutation=='remove':
        sid=prs.slides._sldIdLst[-1];prs.part.drop_rel(sid.rId);prs.slides._sldIdLst.remove(sid)
    elif mutation=='change':prs.slides[-1].shapes[0].text='Invented claim'
    elif mutation=='extra':prs.slides[-1].shapes.add_textbox(0,0,Inches(1),Inches(1)).text='Unapproved'
    else:
        sid=prs.slides._sldIdLst[-1];prs.slides._sldIdLst.remove(sid);prs.slides._sldIdLst.insert(1,sid)
    prs.save(out)
    assert artifact_coverage.audit(path,out,report)['status']=='failed'


def test_pipeline_reviews_added_closing_with_explicit_authorization(ai_session,monkeypatch):
    from app.ai import providers
    prs=Presentation(ai_session.source_path)
    next(s for s in prs.slides[-1].shapes if s.has_text_frame and s.text.strip()=='Thank you').text='Final references'
    prs.save(ai_session.source_path)
    reviews=[]
    def provider(role,system,payload,*args,**kwargs):
        if role=='reviewer':reviews.append(payload)
        return mocked_provider(role,system,payload,*args,**kwargs)
    monkeypatch.setattr(providers,'generate',provider)
    record=generations.build(ai_session,mode='ai')
    assert record['state']=='ready'
    assert len(reviews)==4 and len(record['ai_pipeline']['final_reviews'])==4
    assert reviews[-1]['authorized_addition']==bookends.AUTHORIZATION
    assert reviews[-1]['original_objects']==[]
    assert record['report']['slide_count']==4
    assert generations.public(record)['added_slides'][0]['output_slide']==3
