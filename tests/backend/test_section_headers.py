from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches
import pytest
from app import grounded
from app.ai import source_decisions, layout
from app.qa.artifact_coverage import audit
from slide_engine import inventory, template_policy as T
from rubric_fixtures import source_choice


def section_source(tmp_path):
    p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
    for title in ('Module introduction','Lesson 2: Continuous delivery','Regular content'):
        s=p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_textbox(Inches(1),Inches(1),Inches(7),Inches(1)).text=title
        s.shapes.add_textbox(Inches(1),Inches(3),Inches(6),Inches(1)).text='Supporting text'
    footer=p.slides[1].shapes.add_textbox(Inches(.2),Inches(6.9),Inches(3),Inches(.3))
    footer.text='STEVENS INSTITUTE OF TECHNOLOGY'
    path=tmp_path/'source.pptx';p.save(path)
    digest=inventory.sha256(path);choices={}
    for i in range(3):
        objects=source_decisions.source_objects(p,i,digest)
        choices[str(i)]=source_choice({'objects':objects,'source_slide':i})
    c=choices['1'];c['slide_kind']='section'
    role=c['elements'][-1]
    role.update(role='logo',confidence='high',content_bearing=False,contains_logo=True,
                artwork_action='remove',reason='Separate bottom-left Stevens footer wordmark; omit on section headers.')
    c['remove_ids']=[role['id']]
    return path,p,digest,choices


def test_section_layout_omits_footer_without_removing_regular_content_branding(tmp_path):
    source,_,_,choices=section_source(tmp_path)
    candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions=choices)
    p=Presentation(candidate);s=p.slides[1]
    assert T.is_section(s) and s.slide_layout.name=='Section Header'
    assert s.slide_layout._element.get('showMasterSp')=='0'
    assert T.contract(s)['bottom_left_logo']=='omit'
    assert all(x['id'].split(':')[1]!='master' for x in layout.template_context(s) if x['id'].startswith('template:'))
    assert any(x['id'].startswith('template:layout:') for x in layout.template_context(s))
    assert any(x['id'].startswith('template:master:') for x in layout.template_context(p.slides[2]))
    assert not any('STEVENS INSTITUTE' in x.text for x in s.shapes if x.has_text_frame)
    assert 'Lesson 2: Continuous delivery' in ' '.join(x.text for x in s.shapes if x.has_text_frame)
    assert T.check(candidate)['status']=='passed'
    assert audit(source,candidate,report)['status']=='passed'


@pytest.mark.parametrize('change',['content_slide','top_right','composite','linked'])
def test_section_footer_exception_cannot_remove_other_protected_art(tmp_path,change):
    source,p,digest,choices=section_source(tmp_path);choice=deepcopy(choices['1'])
    if change=='content_slide': choice['slide_kind']='content'
    if change=='top_right':
        p.slides[1].shapes[-1].left=Inches(10);p.slides[1].shapes[-1].top=Inches(.3)
    if change=='composite':
        p.slides[1].shapes[-1].width=Inches(12);p.slides[1].shapes[-1].height=Inches(7)
    if change=='linked':
        p.slides[1].shapes[-1].text_frame.paragraphs[0].runs[0].hyperlink.address='https://example.com'
    with pytest.raises(ValueError): source_decisions.validate(choice,p,1,digest)
