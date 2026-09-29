"""Shared-state integration tests against a disposable PostgreSQL schema and private S3 emulator."""
import json
import os
import time
import uuid
from pathlib import Path
from types import SimpleNamespace
import pytest

pytest.importorskip('psycopg')
pytest.importorskip('moto')
from app import auth, sessions
from app.cloud import db, objects, state, tasks, replay
from app.cloud.api import begin_upload, complete_upload, Upload, UploadComplete, dispatch
from app.ai import providers


@pytest.fixture
def shared(monkeypatch):
    url=os.environ.get('STEVENS_TEST_DATABASE_URL')
    if not url:pytest.skip('Set STEVENS_TEST_DATABASE_URL to a disposable PostgreSQL database.')
    from moto import mock_aws
    name='test_'+uuid.uuid4().hex
    settings={'STEVENS_STORAGE_MODE':'shared','DATABASE_URL':url,'STEVENS_STORAGE_NAMESPACE':name,
        'STEVENS_S3_ACCESS_KEY_ID':'testing','STEVENS_S3_SECRET_ACCESS_KEY':'testing','STEVENS_S3_BUCKET':'studio-test',
        'STEVENS_AUTH_MODE':'invitation','STEVENS_ADMIN_CODE':'synthetic-test-code-only-123456789',
        'STEVENS_AI_PROVIDER':'vercel','STEVENS_OFFLINE':'1','CRON_SECRET':'synthetic-cron-only',
        'STEVENS_S3_REGION':'us-east-1'}
    for key,value in settings.items():monkeypatch.setenv(key,value)
    monkeypatch.delenv('STEVENS_S3_ENDPOINT_URL',raising=False)
    with mock_aws():
        objects.client().create_bucket(Bucket=objects.bucket())
        auth.bootstrap_codes()
        with auth.database() as con:user=dict(con.execute('SELECT * FROM users').fetchone())
        login,csrf=auth.issue_login(user['id'])
        user.update(token=auth.digest(login),csrf=csrf)
        token=auth.current_user.set(user)
        try:
            with state.scope():yield user
        finally:
            auth.current_user.reset(token)
            import psycopg
            from psycopg import sql
            with psycopg.connect(url,autocommit=True) as con:
                con.execute(sql.SQL('DROP SCHEMA {} CASCADE').format(sql.Identifier(name)))


def test_accounts_and_cookie_signing_secret_are_shared(shared):
    secret=auth.cookie_secret()
    db._initialized.clear()  # Simulate a new API process.
    assert auth.cookie_secret()==secret
    with auth.database() as con:
        assert con.execute('SELECT count(*) FROM users').fetchone()[0]==1


def test_sessions_restore_files_and_identity_in_a_fresh_request(shared):
    sess=sessions.create()
    with sessions.job(sess):
        Path(sess.source_path).write_bytes(b'synthetic deck')
        sess.revision_version=3
    with state.scope():
        restored=sessions.get(sess.id)
        assert restored.dir!=sess.dir
        assert restored.owner_user_id==shared['id'] and restored.revision_version==3
        with sessions.read_job(restored):assert Path(restored.source_path).read_bytes()==b'synthetic deck'


def test_concurrent_mutations_and_stale_snapshot_are_rejected(shared):
    sess=sessions.create();state.save(sess)
    stale=state.decode(state.row(sess.id))
    try:
        with sessions.job(sess):
            with pytest.raises(ValueError,match='busy'):
                with state.job(stale):pass
        with pytest.raises(ValueError,match='changed'):state.save(stale)
    finally:
        __import__('shutil').rmtree(stale.dir)


def test_only_one_worker_claims_a_task_and_retry_retains_input(shared):
    sess=sessions.create();state.save(sess)
    tid=tasks.enqueue(sess,'redesign',{},shared)
    assert tasks.enqueue(sess,'redesign',{},shared)==tid
    with pytest.raises(ValueError,match='already queued'):tasks.enqueue(sess,'plan',{},shared)
    first=tasks.claim(tid)
    assert first and tasks.claim(tid) is None
    with db.connect() as con:
        con.execute('UPDATE processing_tasks SET lease_until=0 WHERE id=%s',(tid,))
        con.execute('UPDATE processing_sessions SET lease_until=0 WHERE id=%s',(sess.id,))
    second=tasks.claim(tid)
    assert second['lease']!=first['lease'] and second['baseline']==first['baseline']
    with pytest.raises(ValueError,match='lease'):tasks.fence(tid,first['lease'])


def test_queued_work_cannot_release_a_previously_ready_candidate(shared):
    import asyncio
    from starlette.requests import Request
    sess=sessions.create();sess.analysis={'slide_count':1};state.save(sess)
    tasks.enqueue(sess,'redesign',{},shared)
    request=Request({'type':'http','method':'GET','path':f'/api/sessions/{sess.id}/download',
                     'query_string':b'generation_id=old','headers':[]})
    result=asyncio.run(dispatch(request,shared))
    assert result.status_code==409 and b'QA must finish' in result.body


def test_completed_step_replay_does_not_call_provider_again(shared):
    sess=sessions.create();state.save(sess)
    tid=tasks.enqueue(sess,'redesign',{},shared);task=tasks.claim(tid)
    calls=[]
    value=replay.Replay(task,sess).step({'kind':'synthetic'},lambda:calls.append(1) or {'ok':True},reserve=0)
    assert value=={'ok':True}
    assert replay.Replay(task,sess).step({'kind':'synthetic'},lambda:calls.append(1),reserve=0)==value
    assert calls==[1]
    with pytest.raises(ValueError,match='Replay inputs changed'):
        replay.Replay(task,sess).step({'kind':'different-candidate'},lambda:{},reserve=0)


def test_slice_yields_before_an_external_call_would_exceed_deadline(shared):
    sess=sessions.create();state.save(sess)
    runner=replay.Replay(tasks.claim(tasks.enqueue(sess,'redesign',{},shared)),sess)
    runner.deadline=time.monotonic()
    with pytest.raises(replay.SliceComplete):runner.step('next',lambda:pytest.fail('must yield first'))


def test_cancellation_fences_workers_and_cleanup_retries(shared):
    sess=sessions.create();Path(sess.source_path).write_bytes(b'private');state.save(sess)
    tid=tasks.enqueue(sess,'redesign',{},shared);claim=tasks.claim(tid)
    assert not sessions.delete(sess.id)  # Revoked immediately; active lease defers deletion.
    with pytest.raises(ValueError):tasks.fence(tid,claim['lease'])
    assert tasks.get(tid)['state']=='cancelled'
    with db.connect() as con:con.execute('UPDATE processing_sessions SET lease_until=0 WHERE id=%s',(sess.id,))
    assert state.purge(sess.id)
    assert state.row(sess.id) is None
    assert not objects.client().list_objects_v2(Bucket=objects.bucket(),Prefix=objects.prefix(sess.id)).get('Contents')


def test_revoked_login_stops_existing_worker(shared):
    sess=sessions.create();state.save(sess)
    tid=tasks.enqueue(sess,'redesign',{},shared);claim=tasks.claim(tid)
    with auth.database() as con:con.execute('DELETE FROM logins WHERE token=?',(shared['token'],))
    with pytest.raises(ValueError):tasks.fence(tid,claim['lease'])
    with pytest.raises(ValueError):state.active(sess)
    sess._cloud_new=False


def test_direct_upload_checks_size_and_freezes_source(shared):
    request=SimpleNamespace(state=SimpleNamespace(user=shared))
    ticket=begin_upload(Upload(filename='test.pdf',size=5,target='/api/sessions'),request)
    key=ticket['fields']['key']
    objects.put(key,b'wrong-size')
    with pytest.raises(Exception,match='size'):complete_upload(UploadComplete(upload_id=ticket['upload_id']),request)
    objects.put(key,b'%PDF-')
    result=complete_upload(UploadComplete(upload_id=ticket['upload_id']),request)
    tid=json.loads(result.body)['task_id'];task=tasks.get(tid,shared)
    assert task['operation']=='import' and objects.get(task['body']['key'])==b'%PDF-'
    retry=complete_upload(UploadComplete(upload_id=ticket['upload_id']),request)
    assert json.loads(retry.body)['task_id']==tid
    objects.put(key,b'xxxxx') # Even a reused upload URL cannot change the frozen input.
    assert objects.get(task['body']['key'])==b'%PDF-'


def test_upload_ticket_is_bound_to_authenticated_owner(shared):
    from fastapi import HTTPException
    request=SimpleNamespace(state=SimpleNamespace(user=shared))
    ticket=begin_upload(Upload(filename='test.pdf',size=5,target='/api/sessions'),request)
    objects.put(ticket['fields']['key'],b'%PDF-')
    attacker=SimpleNamespace(state=SimpleNamespace(user={**shared,'id':'other'}))
    with pytest.raises(HTTPException) as exc:complete_upload(UploadComplete(upload_id=ticket['upload_id']),attacker)
    assert exc.value.status_code==404


def test_signin_throttle_survives_process_reset(shared):
    for _ in range(8):db.login_attempt('synthetic-peer',time.time())
    auth._attempts.clear()
    with pytest.raises(Exception,match='Too many'):db.login_attempt('synthetic-peer',time.time())


def test_large_previews_use_private_short_lived_object_urls(shared):
    sess=sessions.create();state.save(sess)
    path=Path(sess.dir,'large.png');path.write_bytes(b'x'*(5*1024*1024))
    response=objects.preview_response(sess,path)
    assert response.status_code==302
    assert 'X-Amz-Expires=60' in response.headers['location']
    assert response.headers['cache-control']=='no-store'


def test_account_migration_preserves_users_without_copying_sessions(shared,tmp_path):
    import sqlite3, hashlib
    from tools.migrate_accounts import migrate
    source=tmp_path/'accounts.sqlite3'
    with sqlite3.connect(source) as con:
        con.executescript('''CREATE TABLE users(id,email,issuer,subject,role,active);
            CREATE TABLE codes(token,user_id); CREATE TABLE invites(token,email,role,expires);
            CREATE TABLE settings(key,value);''')
        con.execute("INSERT INTO users VALUES ('migrated','Synthetic account','invitation','migrated','admin',1)")
        con.execute('INSERT INTO codes VALUES (?,?)',(hashlib.sha256(b'synthetic-long-code').hexdigest(),'migrated'))
        con.execute("INSERT INTO settings VALUES ('codes_bootstrapped','1')")
    before=source.read_bytes()
    with pytest.raises(ValueError,match='already contains users'):migrate(source)
    with db.connect() as con:
        for table in ('codes','logins','users','settings'):con.execute('DELETE FROM '+table)
    migrate(source)
    assert source.read_bytes()==before
    with db.connect() as con:
        assert con.execute('SELECT id,role FROM users').fetchone()=={'id':'migrated','role':'admin'}
        assert con.execute('SELECT count(*) AS n FROM logins').fetchone()['n']==0
    with sqlite3.connect(source) as con:
        con.execute('UPDATE codes SET token=?',(hashlib.sha256(b'admin').hexdigest(),))
    with pytest.raises(ValueError,match='default admin'):migrate(source)


def test_real_subprocess_import_uses_shared_state(shared,monkeypatch):
    """No in-process worker substitution: child imports and talks HTTP to S3."""
    import fitz
    from moto.server import ThreadedMotoServer
    server=ThreadedMotoServer(ip_address='127.0.0.1',port=0,verbose=False)
    server.start()
    try:
        host,port=server.get_host_and_port()
        monkeypatch.setenv('STEVENS_S3_ENDPOINT_URL',f'http://{host}:{port}')
        sess=sessions.create();state.save(sess)
        doc=fitz.open();page=doc.new_page(width=960,height=540)
        page.insert_text((80,100),'Synthetic subprocess import verification',fontsize=28)
        key=objects.prefix(sess.id)+'input.pdf';objects.put(key,doc.tobytes());doc.close()
        tid=tasks.enqueue(sess,'import',{'key':key,'filename':'synthetic.pdf'},shared)
        assert tasks.run_one(tid)['advanced']
        task=tasks.get(tid)
        assert task['state']=='completed',task['error']
        assert task['result']['slide_count']==1
        with state.scope():
            restored=state.get(sess.id)
            with sessions.read_job(restored):
                assert Path(restored.source_path).is_file()
                assert Path(restored.preview_path(0,'before')).is_file()
    finally:server.stop()


def test_http_signin_isolation_upload_and_missing_qa_gate(shared):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app import api, generations
    from app.cloud.api import router
    app=FastAPI();app.include_router(auth.router);app.include_router(api.router);app.include_router(router)
    app.middleware('http')(auth.guard)
    client=TestClient(app)
    assert client.get('/api/auth/status').json()['cloud'] is True
    login=client.post('/api/auth/code',json={'code':os.environ['STEVENS_ADMIN_CODE']})
    assert login.status_code==200
    headers={'X-CSRF-Token':login.json()['csrf']}
    ticket=client.post('/api/cloud/uploads',json={'filename':'test.pdf','size':5,'target':'/api/sessions'},headers=headers)
    assert ticket.status_code==200
    objects.put(ticket.json()['fields']['key'],b'%PDF-')
    submitted=client.post('/api/cloud/uploads/complete',json={'upload_id':ticket.json()['upload_id']},headers=headers)
    assert submitted.status_code==202
    tid=submitted.json()['task_id'];task=tasks.get(tid);sid=task['session_id']
    assert TestClient(app).get('/api/cloud/tasks/'+tid).status_code==401
    assert client.get('/api/cloud/tasks/'+tid).status_code==200
    # Clear the synthetic queued import, then verify a candidate with no QA cannot export.
    with db.connect() as con:con.execute("UPDATE processing_tasks SET state='failed' WHERE id=%s",(tid,))
    with state.scope():
        sess=state.get(sid)
        with state.job(sess):
            sess.analysis={'slide_count':1}
            sess.generation={'generation_id':'not-verified','state':'checking','checks':{}}
    denied=client.get(f'/api/sessions/{sid}/download?generation_id=not-verified')
    assert denied.status_code==409 and 'download_url' not in denied.json()
    assert client.get('/api/cloud/tick').status_code==401
    cancelled=client.post(f'/api/sessions/{sid}/finalize',json={},headers=headers)
    assert cancelled.status_code==200 and cancelled.json()['deleted']


def test_real_renderer_and_redesign_resume_after_interruption(shared,monkeypatch):
    """Real engine/render/storage/replay; synthetic model responses, never live QA evidence."""
    from app import api, generations, rendering
    from app.cloud import worker
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from rubric_fixtures import source_choice, passed_checks, qa_review
    from slide_engine.package import save_deck
    if not rendering.available():pytest.fail('This deployment check requires a real LibreOffice renderer.')
    sess=sessions.create()
    prs=Presentation()
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    box=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1))
    box.text='Synthetic deployment verification'
    box.text_frame.paragraphs[0].runs[0].font.size=Pt(28)
    save_deck(prs,sess.source_path)
    api.analyze_into(sess,'synthetic.pptx',Path(sess.source_path).read_bytes())
    calls=[]
    def model(role,system,payload,images=(),max_tokens=0):
        calls.append(role)
        assert all(Path(p).is_file() for _,p in images)
        if role=='output_qa':data=qa_review(payload)
        elif payload.get('stage')=='source_decisions':data=source_choice(payload)
        elif role=='reviewer':data={'verdict':'passed','summary':'Synthetic contract fixture','findings':[],'rubric':passed_checks()}
        else:raise AssertionError('Unexpected model role: '+role)
        return {'status':'completed','provider':'synthetic','model':'contract-only','data':data}
    monkeypatch.setattr(providers,'_generate_attempts',model)
    # Enable capabilities without sending a network request; transport is replaced above.
    monkeypatch.setenv('STEVENS_OFFLINE','0');monkeypatch.setenv('AI_GATEWAY_API_KEY','synthetic-not-a-real-key')
    tid=tasks.enqueue(sess,'redesign',{'mode':'ai','repair_passes':0},shared)
    first=tasks.claim(tid)
    original_step=replay.Replay.step
    def interrupted(self,*args,**kwargs):
        if self.ordinal==3:raise replay.SliceComplete()
        return original_step(self,*args,**kwargs)
    with monkeypatch.context() as patch:
        patch.setattr(replay.Replay,'step',interrupted)
        assert worker.run(tid,first['lease'])==75
    checkpoint_calls=len(calls)
    assert state.row(sess.id)['state']['generation'] is None
    with db.connect() as con:
        con.execute('UPDATE processing_tasks SET lease_until=0 WHERE id=%s',(tid,))
        con.execute('UPDATE processing_sessions SET lease_until=0 WHERE id=%s',(sess.id,))
    second=tasks.claim(tid)
    assert worker.run(tid,second['lease'])==0,tasks.get(tid)['error']
    task=tasks.get(tid)
    assert task['state']=='completed'
    record=task['result']['generation']
    assert all(record['checks'][name]['status']=='passed' for name in generations.QA_REQUIRED),record['findings']
    assert calls.count('element_roles')==1  # Completed source decision was replayed, not purchased twice.
    assert len(calls)>checkpoint_calls
    with db.connect() as con:
        con.execute('UPDATE processing_sessions SET lease=NULL,lease_until=0 WHERE id=%s',(sess.id,))
    with state.scope():
        restored=state.get(sess.id)
        with state.job(restored):
            record=restored.generation
            eligible=[f['id'] for f in record['findings'] if generations.can_approve(record,f)]
            if eligible:generations.decide(restored,generations.Decision(generation_id=record['generation_id'],candidate_sha256=record['candidate_sha256'],finding_ids=eligible,rationale='Synthetic test only; transport fixture is not a live visual assessment.'))
            assert generations.download_allowed(record),record['findings']
            data=api.artifact_bytes(restored,True,'pptx')
            assert __import__('hashlib').sha256(data).hexdigest()==record['candidate_sha256']
            assert api.artifact_bytes(restored,True,'pdf').startswith(b'%PDF-')
            Path(record['candidate']).write_bytes(data+b'tampered')
            with pytest.raises(ValueError):api.artifact_bytes(restored,True,'pptx')
