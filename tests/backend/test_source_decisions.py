from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt
import pytest
from app import grounded, generations
from app.ai import source_decisions, output_qa
from app.qa.artifact_coverage import audit
from slide_engine import inventory, template_policy as T
from rubric_fixtures import role_map, qa_review, source_choice
from test_ai_pipeline import ai_session, config, mocked_provider
REAL_QA_RUN=output_qa.run


def decision(source,action='redesign'):
    p=Presentation(source); digest=inventory.sha256(source)
    objects=source_decisions.source_objects(p,0,digest)
    value={**source_choice({'objects':objects,'source_slide':0}),'action':action,'matches_template':action=='keep_original',
           'reason':'Synthetic source-first decision.','remove_ids':[]}
    return p,digest,objects,value


def test_matching_slide_is_native_unchanged_and_tampering_is_detected(tmp_path):
    # A genuine bundled template slide: includes source theme, layout, master and art.
    p=Presentation(grounded.TEMPLATE_PATH)
    for sid,slide in list(zip(p.slides._sldIdLst,p.slides)):
        if slide.slide_layout.name=='1_Title Slide': continue
        p.part.drop_rel(sid.rId); p.slides._sldIdLst.remove(sid)
    source=tmp_path/'source.pptx'; p.save(source)
    _,_,_,choice=decision(source,'keep_original')
    candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions={'0':choice})
    result=Presentation(candidate)
    assert T.is_preserved(result.slides[0])
    assert T.unchanged_blob(result.slides[0].part)==T.unchanged_blob(Presentation(source).slides[0].part)
    assert audit(source,candidate,report)['status']=='passed'
    s=next(s for s in result.slides[0].shapes if s.has_text_frame)
    s.left+=Inches(.1); result.save(candidate)
    assert any(f['code']=='UNCHANGED_SLIDE_ALTERED' for f in audit(source,candidate,report)['findings'])


def image_source(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text='A required title'
    image=tmp_path/'decoration.png';Image.new('RGB',(200,100),'red').save(image)
    s.shapes.add_picture(str(image),Inches(1),Inches(3),Inches(4),Inches(2))
    source=tmp_path/'source.pptx';p.save(source)
    return source


@pytest.mark.parametrize('damage',['logo','content','unknown','text','dependency'])
def test_removal_cannot_exempt_protected_source_elements(tmp_path,damage):
    source=image_source(tmp_path);p,digest,objects,value=decision(source)
    art=next(e for e in value['elements'] if e['id']==objects[-1]['id'])
    art.update(role='decoration',confidence='high',content_bearing=False,contains_logo=False)
    value['remove_ids']=[art['id']]
    art['artwork_action']='remove'
    if damage=='logo':art['contains_logo']=True
    if damage=='content':art['content_bearing']=True
    if damage=='unknown':art['role']='unknown'
    if damage=='text':
        value['elements'][0].update(role='decoration',confidence='high',content_bearing=False)
        value['remove_ids']=[value['elements'][0]['id']]
    if damage=='dependency':value['elements'][0]['related_ids']=[art['id']]
    with pytest.raises(ValueError):source_decisions.validate(value,p,0,digest)


def test_removed_background_is_audited_and_protected_logo_survives(tmp_path):
    source=image_source(tmp_path);p,digest,objects,value=decision(source)
    art=value['elements'][-1];art.update(role='decoration',confidence='high',content_bearing=False,contains_logo=False)
    value['remove_ids']=[art['id']]
    art['artwork_action']='remove'
    candidate=tmp_path/'removed.pptx'; report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    assert audit(source,candidate,report)['status']=='passed'
    assert not any(s.shape_type==13 for s in Presentation(candidate).slides[0].shapes)
    retained=tmp_path/'retained.pptx'; other=grounded.build_deck(source,retained)
    other['source_decisions']={'0':value}
    assert any(f['code']=='REJECTED_ARTWORK_RETAINED' for f in audit(source,retained,other)['findings'])
    value['remove_ids']=[];art.update(role='logo',contains_logo=True,content_bearing=True,artwork_action='retain')
    report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    actual=next(s for s in Presentation(candidate).slides[0].shapes if s.shape_type==13)
    original=next(s for s in p.slides[0].shapes if s.shape_type==13)
    assert actual.image.blob==original.image.blob
    assert audit(source,candidate,report)['status']=='passed'


def test_redesign_qa_pairs_originals_in_output_order_including_splits(tmp_path,monkeypatch):
    # Bypass only the legacy autouse QA mock; exercise the real QA orchestrator.
    module=output_qa
    p=Presentation(); pages=[]
    for i in range(3):
        p.slides.add_slide(p.slide_layouts[6])
        png=tmp_path/f'slide-{i}.png';Image.new('RGB',(60,40),'white').save(png)
        pages.append({'output_slide':i,'png':str(png),'success':True})
    for i in range(2):Image.new('RGB',(60,40),'red').save(tmp_path/f'before-{i}.png')
    candidate=tmp_path/'candidate.pptx';p.save(candidate)
    original_deck=Presentation()
    for i in range(2):
        s=original_deck.slides.add_slide(original_deck.slide_layouts[6])
        s.notes_slide.notes_text_frame.text=f'Original note {i}'
    source=tmp_path/'source.pptx';original_deck.save(source)
    record={'mode':'ai','candidate':str(candidate),'report':{},'source_to_output_slides':{'0':[0,1],'1':[2]},
            'checks':{'render_verification':{'candidate_sha256':inventory.sha256(candidate),'pages':pages}}}
    reference=tmp_path/'reference.png';Image.new('RGB',(60,40),'blue').save(reference)
    record['ai_pipeline']={'template_references':[{'layout':'Blank','label':'APPROVED TEMPLATE: Blank',
                              'image':str(reference),'sha256':inventory.sha256(reference)}]}
    sess=SimpleNamespace(source_path=source,ensure_active=lambda:None,preview_path=lambda i,v:str(tmp_path/f'before-{i}.png'))
    requests=[]
    def provider(role,system,payload,images,**kwargs):
        requests.append(images)
        assert [Path(path).name for _,path in images]==['reference.png','before-0.png','slide-0.png','before-0.png','slide-1.png','before-1.png','slide-2.png']
        assert payload['slides'][0]['note_comparison']['expected_on_this_output']=='Original note 0'
        assert payload['slides'][1]['note_comparison']['expected_on_this_output']==''
        return {'status':'completed','data':qa_review(payload)}
    monkeypatch.setattr(module.providers,'generate',provider)
    assert all(c['status']=='passed' for c in REAL_QA_RUN(sess,record).values())
    assert len(requests)==2
    Image.new('RGB',(60,40),'green').save(reference)
    assert REAL_QA_RUN(sess,record)['output_qa_coverage']['status']=='error'
    (tmp_path/'before-1.png').unlink()
    assert REAL_QA_RUN(sess,record)['output_qa_coverage']['status']=='error'


@pytest.mark.renderer
def test_unchanged_template_slide_renders_identically(tmp_path):
    from app.qa.render_verify import check
    from PIL import ImageChops
    p=Presentation(grounded.TEMPLATE_PATH)
    for sid,slide in list(zip(p.slides._sldIdLst,p.slides)):
        if slide.slide_layout.name=='1_Title Slide': continue
        p.part.drop_rel(sid.rId);p.slides._sldIdLst.remove(sid)
    source=tmp_path/'source.pptx';p.save(source)
    _,_,_,choice=decision(source,'keep_original')
    candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions={'0':choice})
    assert audit(source,candidate,report)['status']=='passed'
    check(source,tmp_path/'before'); check(candidate,tmp_path/'after')
    with Image.open(tmp_path/'before/slide-0.png') as before,Image.open(tmp_path/'after/slide-0.png') as after:
        assert before.size==after.size
        assert ImageChops.difference(before.convert('RGB'),after.convert('RGB')).getbbox() is None


def test_qa_can_reverse_an_incorrect_artwork_removal(ai_session,tmp_path,monkeypatch):
    from app.ai import pipeline
    p=Presentation(ai_session.source_path)
    png=tmp_path/'meaningful.png';Image.new('RGB',(160,90),'blue').save(png)
    p.slides[0].shapes.add_picture(str(png),Inches(1),Inches(4),Inches(2),Inches(1.125))
    p.save(ai_session.source_path)
    requests=[]
    def provider(role,system,payload,*args,**kwargs):
        result=mocked_provider(role,system,payload,*args,**kwargs)
        if payload.get('stage')=='source_decisions' and payload['source_slide']==0:
            requests.append(payload)
            image=next(o for o in payload['objects'] if o['kind']=='picture')
            entry=next(e for e in result['data']['elements'] if e['id']==image['id'])
            entry.update(role='image',confidence='high',content_bearing=bool(payload['findings']),contains_logo=False)
            entry['artwork_action']='retain' if payload['findings'] else 'remove'
            result['data'].update(action='redesign',matches_template=False,reason='Synthetic decision reversal.',
                remove_ids=[] if payload['findings'] else [image['id']])
        return result
    def qa(sess,record,evidence):
        prs=Presentation(record['candidate'])
        missing=not any(s.shape_type==13 for s in prs.slides[0].shapes)
        result={n:{'status':'passed','findings':[]} for n in output_qa.CHECKS+('output_qa_visual',)}
        if missing:result['output_qa_accuracy']={'status':'failed','findings':[{'code':'OUTPUT_ACCURACY',
            'criterion':'content_presence','severity':'blocking','output_slide':0,'message':'Restore the meaningful source picture.'}]}
        return result
    monkeypatch.setattr(pipeline.providers,'generate',provider)
    monkeypatch.setattr(output_qa,'run',qa)
    record=generations.build(ai_session,mode='ai',repair_passes=1)
    assert len(requests)==2 and requests[-1]['findings']
    assert record['output_qa_repairs'][0]['accepted'],record['output_qa_repairs']
    assert any(s.shape_type==13 for s in Presentation(record['candidate']).slides[0].shapes)
    assert record['checks']['artifact_coverage']['status']=='passed'


def test_first_page_type_and_explicit_picture_decision_are_required(tmp_path):
    source=image_source(tmp_path);p,digest,objects,value=decision(source)
    value['slide_kind']='content'
    with pytest.raises(ValueError,match='first source slide'):source_decisions.validate(value,p,0,digest)
    value['slide_kind']='cover';value['remove_ids']=[objects[-1]['id']]
    with pytest.raises(ValueError,match='explicit element'):source_decisions.validate(value,p,0,digest)


def test_invalid_source_role_gets_one_bounded_correction(ai_session):
    attempts=[]
    def provider(role,system,payload,*args,**kwargs):
        attempts.append(payload.get('validation_error'))
        result=mocked_provider(role,system,payload,*args,**kwargs)
        if len(attempts)==1:result['data']['elements'][0]['alignment_reference']='parent'
        return result
    choices,calls=source_decisions.run(ai_session,lambda **kw:None,generate=provider,indices={0})
    assert len(calls)==2 and attempts[0] is None and 'parent' in attempts[1]
    assert choices['0']['slide_kind']=='cover'


def test_cover_uses_semantics_instead_of_largest_wordmark(tmp_path):
    source=image_source(tmp_path);p=Presentation(source)
    logo=p.slides[0].shapes.add_textbox(Inches(8),Inches(.3),Inches(2),Inches(.5))
    logo.text='STEVENS';logo.text_frame.paragraphs[0].runs[0].font.size=Pt(60)
    p.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size=Pt(16)
    p.save(source);_,_,objects,value=decision(source)
    next(e for e in value['elements'] if e['id']==objects[-1]['id']).update(role='logo',contains_logo=True)
    art=next(e for e in value['elements'] if e['id']==objects[-2]['id'])
    art.update(role='background',confidence='high',content_bearing=False,artwork_action='remove')
    value['remove_ids']=[art['id']]
    candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions={'0':value})
    output=Presentation(candidate).slides[0]
    title=next(s for s in output.shapes if s.has_text_frame and s.text=='A required title')
    wordmark=next(s for s in output.shapes if s.has_text_frame and s.text=='STEVENS')
    assert title.name.endswith('|title') and wordmark.name.endswith('|logo')
    assert T.contains([v/T.EMU for v in (title.left,title.top,title.width,title.height)],T.COVER_TITLE)
    assert not any(s.shape_type==13 for s in output.shapes)
    assert audit(source,candidate,report)['status']=='passed'


@pytest.mark.parametrize('role,damage',[('code','indent'),('code','font'),('logo','font')])
def test_protected_code_and_logo_typography_survive_and_tampering_blocks(tmp_path,role,damage):
    from slide_engine.text_style import protected_text
    from app.ai import layout
    from test_ai_pipeline import identity_plan
    p=Presentation();p.slides.add_slide(p.slide_layouts[6])
    s=p.slides.add_slide(p.slide_layouts[6])
    shape=s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(3))
    shape.text='if ready:\n    execute()\n\n    finish()' if role=='code' else 'STEVENS\nINSTITUTE OF TECHNOLOGY'
    for para in shape.text_frame.paragraphs:
        for run in para.runs:run.font.name='Consolas' if role=='code' else 'Georgia';run.font.size=Pt(18)
    source=tmp_path/'source.pptx';p.save(source);digest=inventory.sha256(source)
    objects=source_decisions.source_objects(p,1,digest)
    choice=source_choice({'source_slide':1,'objects':objects});choice['elements'][0]['role']=role
    candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate,source_decisions={'1':choice})
    out=Presentation(candidate);copied=next(s for s in out.slides[1].shapes if s.name.endswith('|'+role))
    assert protected_text(shape)==protected_text(copied)
    assert audit(source,candidate,report)['status']=='passed'
    plan=identity_plan(layout.describe(out.slides[1]));plan['objects'][0]['color']='red'
    with pytest.raises(ValueError,match='protected'):
        layout.validate(layout.LayoutPlan.model_validate(plan),out.slides[1],out.slide_width/T.EMU,out.slide_height/T.EMU)
    if damage=='indent':copied.text_frame.paragraphs[1].runs[0].text='execute()'
    else:copied.text_frame.paragraphs[0].runs[0].font.name='Arial'
    out.save(candidate)
    assert any(f['code']=='PROTECTED_TYPOGRAPHY_ALTERED' for f in audit(source,candidate,report)['findings'])


@pytest.mark.parametrize('difference',[1,10000])
def test_keep_original_canvas_rounding_is_validated_early(tmp_path,difference):
    p=Presentation(grounded.TEMPLATE_PATH)
    for sid,slide in list(zip(p.slides._sldIdLst,p.slides)):
        if slide.slide_layout.name=='1_Title Slide': continue
        p.part.drop_rel(sid.rId);p.slides._sldIdLst.remove(sid)
    p.slide_width-=difference
    source=tmp_path/'rounded.pptx';p.save(source)
    p,digest,objects,value=decision(source,'keep_original')
    if difference==1:
        source_decisions.validate(value,p,0,digest)
        report=grounded.build_deck(source,tmp_path/'result.pptx',source_decisions={'0':value})
        assert audit(source,tmp_path/'result.pptx',report)['status']=='passed'
    else:
        with pytest.raises(ValueError,match='canvas differs'):
            source_decisions.validate(value,p,0,digest)
