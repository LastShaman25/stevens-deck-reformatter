from copy import deepcopy
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from app import grounded, generations
from app.ai import layout, output_qa, pipeline
from app.qa.artifact_coverage import audit
from slide_engine import template_policy as T
from test_ai_pipeline import ai_session, config, mocked_provider, identity_plan
from test_regressions import fixture
from rubric_fixtures import repair_evidence


def test_cover_extracts_ordinary_title_and_interior_preserves_contrast(tmp_path):
    source=fixture(tmp_path); p=Presentation(source)
    title=p.slides[0].shapes[0]
    title.text='Native cover title'
    title.text_frame.paragraphs[0].runs[0].font.size=Pt(32)
    assert p.slides[0].shapes.title is None
    middle=p.slides[1]; middle.background.fill.solid()
    middle.background.fill.fore_color.rgb=RGBColor.from_string('A22537')
    for s in middle.shapes:
        if s.has_text_frame:
            for para in s.text_frame.paragraphs:
                for run in para.runs: run.font.color.rgb=RGBColor(255,255,255)
    p.save(source); candidate=tmp_path/'candidate.pptx'
    report=grounded.build_deck(source,candidate); result=Presentation(candidate)
    assert T.is_cover(result.slides[0])
    cover_title=next(s for s in result.slides[0].shapes if s.text=='Native cover title')
    assert cover_title.name.endswith('|title')
    assert T.contains([v/T.EMU for v in (cover_title.left,cover_title.top,cover_title.width,cover_title.height)],T.COVER_TITLE)
    panel=result.slides[1].shapes[0]
    assert panel.name==T.BACKGROUND_NAME and str(panel.fill.fore_color.rgb)=='A22537'
    assert T.check(candidate)['status']=='passed'
    assert audit(source,candidate,report)['status']=='passed'
    panel.fill.fore_color.rgb=RGBColor(255,255,255); result.save(candidate)
    assert any(f['code']=='SOURCE_BACKGROUND_ALTERED' for f in audit(source,candidate,report)['findings'])


def test_logo_intrusion_rejected_before_edit_and_after_export(tmp_path):
    source=fixture(tmp_path); candidate=tmp_path/'candidate.pptx'
    grounded.build_deck(source,candidate); p=Presentation(candidate); slide=p.slides[1]
    value=identity_plan(layout.describe(slide)); value['objects'][0].update(x=.1,y=6.9,w=2,h=.3)
    with pytest.raises(ValueError,match='protected logo'):
        layout.validate(layout.LayoutPlan.model_validate(value),slide,p.slide_width/T.EMU,p.slide_height/T.EMU)
    slide.shapes[0].left=Inches(.1); slide.shapes[0].top=Inches(6.9); p.save(candidate)
    assert any(f['code']=='TEMPLATE_CONTENT_BOUNDS' and f['output_slide']==1 for f in pipeline.structural(candidate)['findings'])


@pytest.mark.parametrize('outcome',['fixed','unchanged','error'])
def test_final_qa_routes_only_failed_slides_and_rechecks_all(ai_session,monkeypatch,outcome):
    reviewed=[]; planned=[]; repair_planned=[]
    def qa(sess,record,evidence):
        p=Presentation(record['candidate'])
        reviewed.append([s._element.xml for s in p.slides])
        if len(reviewed)>1:
            assert record['repair_acceptance_checks'][0]['acceptance_condition']
        checks={n:{'status':'passed','findings':[]} for n in output_qa.CHECKS+('output_qa_visual',)}
        if len(reviewed)==1 or outcome=='unchanged':
            checks['output_qa_visual']={'status':'failed','findings':[{**repair_evidence(),'code':'OUTPUT_VISUAL',
                'criterion':'spatial_layout','severity':'blocking','affected_slides':[1],
                'message':'Move this body slightly right inside its box.','evidence':'Synthetic positional defect.'}]}
        elif outcome=='error': checks['output_qa_coverage']={'status':'error','findings':[]}
        return checks
    def provider(role,system,payload,*args,**kwargs):
        result=mocked_provider(role,system,payload,*args,**kwargs)
        if role=='planner':
            planned.append(payload['slide'])
            if any(f.get('from_check')=='output_qa_visual' for f in payload['findings']):
                repair_planned.append(payload['slide'])
                assert payload['slide']==1
                assert any(f['criterion']=='spatial_layout' and f['evidence'] for f in payload['findings'] if f.get('from_check'))
                assert any(f.get('required_correction') and f.get('acceptance_condition') for f in payload['findings'])
                result['data']['objects'][0]['x']+=.003
        if role=='reviewer' and repair_planned:
            incoming=[f for f in payload['repair_acceptance_checks'] if f.get('from_check')]
            if incoming:
                assert incoming[0]['acceptance_condition']
        return result
    monkeypatch.setattr(output_qa,'run',qa)
    monkeypatch.setattr(pipeline.providers,'generate',provider)
    record=generations.build(ai_session,mode='ai',repair_passes=1)
    assert planned==[1]  # Already-passing slides are never sent to the planner.
    assert repair_planned==[1]
    assert len(reviewed)==2 and len(reviewed[1])==3
    assert reviewed[0][0]==reviewed[1][0] and reviewed[0][2]==reviewed[1][2]
    assert reviewed[0][1]!=reviewed[1][1]
    repair=record['output_qa_repairs'][0]
    assert repair['targets']==[1] and repair['accepted']==(outcome=='fixed')
    if outcome=='fixed':
        assert not any(f['code']=='OUTPUT_VISUAL' for f in record['findings'])
        assert record['candidate_sha256']==repair['candidate_sha256']
    else:
        assert record['state']=='failed'
        assert record['candidate_sha256']==repair['before_sha256']


def test_replacing_check_drops_only_its_stale_findings():
    r={'checks':{},'findings':[]}
    generations.add_check(r,'first',{'status':'failed','findings':[{'code':'OLD'}]})
    generations.add_check(r,'other',{'status':'failed','findings':[{'code':'KEEP'}]})
    generations.add_check(r,'first',{'status':'passed','findings':[]})
    assert [f['code'] for f in r['findings']]==['KEEP']


def test_final_qa_repair_respects_remaining_call_budget(ai_session,monkeypatch):
    monkeypatch.setenv('STEVENS_AI_MAX_CALLS','14')
    monkeypatch.setattr(pipeline,'structural',lambda path:{'status':'passed','findings':[]})
    calls=[]
    def provider(role,*args,**kwargs):
        calls.append(role); return mocked_provider(role,*args,**kwargs)
    monkeypatch.setattr(pipeline.providers,'generate',provider)
    def qa(*args):
        result={n:{'status':'passed','findings':[]} for n in output_qa.CHECKS+('output_qa_visual',)}
        result['output_qa_visual']={'status':'failed','findings':[{'code':'OUTPUT_VISUAL','criterion':'spatial_layout',
            'severity':'blocking','output_slide':1,'message':'Synthetic defect remains.'}]}
        return result
    monkeypatch.setattr(output_qa,'run',qa)
    record=generations.build(ai_session,mode='ai',repair_passes=1)
    assert len(calls)==9  # Baseline and unchanged-slide evidence are reused; edited slide is reviewed.
    assert record['output_qa_repairs'][0]['ai_pipeline']['baseline_review_reused']['reviewed_slides']==3
    assert record['state']=='failed'
    assert not record['output_qa_repairs'][0]['accepted']
    assert record['candidate_sha256']==record['output_qa_repairs'][0]['before_sha256']


def test_mandatory_repairs_continue_until_second_recheck_passes(ai_session,monkeypatch):
    rounds=[]
    monkeypatch.setattr(pipeline,'structural',lambda path:{'status':'passed','findings':[]})
    def qa(sess,record,evidence):
        rounds.append(record['candidate_sha256'])
        checks={n:{'status':'passed','findings':[]} for n in output_qa.CHECKS+('output_qa_visual',)}
        remaining=max(0,3-len(rounds))
        if remaining:
            checks['output_qa_visual']={'status':'needs_review','findings':[
                {**repair_evidence(),'code':f'SPACING_{i}','criterion':'spatial_layout',
                 'severity':'review','output_slide':1,'message':f'Correct gap {i}.'}
                for i in range(remaining)]}
        return checks
    def provider(role,system,payload,*a,**kw):
        result=mocked_provider(role,system,payload,*a,**kw)
        if role=='planner' and any(f.get('from_check')=='output_qa_visual' for f in payload['findings']):
            result['data']['objects'][0]['x']+=.003
        return result
    monkeypatch.setattr(output_qa,'run',qa)
    monkeypatch.setattr(pipeline.providers,'generate',provider)
    # Even a legacy zero setting cannot disable mandatory final QA repair.
    record=generations.build(ai_session,mode='ai',repair_passes=0)
    assert len(rounds)==3 and len(set(rounds))==3
    assert len(record['output_qa_repairs'])==2
    assert all(r['accepted'] for r in record['output_qa_repairs'])
    assert generations.qa_passed(record) and generations.download_allowed(record)


@pytest.mark.parametrize('always_reject',[False,True])
def test_rejected_render_feedback_is_retried_without_releasing_it(ai_session,monkeypatch,always_reject):
    from rubric_fixtures import passed_checks
    planned=[];final_reviews=[]
    monkeypatch.setattr(pipeline,'structural',lambda path:{'status':'passed','findings':[]})
    def qa(sess,record,evidence):
        final_reviews.append(record['candidate_sha256'])
        result={n:{'status':'passed','findings':[]} for n in output_qa.CHECKS+('output_qa_visual',)}
        if len(final_reviews)==1:
            result['output_qa_visual']={'status':'needs_review','findings':[
                {**repair_evidence(),'code':'SPACING','criterion':'spatial_layout',
                 'severity':'review','output_slide':1,'message':'Separate this paragraph.'}]}
        return result
    def provider(role,system,payload,*a,**kw):
        result=mocked_provider(role,system,payload,*a,**kw)
        if role=='planner':
            planned.append(payload['findings'])
            result['data']['objects'][0]['x']+=.03*len(planned)
        if role=='reviewer' and planned and (always_reject or len(planned)==1):
            f={**repair_evidence(),'criterion':'spatial_layout','severity':'blocking','message':'The proposed repair creates a collision.'}
            result['data'].update(verdict='failed',findings=[f],rubric=passed_checks([f]))
        return result
    monkeypatch.setattr(output_qa,'run',qa)
    monkeypatch.setattr(pipeline.providers,'generate',provider)
    record=generations.build(ai_session,mode='ai',repair_passes=0)
    assert len(planned)==(3 if always_reject else 2)
    assert any(f.get('from_check')=='rejected_repair' and 'collision' in f['message'] for f in planned[1])
    assert not record['output_qa_repairs'][0]['accepted']
    if always_reject:
        assert not generations.checks_satisfied(record)
        assert len(final_reviews)==1
        assert record['candidate_sha256']==record['output_qa_repairs'][0]['before_sha256']
    else:
        assert record['output_qa_repairs'][1]['accepted']
        assert len(final_reviews)==2
        assert generations.download_allowed(record)


def test_first_page_layout_is_checked_even_for_preserved_slide(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    s._element.cSld.set('name','sss:preserved:test')
    path=tmp_path/'wrong-cover.pptx';p.save(path)
    assert T.check(path)['findings'][0]['code']=='FIRST_PAGE_TEMPLATE'


def test_planner_cannot_introduce_text_stacking_or_overflow(tmp_path):
    p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
    for i,y in enumerate((1,2)):
        box=s.shapes.add_textbox(Inches(1),Inches(y),Inches(5),Inches(.5))
        box.name=f'text-{i}';box.text='A short equation annotation'
        box.text_frame.paragraphs[0].runs[0].font.size=Pt(14)
    value=identity_plan(layout.describe(s))
    value['objects'][1]['y']=1.15
    with pytest.raises(ValueError,match='Text collision'):
        layout.validate(layout.LayoutPlan.model_validate(value),s,13.33,7.5)
    value=identity_plan(layout.describe(s));value['objects'][0]['font_size']=55
    with pytest.raises(ValueError,match='Text fit worsens'):
        layout.validate(layout.LayoutPlan.model_validate(value),s,13.33,7.5)


def test_baseline_preserves_small_line_font_proportions(tmp_path):
    p=Presentation();cover=p.slides.add_slide(p.slide_layouts[6])
    cover.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Mathematics'
    s=p.slides.add_slide(p.slide_layouts[6])
    box=s.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(.2))
    box.text_frame.word_wrap=False
    para=box.text_frame.paragraphs[0]
    for text,size in [('x',10),('2',6)]:
        r=para.add_run();r.text=text;r.font.size=Pt(size)
    source=tmp_path/'small-lines.pptx';p.save(source)
    candidate=tmp_path/'candidate.pptx';grounded.build_deck(source,candidate)
    result=next(sh for sh in Presentation(candidate).slides[1].shapes if sh.has_text_frame and sh.text=='x2')
    sizes=[r.font.size.pt for r in result.text_frame.paragraphs[0].runs]
    assert max(sizes)<12 and sizes[0]/sizes[1]==pytest.approx(10/6,abs=.01)
    assert result.text_frame.word_wrap is False
