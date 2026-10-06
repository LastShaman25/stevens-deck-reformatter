from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from fastapi.testclient import TestClient
import pytest
from app import grounded, sessions, generations
from app.main import app
from app.qa.artifact_coverage import audit
from slide_engine import templates,template_policy as T


def source_deck(tmp_path):
    p=Presentation()
    for title in ('Source title','Source content'):
        s=p.slides.add_slide(p.slide_layouts[6])
        s.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(1)).text=title
        s.shapes.add_textbox(Inches(1),Inches(3),Inches(6),Inches(1)).text='Details stay intact.'
    path=tmp_path/'source.pptx';p.save(path)
    return path


@pytest.mark.parametrize('template_id,names',[
    ('cpe',['Title slide 1','1-line - 1 tect box','end slide']),
    ('stevens',['Title Slide','Title Only','1_Title Slide']),
])
def test_upload_routes_build_and_coverage_to_selected_template(tmp_path,template_id,names):
    source=source_deck(tmp_path)
    response=TestClient(app).post('/api/sessions',files={'file':('source.pptx',source.read_bytes())},data={'template_id':template_id})
    assert response.status_code==200,response.text
    assert response.json()['template_id']==template_id
    sess=sessions.get(response.json()['session_id'])
    try:
        with sessions.job(sess):
            output=tmp_path/'candidate.pptx'
            report=grounded.build_deck(sess.source_path,output,require_closing=True)
            assert [s.slide_layout.name for s in Presentation(output).slides]==names
            assert T.check(output)['status']=='passed'
            assert audit(sess.source_path,output,report)['status']=='passed'
            assert templates.current_id()==template_id
    finally:sessions.delete(sess.id)
    assert templates.current_id()=='stevens'


def test_template_context_is_isolated_between_concurrent_jobs():
    barrier=Barrier(2)
    def inspect(ident):
        with templates.use(ident):
            barrier.wait(timeout=5)
            return templates.current_id(),T.OPENING_LAYOUT,T.COVER_TITLE
    with ThreadPoolExecutor(2) as pool:
        a=pool.submit(inspect,'cpe');b=pool.submit(inspect,'stevens')
        assert a.result()==('cpe','Title slide 1',(.73,3.05,7.84,2.25))
        assert b.result()[0:2]==('stevens','Title Slide')
    assert templates.current_id()=='stevens'


def test_unknown_format_rejected_before_upload(tmp_path):
    response=TestClient(app).post('/api/sessions',files={'file':('source.pptx',source_deck(tmp_path).read_bytes())},data={'template_id':'unknown'})
    assert response.status_code==422


def test_cpe_prompts_use_selected_artwork_and_footer():
    from app.ai import pipeline,source_decisions,output_qa,rubric
    with templates.use('cpe'):
        for system in (pipeline.PLANNER_SYSTEM,source_decisions.SYSTEM,output_qa.SYSTEM,rubric.QA,rubric.GENERATOR):
            prompt=templates.prompt(system)
            assert 'Selected format: CPE.' in prompt
            assert 'mostly red' not in prompt
            assert 'statue-photo' not in prompt
            assert 'bottom-left' not in prompt
            assert templates.prompt(prompt)==prompt


def test_cpe_authoring_uses_selected_layouts_and_native_bounds(tmp_path):
    from app.authoring import service,composer
    from app.authoring.models import CreationRequest,DeckSpec,SlideSpec
    sess=service.create(CreationRequest(topic='Example',audience='Students',template_id='cpe'))
    try:
        with sessions.job(sess):
            spec=DeckSpec(slides=[SlideSpec(id=k,title=k.title(),bullets=['Details']) for k in ['opening','content','closing']])
            output=tmp_path/'authored.pptx'
            manifest=composer.compose(spec,output,tmp_path/'assets',kinds=['opening','content','closing'])
            assert [s.slide_layout.name for s in Presentation(output).slides]==['Title slide 1','1-line - 1 tect box','end slide']
            assert composer.audit(output,manifest)['status']=='passed'
            assert T.check(output)['status']=='passed'
            assert not any(s.has_text_frame and s.text.strip()=='Presenter:' for s in Presentation(output).slides[0].slide_layout.shapes)
    finally:sessions.delete(sess.id)
