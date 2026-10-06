"""Regressions reproduced from the October timing report, without live AI."""
from copy import deepcopy
from types import SimpleNamespace
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches
from lxml import etree
from app import grounded
from app.ai import output_qa, providers, layout
from app.authoring import composer
from app.authoring.models import DeckSpec, SlideSpec
from app.qa.artifact_coverage import audit
from slide_engine import preserve
from test_implementation import qa_fixture
from rubric_fixtures import qa_review, repair_evidence

REAL_QA = output_qa.run


@pytest.mark.parametrize('kind', [XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.LINE, XL_CHART_TYPE.PIE])
@pytest.mark.parametrize('explicit', [False, True])
def test_chart_source_palette_is_preserved_and_tampering_fails(tmp_path, kind, explicit):
    p = Presentation(); p.slides.add_slide(p.slide_layouts[6])
    slide = p.slides.add_slide(p.slide_layouts[6])
    data = CategoryChartData(); data.categories = ['A', 'B']
    data.add_series('Positive', [4, 8]); data.add_series('Negative', [8, 4])
    chart = slide.shapes.add_chart(kind, Inches(1), Inches(1), Inches(6), Inches(3), data).chart
    if explicit:
        for series, color in zip(chart.series, [MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.ACCENT_2]):
            series.format.fill.solid(); series.format.fill.fore_color.theme_color = color
    source = tmp_path/'source.pptx'; candidate = tmp_path/'candidate.pptx'; p.save(source)
    report = preserve.build(source, candidate, grounded.TEMPLATE_PATH)
    assert audit(source, candidate, report)['status'] == 'passed'
    result = Presentation(candidate)
    chart = next(s.chart for s in result.slides[1].shapes if s.has_chart)
    override = next(r.target_part for r in chart.part.rels.values() if r.reltype == RT.THEME_OVERRIDE)
    theme = etree.fromstring(override.blob)
    assert theme.find('.//'+qn('a:accent1'))[0].get('val') == '4F81BD'
    assert theme.find('.//'+qn('a:accent2'))[0].get('val') == 'C0504D'
    theme.find('.//'+qn('a:accent1'))[0].set('val', 'FFFFFF')
    override._blob = etree.tostring(theme)
    result.save(candidate)
    assert audit(source, candidate, report)['status'] == 'failed'


def test_table_grid_and_caption_stay_consistent_after_preservation_and_ai_resize(tmp_path):
    p = Presentation(); p.slide_width = Inches(13.333333); p.slide_height = Inches(7.5)
    p.slides.add_slide(p.slide_layouts[6]); slide = p.slides.add_slide(p.slide_layouts[6])
    shape = slide.shapes.add_table(5, 2, Inches(1), Inches(1), Inches(10), Inches(4))
    shape.table.columns[0].width = Inches(3); shape.table.columns[1].width = Inches(7)
    shape.table.cell(0,0).merge(shape.table.cell(0,1))
    for i in range(5): shape.table.cell(i,0).text = str(i)
    slide.shapes.add_textbox(Inches(1), Inches(5.05), Inches(10), Inches(.5)).text = 'Caption'
    source=tmp_path/'source.pptx'; candidate=tmp_path/'candidate.pptx'; p.save(source)
    report=preserve.build(source,candidate,grounded.TEMPLATE_PATH)
    slide=Presentation(candidate).slides[1]
    table=next(s for s in slide.shapes if s.has_table)
    caption=next(s for s in slide.shapes if s.has_text_frame)
    assert table.top + sum(r.height for r in table.table.rows) < caption.top
    assert sum(c.width for c in table.table.columns) == table.width
    assert table.table.columns[0].width/table.width == pytest.approx(.3, abs=1e-6)
    edits=[layout.Edit(id=ident,x=box[0],y=box[1],w=box[2]*.8 if s.has_table else box[2],
                       h=box[3]*.8 if s.has_table else box[3]) for ident,s,box,_,_ in layout.nodes(slide)]
    updated=tmp_path/'updated.pptx'
    updated_report,_=layout.apply(candidate,updated,{1:layout.LayoutPlan(layout='Title Only',rationale='Resize table',objects=edits)},report)
    table=next(s for s in Presentation(updated).slides[1].shapes if s.has_table)
    assert sum(r.height for r in table.table.rows) == table.height
    assert sum(c.width for c in table.table.columns) == table.width
    assert audit(source,updated,updated_report)['status']=='passed'


def test_long_permitted_closing_is_fitted_without_changing_text(tmp_path):
    bullets=['Apply interdisciplinary methodologies to strengthen collaborative engineering decisions.',
             'Evaluate implementation alternatives through systematic experimentation and reflection.']
    assert sum(map(len,bullets)) == 175
    deck=DeckSpec(slides=[SlideSpec(id='a',title='Topic'),SlideSpec(id='b',title='Thank you!',bullets=bullets)])
    candidate=tmp_path/'closing.pptx'
    manifest=composer.compose(deck,candidate,tmp_path/'assets',kinds=['opening','closing'])
    body=next(s for s in Presentation(candidate).slides[-1].shapes if s.name=='authored-body')
    assert body.text=='\n'.join(bullets)
    assert 16 <= body.text_frame.paragraphs[0].font.size.pt < 20
    assert composer.audit(candidate,manifest)['status']=='passed'


@pytest.mark.parametrize('count',[5,6,10,11])
def test_qa_batch_scope_and_global_closing_checks(tmp_path, monkeypatch, count):
    record=qa_fixture(tmp_path,count); seen=[]
    finding={**repair_evidence(),'slides':[count],'criterion':'structure_sequence','severity':'blocking',
             'accuracy':'not_applicable','message':'Final slide is not the required closing.'}
    def provider(role,system,payload,**kwargs):
        scope=payload['review_scope']; seen.append(deepcopy(scope))
        assert scope['visible_ordinals']==payload['expected']
        assert scope['total_slides']==count
        assert scope['includes_closing_position']==(count in payload['expected'])
        assert 'Never infer a missing opening' in system
        return {'status':'completed','data':qa_review(payload,[finding] if scope['whole_deck_checks'] else [])}
    monkeypatch.setattr(providers,'generate',provider)
    checks=REAL_QA(SimpleNamespace(ensure_active=lambda:None),record)
    assert checks['output_qa_sequence']['status']=='failed'
    assert checks['output_qa_coverage']['status']=='passed'
    assert seen[-1]['whole_deck_checks'] and all(not s['whole_deck_checks'] for s in seen[:-1])


def test_unseen_slide_findings_still_fail_after_correction(tmp_path,monkeypatch):
    record=qa_fixture(tmp_path,6)
    finding={**repair_evidence(),'slides':[6],'criterion':'structure_sequence','severity':'blocking',
             'accuracy':'not_applicable','message':'Unseen closing'}
    monkeypatch.setattr(providers,'generate',lambda role,system,payload,**kw:
        {'status':'completed','data':qa_review(payload,[finding])})
    checks=REAL_QA(SimpleNamespace(ensure_active=lambda:None),record)
    assert checks['output_qa_coverage']['status']=='error'
    assert len(record['output_qa_validation_errors'])==2


def test_chart_theme_copy_does_not_mutate_shared_source_part():
    from slide_engine.chart_theme import preserved_part
    from slide_engine.inventory import relationship_signature
    p=Presentation(); s=p.slides.add_slide(p.slide_layouts[6])
    data=CategoryChartData(); data.categories=['A']; data.add_series('One',[1])
    shape=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0,Inches(5),Inches(3),data)
    before=relationship_signature(shape.chart.part)
    clone=preserved_part(shape)
    assert clone is not shape.chart.part
    assert clone.blob==shape.chart.part.blob
    assert relationship_signature(shape.chart.part)==before
    # A shared dependency can already have a new package name while the
    # relationship remembers its old relative target. Do not lose the workbook.
    from pptx.opc.packuri import PackURI
    workbook=shape.chart.part.chart_workbook.xlsx_part
    workbook._partname=PackURI('/ppt/embeddings/renamed.xlsx')
    clone=preserved_part(shape)
    assert clone.chart_workbook.xlsx_part is workbook


def test_table_grid_mismatch_is_blocking():
    from app.qa.style_audit import audit as style_audit
    p=Presentation(); s=p.slides.add_slide(p.slide_layouts[6])
    shape=s.shapes.add_table(2,2,0,0,Inches(5),Inches(2))
    shape.height=Inches(1)
    assert any(f['type']=='table_geometry' and f['sev']=='fail' for f in style_audit(s,p.slide_width,p.slide_height))


def test_unfittable_bookend_is_rejected_without_truncation(tmp_path):
    bullets=['\n'.join(['word']*20)]
    deck=DeckSpec(slides=[SlideSpec(id='a',title='Topic'),SlideSpec(id='b',title='Thank you!',bullets=bullets)])
    with pytest.raises(ValueError,match='cannot fit readably'):
        composer.compose(deck,tmp_path/'deck.pptx',tmp_path/'assets',kinds=['opening','closing'])
    assert deck.slides[-1].bullets==bullets


@pytest.mark.renderer
def test_real_renderer_preserves_chart_colors_and_fits_closing(tmp_path):
    import fitz
    from PIL import Image
    from app import rendering
    if not rendering.available(): pytest.skip('A real slide renderer is required.')
    p=Presentation();p.slides.add_slide(p.slide_layouts[6]);s=p.slides.add_slide(p.slide_layouts[6])
    data=CategoryChartData();data.categories=['A','B'];data.add_series('Positive',[4,8]);data.add_series('Negative',[8,4])
    chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(1),Inches(1),Inches(6),Inches(3),data).chart
    for series,color in zip(chart.series,[MSO_THEME_COLOR.ACCENT_1,MSO_THEME_COLOR.ACCENT_2]):
        series.format.fill.solid();series.format.fill.fore_color.theme_color=color
    source=tmp_path/'source.pptx';candidate=tmp_path/'candidate.pptx';p.save(source)
    preserve.build(source,candidate,grounded.TEMPLATE_PATH)
    with fitz.open(rendering.pptx_to_pdf(candidate,tmp_path/'chart-render')) as doc:
        pix=doc[1].get_pixmap()
        colors=Image.frombytes('RGB',(pix.width,pix.height),pix.samples).getdata()
        # PDF color conversion can round one channel by a level. A two-level
        # tolerance still clearly distinguishes the original and template palettes.
        assert sum(max(abs(a-b) for a,b in zip(c,(79,129,189)))<=2 for c in colors)>100
        assert sum(max(abs(a-b) for a,b in zip(c,(192,80,77)))<=2 for c in colors)>100
    bullets=['Apply interdisciplinary methodologies to strengthen collaborative engineering decisions.',
             'Evaluate implementation alternatives through systematic experimentation and reflection.']
    deck=DeckSpec(slides=[SlideSpec(id='a',title='Topic'),SlideSpec(id='b',title='Thank you!',bullets=bullets)])
    closing=tmp_path/'closing.pptx';composer.compose(deck,closing,tmp_path/'assets',kinds=['opening','closing'])
    body=next(s for s in Presentation(closing).slides[-1].shapes if s.name=='authored-body')
    with fitz.open(rendering.pptx_to_pdf(closing,tmp_path/'closing-render')) as doc:
        spans=[span for block in doc[-1].get_text('dict')['blocks'] if 'lines' in block
               for line in block['lines'] for span in line['spans'] if span['bbox'][1]>body.top/12700]
        assert ' '.join(' '.join(span['text'] for span in spans).split())==' '.join(bullets)
        assert max(span['bbox'][3] for span in spans)<(body.top+body.height)/12700
