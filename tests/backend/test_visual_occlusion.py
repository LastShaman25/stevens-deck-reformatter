"""Image bounds alone are not evidence that their visible pixels hide text."""
from io import BytesIO

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.qa.style_audit import audit


def specimen(*, point=None, color='363D45', image_above=False, rotation=0,
             crop_left=0, alpha=255):
    image=Image.new('RGBA',(400,300),(255,255,255,alpha))
    draw=ImageDraw.Draw(image)
    draw.rectangle((0,0,399,299),outline='red',width=4)
    draw.rectangle((100,120,180,250),fill='blue')
    if point: draw.point(point,fill='black')
    blob=BytesIO();image.save(blob,format='PNG');blob.seek(0)
    prs=Presentation();slide=prs.slides.add_slide(prs.slide_layouts[6])
    def add_text():
        text=slide.shapes.add_textbox(Inches(2),Inches(1.6),Inches(3),Inches(.4))
        run=text.text_frame.paragraphs[0].add_run();run.text='Forward Pass - MLP'
        run.font.name='Arial';run.font.size=Pt(18)
        if color: run.font.color.rgb=RGBColor.from_string(color)
        return text
    if image_above: text=add_text()
    picture=slide.shapes.add_picture(blob,Inches(1),Inches(1),Inches(6),Inches(4))
    picture.crop_left=crop_left;picture.rotation=rotation
    if not image_above: text=add_text()
    return prs,slide,picture,text


def occlusions(prs,slide):
    return [f for f in audit(slide,prs.slide_width,prs.slide_height)
            if f['type']=='visual_occlusion']


def test_native_title_over_proven_white_image_padding_is_not_occluded():
    prs,slide,_,_=specimen()
    assert not occlusions(prs,slide)


@pytest.mark.parametrize('kwargs',[
    {'point':(150,50)},                  # Even one visible artwork pixel stays reviewable.
    {'point':(66,50)},                   # Include antialiasing immediately outside text bounds.
    {'color':'FFFFFF'},                 # White on white is not legible.
    {'color':None},                     # Unknown inherited color is not affirmative evidence.
    {'image_above':True},               # Opaque image above native text hides it.
    {'rotation':10},                    # Unknown transformed pixel coordinates stay reviewable.
    {'alpha':0},                        # Transparent pixels do not prove the underlying color.
    {'crop_left':.5,'point':(300,50)},   # Map source pixels through the actual picture crop.
])
def test_visible_or_uncertain_image_text_intersections_remain_flagged(kwargs):
    prs,slide,_,_=specimen(**kwargs)
    assert len(occlusions(prs,slide))==1


def test_artwork_outside_visible_picture_crop_does_not_create_collision():
    prs,slide,_,_=specimen(crop_left=.5,point=(150,50))
    assert not occlusions(prs,slide)


def test_flipped_and_effect_processed_images_remain_reviewable():
    from pptx.oxml.xmlchemy import OxmlElement
    prs,slide,picture,_=specimen()
    picture._element.spPr.xfrm.set('flipH','1')
    assert occlusions(prs,slide)
    picture._element.spPr.xfrm.attrib.pop('flipH')
    picture._element.blipFill.blip.append(OxmlElement('a:grayscl'))
    assert occlusions(prs,slide)


def test_grouped_geometry_is_not_assumed_to_use_root_picture_coordinates():
    prs,slide,picture,text=specimen()
    slide.shapes.add_group_shape([picture,text])
    assert occlusions(prs,slide)
