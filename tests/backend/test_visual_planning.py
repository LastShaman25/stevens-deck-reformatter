from pathlib import Path
import pytest
from pptx import Presentation
from app.authoring import service, composer
from app.authoring.models import (
    CreationRequest, Outline, OutlineSlide, VisualPlan, SlideSpec, DeckSpec,
    DiagramSpec, DiagramNode, TableSpec,
)


def planned(kind='diagram'):
    return OutlineSlide(id='body', title='Review process', points=['Draft, review, release'],
        visual=VisualPlan(kind=kind, description='Draft then review then release.',
                          reason='Shows the order of work.'))


def diagram():
    return DiagramSpec(kind='process', nodes=[DiagramNode(label=s, detail='One supported step.')
                                             for s in ('Draft', 'Review', 'Release')])


def test_native_table_header_contrast_survives_template_and_save(tmp_path):
    from app import brand
    path=tmp_path/'table.pptx'
    deck=DeckSpec(slides=[SlideSpec(id='opening',title='Review'),
        SlideSpec(id='body',title='Review results',table=TableSpec(headers=['Result','Passed'],rows=[['First','9 of 12']])),
        SlideSpec(id='closing',title='Thank you!')])
    manifest=composer.compose(deck,path,tmp_path/'assets',kinds=['opening','content','closing'])
    assert composer.audit(path,manifest)['status']=='passed'
    table=next(s.table for s in Presentation(path).slides[1].shapes if s.has_table)
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            assert str(cell.fill.fore_color.rgb)==(brand.RED if ri==0 else brand.WHITE)
            assert all(str(p.font.color.rgb)==(brand.WHITE if ri==0 else brand.INK) for p in cell.text_frame.paragraphs)


def test_missing_visual_plan_and_missing_planned_visual_retry(monkeypatch):
    sess=service.create(CreationRequest(topic='Explain a process', audience='Students'))
    expected=service.with_bookends(Outline(title='Review', rationale='A process benefits from a diagram.', slides=[planned()])).model_dump()
    for slide in expected['slides']:
        slide['unit_ids']=['process'] if slide['kind']=='content' else []
    import copy
    missing=copy.deepcopy(expected); missing['slides'][1]['visual']=None
    content={'units':[{'id':'process','section':'Review','objective':'Explain draft, review, release.',
                      'kind':'concept','priority':'essential','source_pages':[]}], 'context_pages':[]}
    responses=[content, missing, expected]
    payloads=[]
    def provider(*args, **kwargs):
        payloads.append(args[2]); return {'status':'completed','data':responses.pop(0)}
    monkeypatch.setattr(service.providers,'generate',provider)
    service.plan(sess)
    assert 'explicit visual decision' in payloads[2]['repair_instruction']
    item=Outline.model_validate(sess.creation['outline']).slides[1]
    assert item.visual.kind=='diagram'
    responses.extend([SlideSpec(id='body', title=item.title).model_dump(),
                      SlideSpec(id='body', title=item.title, diagram=diagram()).model_dump()])
    result=service.call(sess,'author','Write the slide.',{},SlideSpec,
                        validate=lambda slide:service.validate_visual_content(item,slide))
    assert result.diagram and 'approved diagram' in payloads[-1]['repair_instruction']


def test_visual_rules_prevent_silent_removal_and_invalid_source_figure():
    before=SlideSpec(id='body',title='Process',diagram=diagram())
    with pytest.raises(ValueError,match='retain'):
        service.validate_repair_visual(before,SlideSpec(id='body',title='Process'))
    service.validate_repair_visual(before,before)
    with pytest.raises(ValueError,match='source page'):
        service.validate_visual_plans(Outline(title='x',rationale='x',slides=[planned('source_figure')]))
    item=planned('source_figure');item.source_pages=[2]
    with pytest.raises(ValueError,match='assigned'):
        service.validate_visual_content(item,SlideSpec(id='body',title='Process',figure_page=1))
    item.kind='opening'
    with pytest.raises(ValueError,match='template artwork'):
        service.validate_visual_plans(Outline(title='x',rationale='x',slides=[item]))


@pytest.mark.parametrize('mutation',['label','arrow'])
@pytest.mark.parametrize('node_count',[3,6])
def test_native_diagram_is_verified_after_save(tmp_path,mutation,node_count):
    path=tmp_path/'diagram.pptx'
    visual=diagram()
    if node_count==6:visual.nodes+= [DiagramNode(label=f'Step {i}',detail='A short description.') for i in range(4,7)]
    deck=DeckSpec(slides=[SlideSpec(id='opening',title='Review'),
        SlideSpec(id='body',title='Review process',bullets=['Release follows review.'],diagram=visual),
        SlideSpec(id='closing',title='Thank you!')])
    manifest=composer.compose(deck,path,tmp_path/'assets',kinds=['opening','content','closing'])
    assert composer.audit(path,manifest)['status']=='passed'
    prs=Presentation(path); slide=prs.slides[1]
    target=next(s for s in slide.shapes if s.name==f'authored-diagram-{mutation}-0')
    if mutation=='label':target.text='Unapproved meaning'
    else:target._element.getparent().remove(target._element)
    prs.save(path)
    assert composer.audit(path,manifest)['status']=='failed'
