"""Opt-in diagnostic replay of a budget-blocked final review, without changing the job.

Reuses this exact artifact's recorded batch responses to reconstruct the final
request. Only deck synthesis is sent live. This does not approve or release a deck.
"""
import argparse
from copy import deepcopy
import json
from pathlib import Path
import sys
import time
from types import SimpleNamespace

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from pptx import Presentation
from app import sessions
from app.ai import output_qa, providers
from slide_engine.inventory import sha256


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--record',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--live',action='store_true',required=True)
    args=parser.parse_args()
    if args.output.exists():parser.error('Choose a fresh report path.')
    original=json.loads(args.record.read_text(encoding='utf-8')); record=deepcopy(original)
    root=args.record.parents[2];source=root/'source.pptx'
    assert sha256(source)==record['source_sha256']
    assert sha256(record['candidate'])==record['candidate_sha256']
    batches={tuple(b['request_ordinals']):b for b in original['output_qa']['batches']}
    deadline=time.time()+600
    def active():
        if time.time()>=deadline:raise ValueError('Diagnostic replay expired.')
    sess=SimpleNamespace(source_path=str(source),preview_path=lambda i,v:str(root/'previews'/f'{v}-{i}.png'),
        ensure_active=active,execution_deadline=deadline,calls=original['usage']['upload_requests'],
        tokens=original['usage']['upload_tokens'])
    providers.reserve_output_qa(sess,record['report']['slide_count'],redesign=True)
    prs=Presentation(source)
    evidence={'source_slides':[{'ordinal':i+1,'text':'\n'.join(s.text for s in slide.shapes if s.has_text_frame),
        'tables':[[[c.text for c in row.cells] for row in s.table.rows] for s in slide.shapes if s.has_table],
        'charts':[{'series':[{'name':v.name,'values':list(v.values)} for v in s.chart.series]} for s in slide.shapes if s.has_chart],
        'notes':slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''} for i,slide in enumerate(prs.slides)],
        'source_to_output_slides':record['source_to_output_slides']}
    send=providers.generate;live=[]
    def replay(role,system,payload,**kwargs):
        if payload['stage']=='slide_review':
            batch=batches[tuple(payload['expected'])]
            response=deepcopy(batch['response'])
            # The saved receipt includes this application-derived field; it is not model output.
            for finding in response['findings']:finding.pop('category',None)
            return {'status':'completed','data':response,'recorded_batch_reused':True}
        if record['output_manifest_sha256']!=original['output_manifest_sha256']:
            raise ValueError('Reconstructed output manifest differs from the recorded audits; fresh batch reviews are required.')
        estimate=providers.request_token_estimate(system,payload,kwargs.get('images',()),kwargs.get('max_tokens',16000))
        result=send(role,system,payload,**kwargs)
        live.append({'status':result['status'],'message':result.get('message'),'estimated_tokens':estimate,
                     'request_attempts':result.get('request_attempts'),'usage':result.get('usage')})
        return result
    providers.generate=replay
    token=sessions.active_session.set(sess)
    try:checks=output_qa.run(sess,record,evidence)
    finally:
        providers.generate=send;sessions.active_session.reset(token)
    summary={'starting_usage':original['usage'],'token_limit':providers.token_limit(sess),'request_limit':providers.request_limit(),
        'recorded_batches_reused':len(original['output_qa']['batches']),'live_final_requests':live,
        'checks':{k:v['status'] for k,v in checks.items()},
        'errors':list({f['message'] for value in checks.values() if value['status']=='error' for f in value.get('findings',[])}),
        'candidate_unchanged':sha256(record['candidate'])==original['candidate_sha256'],
        'source_unchanged':sha256(source)==original['source_sha256'],'active_job_modified':False}
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps(summary),flush=True)
    return 0 if live and all(r['status']=='completed' for r in live) and all(v['status']!='error' for v in checks.values()) else 1


if __name__=='__main__':raise SystemExit(main())
