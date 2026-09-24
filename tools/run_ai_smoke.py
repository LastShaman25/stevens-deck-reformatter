"""Explicit synthetic AI acceptance run; never reads the private course corpus."""
import argparse
import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'backend')]
from pptx import Presentation
from pptx.util import Inches, Pt
from app import generations, grounded, rendering, sessions
from app.ai import providers


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    mode=parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--live',action='store_true',help='Call configured providers with synthetic content; API charges may apply.')
    mode.add_argument('--mock-providers',action='store_true',help='Test native edits and real rendering with simulated AI responses.')
    parser.add_argument('--output',type=Path,required=True,help='New evidence directory; must not exist.')
    args=parser.parse_args()
    if args.output.exists():parser.error('Output directory already exists; choose a new directory.')
    if args.live and not providers.capabilities()['configured']:
        print('Live checkpoint blocked: configure provider keys in backend/.env.');return 2
    args.output.mkdir(parents=True)
    if args.mock_providers:
        config={'configured':True,'independent_providers':False,
                **{role:{'provider':'mock','model':'synthetic-test','configured':True} for role in ('planner','reviewer')}}
        providers.capabilities=lambda:config
        def mock(role,system,payload,images=(),max_tokens=0):
            if role=='planner':
                objects=[{'id':o['id'],**dict(zip(('x','y','w','h'),o['box']))} for o in payload['objects']]
                for obj in objects:obj['x']+=.04
                data={'layout':'Synthetic horizontal shift','rationale':'Exercise native edits without live AI','objects':objects}
            else:data={'verdict':'passed','summary':'Simulated review; not evidence of live model quality','findings':[]}
            return {'status':'completed','provider':'mock','model':'synthetic-test','data':data}
        providers.generate=mock
    s=sessions.create()
    try:
        p=Presentation();slide=p.slides.add_slide(p.slide_layouts[6])
        for value,y,h,size in [('Synthetic project update',.3,.8,30),
                             ('Research completed: 12 interviews\nPrototype ready for review\nNext milestone: usability testing',2,2.5,22)]:
            shape=slide.shapes.add_textbox(Inches(1),Inches(y),Inches(7),Inches(h));shape.text=value
            for para in shape.text_frame.paragraphs:
                for run in para.runs:run.font.name='Arial';run.font.size=Pt(size)
        slide.notes_slide.notes_text_frame.text='Synthetic notes must survive the redesign.'
        p.save(s.source_path);s.analysis=grounded.analyze(s.source_path)
        rendering.render_to_pdf(s.source_path,s.source_pdf)
        rendering.rasterize_pdf(s.source_pdf,lambda i:s.preview_path(i,'before'))
        s.revisions={'0':{'instruction':'Create a clear title and readable body with generous spacing.'}}
        record=generations.build(s,mode='ai',repair_passes=1)
        shutil.copytree(record['directory'],args.output/'generation')
        shutil.copyfile(s.source_path,args.output/'source.pptx')
        shutil.copyfile(s.preview_path(0,'before'),args.output/'source.png')
        result={'recorded_at':datetime.now(timezone.utc).isoformat(),
                'providers':'live' if args.live else 'mock','renderer':'real',
                'generation':generations.public(record),'human_approval':'not_performed'}
        (args.output/'result.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
        print(json.dumps({'providers':result['providers'],'state':record['state'],
                          'checks':result['generation']['checks'],'evidence':str(args.output.resolve())},indent=2))
        return 0 if record['state']=='ready' else 3
    finally:sessions.delete(s.id)


if __name__=='__main__':raise SystemExit(main())
