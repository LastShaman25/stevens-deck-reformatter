"""Opt-in live checkpoint using synthetic text and bundled template art only."""
import argparse
from io import BytesIO
import json
from pathlib import Path
import sys
from types import SimpleNamespace
from zipfile import ZipFile

ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT),str(ROOT/'backend')]
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches,Pt
from app import grounded
from app.ai import providers,source_decisions,layout,pipeline,rubric
from app.qa import artifact_coverage,render_verify
from slide_engine import inventory,template_policy as T


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live',action='store_true',required=True,help='Send synthetic images to configured providers; API charges apply.')
    parser.add_argument('--output',type=Path,required=True,help='New local evidence directory.')
    args=parser.parse_args()
    if args.output.exists():parser.error('Choose a new evidence directory.')
    if not providers.capabilities()['configured']:parser.error('Configure backend/.env first.')
    directory=args.output.resolve();directory.mkdir(parents=True)
    source=directory/'source.pptx';candidate=directory/'candidate.pptx'
    p=Presentation();p.slide_width=Inches(13.333333);p.slide_height=Inches(7.5)
    slide=p.slides.add_slide(p.slide_layouts[6])
    with ZipFile(grounded.TEMPLATE_PATH) as z:
        im=Image.open(BytesIO(z.read('ppt/media/image2.png'))).convert('RGBA')
        im.thumbnail((1920,1080));stream=BytesIO();im.save(stream,format='PNG')
    slide.shapes.add_picture(stream,0,0,p.slide_width,p.slide_height)
    for text,y,size in [('Synthetic project update',3.5,32),('A logo extraction checkpoint',4.5,20),('Test presenter',6.2,18)]:
        sh=slide.shapes.add_textbox(Inches(7.2),Inches(y),Inches(5.6),Inches(.9));sh.text=text
        for run in sh.text_frame.paragraphs[0].runs:
            run.font.name='Arial';run.font.size=Pt(size);run.font.color.rgb=RGBColor(255,255,255)
    p.save(source)
    before=render_verify.check(source,directory/'original-render')
    template=Presentation(grounded.TEMPLATE_PATH)
    for sid,s in list(zip(template.slides._sldIdLst,template.slides)):
        if s.slide_layout.name=='1_Title Slide':continue
        template.part.drop_rel(sid.rId);template.slides._sldIdLst.remove(sid)
    template_path=directory/'template-reference.pptx';template.save(template_path)
    refs=render_verify.check(template_path,directory/'template-render')
    reference=('APPROVED TEMPLATE: 1_Title Slide',Path(refs['pages'][0]['png']))
    calls=[]
    def generate(role,system,payload,images,max_tokens):
        response=providers.generate(role,system,payload,images,max_tokens=max_tokens)
        calls.append({'stage':payload.get('stage',role),'validation_error':payload.get('validation_error'),**response})
        (directory/'calls.json').write_text(json.dumps(calls,indent=2),encoding='utf-8')
        print(json.dumps({'stage':payload.get('stage',role),'status':response['status']}),flush=True)
        return response
    sess=SimpleNamespace(source_path=str(source),revisions={'0':{'instruction':
        'Use the required first-page template, remove the obsolete source background, and preserve its required logo with the verified extraction tool. Reuse the identical template logo if verified, rather than adding a duplicate.'}},
        ensure_active=lambda:None,preview_path=lambda *args:before['pages'][0]['png'])
    result={'rubric_version':rubric.VERSION,'providers':'live','renderer':'PowerPoint','human_approval':'not_performed',
            'configuration':providers.capabilities()}
    try:
        decisions,_=source_decisions.run(sess,lambda **kwargs:None,generate=generate,template_images=[reference])
        report=grounded.build_deck(source,candidate,source_decisions=decisions)
        result['decisions']=decisions
        result['extractions']=sum(len(d.get('logo_extractions',[])) for d in decisions.values())
        result['artifact_coverage']=artifact_coverage.audit(source,candidate,report)
        result['render']=render_verify.check(candidate,directory/'candidate-render')
        final=Presentation(candidate).slides[0]
        payload={'stage':'paired_final_review','source_decision':decisions['0'],
            'required_objects':layout.describe(final),'template_context':layout.template_context(final),
            'template_contract':T.contract(final),'original_objects':source_decisions.source_objects(Presentation(source),0,inventory.sha256(source)),
            'schema':pipeline.VisualReview.model_json_schema()}
        review=generate('reviewer',pipeline.REVIEW_SYSTEM,payload,
            [('ORIGINAL source',Path(before['pages'][0]['png'])),reference,
             ('FINAL CANDIDATE TO AUDIT (last image)',Path(result['render']['pages'][0]['png']))],max_tokens=10000)
        if review['status']!='completed':raise ValueError('Final live review did not complete.')
        validated=pipeline.VisualReview.model_validate(review['data'])
        rubric.validate_checks(validated.rubric,validated.findings)
        result['paired_review']=validated.model_dump()
        result['passed']=(result['extractions']>0 and result['artifact_coverage']['status']=='passed'
                          and result['render']['status']=='passed' and validated.verdict=='passed' and not validated.findings)
    except Exception as exc:
        result.update(passed=False,error=str(exc)[:1500])
    (directory/'result.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps({'passed':result['passed'],'error':result.get('error'),'extractions':result.get('extractions'),'evidence':str(directory)},indent=2))
    return 0 if result['passed'] else 3


if __name__=='__main__':raise SystemExit(main())
