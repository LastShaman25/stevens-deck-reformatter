"""Opt-in live authoring reproduction: unprompted visual decisions, actual artifacts and QA."""
import argparse
import json
import shutil
import sys
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'backend'),str(ROOT)]
from app import sessions, generations
from app.authoring import service
from app.authoring.models import CreationRequest, Outline, DeckSpec
from pptx import Presentation
from slide_engine import template_policy as T, bookends


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live',action='store_true',required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--replay',type=Path,help='Replay an earlier run’s exact outline and authored content to verify its repair.')
    args=parser.parse_args()
    if args.output.exists():parser.error('Use a fresh output directory.')
    args.output.mkdir(parents=True)
    # Deliberately does not request any visual type; the planner must decide.
    sess=service.create(CreationRequest(audience='Instructors learning a publication workflow',length_preference='brief',
        topic='Explain this synthetic slide publication workflow: draft slides, review every slide, '
              'correct failing slides, recheck, then release only when all pass. '
              'Compare these supplied example results: first review 12 slides, 9 passed, 3 needed corrections; '
              'after correction 12 passed, 0 needed corrections. These numbers are illustrative, not research. '
              'Explain the order and what the two review results mean. Do not add other claims.'))
    try:
        with sessions.job(sess):
            deck=None
            if args.replay:
                service.save_outline(sess,Outline.model_validate_json((args.replay/'outline.json').read_text(encoding='utf-8')),0)
                previous=Presentation(args.replay/'generation'/'candidate.pptx')
                deck=DeckSpec(slides=[json.loads(s.notes_slide.notes_text_frame.text.split('Authoring specification:\n',1)[1]) for s in previous.slides])
            else:service.plan(sess)
            outline=sess.creation['outline']
            (args.output/'outline.json').write_text(json.dumps(outline,indent=2),encoding='utf-8')
            visuals=[s['visual']['kind'] for s in outline['slides'] if s['kind']=='content']
            print('PLANNED '+json.dumps(visuals),flush=True)
            if not any(v!='text_only' for v in visuals):raise RuntimeError('Planner chose no visuals for a process and numerical comparison.')
            service.approve(sess,sess.creation['revision'])
            service.generate(sess,deck)
            record=sess.generation
            shutil.copytree(record['directory'],args.output/'generation')
            prs=Presentation(record['candidate'])
            closing=T.is_closing(prs.slides[-1]) and bookends.has_thanks(prs.slides[-1])
            actual=[next((k for k in ('diagram','chart','plot','equation','table','figure_page') if s.get(k) is not None),'text_only')
                    for s in sess.creation['deck']['slides']]
            summary={'state':record['state'],'download_allowed':generations.download_allowed(record),
                     'planned_visuals':visuals,'actual_visuals':actual,'closing_present':bool(closing),
                     'slide_count':len(prs.slides),'checks':{k:v['status'] for k,v in record['checks'].items()},
                     'repair_count':len(record.get('repair_history',[])),
                     'findings':record['findings'],'calls':sess.calls,'tokens':sess.tokens}
            (args.output/'summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
            print(json.dumps(summary),flush=True)
            return 0 if summary['download_allowed'] and closing else 1
    finally:sessions.delete(sess.id)


if __name__=='__main__':raise SystemExit(main())
